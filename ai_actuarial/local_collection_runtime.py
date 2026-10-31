"""Restartable local collection and publication runtime for the 32 site plan."""

from __future__ import annotations

import argparse
import base64
import contextlib
import datetime as dt
import hashlib
import html
import http.client
import importlib.util
import json
import os
import re
import shlex
import shutil
import socket
import ssl
import subprocess
import sys
import tempfile
import time
import xml.etree.ElementTree as ET
from dataclasses import dataclass
from html.parser import HTMLParser
from pathlib import Path
from types import SimpleNamespace
from typing import Any, Callable, Iterator
from urllib.parse import parse_qsl, urlencode, urljoin, urlsplit, urlunsplit
from zoneinfo import ZoneInfo

from ai_actuarial.collectors.web_page import WebPageCollector
from ai_actuarial.content_hub import (
    apply_content_hub_schema,
    content_hub_ai_url_versions,
    content_hub_list_ai_versions,
    content_hub_read_version,
    content_hub_schema_plan,
    content_hub_status,
    content_hub_url_exclusion,
    ingest_live_article,
    save_live_enrichment,
)
from ai_actuarial.crawler import Crawler
from ai_actuarial.postgres import DatabaseConfig
from ai_actuarial.security import resolve_safe_http_url

PROJECT_ROOT = Path(__file__).resolve().parents[1]
SKILL_ROOT = PROJECT_ROOT / ".agents" / "skills" / "local-actuarial-flow"
FLOW_SCRIPT = SKILL_ROOT / "scripts" / "local_flow.py"
TARGET_WORKER = Path(__file__).with_name("local_collection_target_worker.py")
TIME_ZONE = ZoneInfo("America/New_York")
MAX_SITE_REQUESTS = 30
MAX_SWISS_SITE_REQUESTS = 60
MAX_DAILY_SITE_RUNS = 8
MAX_RECONCILE_PAGE = 100
MAX_ENRICHMENTS_PER_SITE_RUN = 8
MAX_CHEAP_ARTICLES_PER_SITE_RUN = 24
MAX_ARTICLE_SCAN_PER_SITE_RUN = MAX_CHEAP_ARTICLES_PER_SITE_RUN + MAX_ENRICHMENTS_PER_SITE_RUN
ALGOLIA_HOST = "54YLV1QRFB-dsn.algolia.net"
AUSTRALIA_RESEARCH_TYPES = (
    "Publication > Report",
    "Publication > Dialogue Paper",
    "Publication > Discussion Paper",
    "Publication > Public Policy Position Statement",
    "Papers",
    "Professional Practice Document > Technical Paper",
    "Professional Practice Document > Information Note",
    "Professional Practice Document > Practice Guidelines",
)
MAX_ARTICLE_BYTES = 40 * 1024 * 1024
MAX_MARKDOWN_CHARS = 500_000
MAX_EVIDENCE_QUOTE_CHARS = 240
MAX_EVIDENCE_SPAN_CHARS = 200
EVIDENCE_SPAN_OVERLAP_CHARS = 40
EVIDENCE_SPAN_ID_PATTERN = r"^span-[0-9a-f]{16}$"
RECIPE = "gpt-6-luna-high-local-live-v1"

_FILE_SUFFIXES = {".pdf", ".docx", ".doc", ".pptx", ".ppt", ".xlsx", ".xls", ".txt", ".rtf"}
_PATH_HINTS = (
    "publication",
    "research",
    "report",
    "paper",
    "study",
    "document",
    "technical",
    "standard",
    "guideline",
    "journal",
    "resource",
    "white-paper",
    "whitepaper",
    "bulletin",
    "analysis",
    "library",
    "download",
    "upload",
    "proceeding",
    "actuarial",
    "practice",
)
_EXCLUDE_HINTS = (
    "recruit",
    "vacanc",
    "career",
    "tender",
    "rfp",
    "privacy",
    "login",
    "account",
    "donat",
)
_IAA_ADMIN_DOCUMENT_NAMES = frozenset(
    {
        "statutes.pdf",
        "statutes_fr.pdf",
        "2012-subsidiarity-principle-en.pdf",
        "2012-subsidiarity-principle-fr.pdf",
        "internalregulations.pdf",
    }
)


@dataclass(frozen=True)
class Site:
    site_id: str
    action: str
    seeds: tuple[str, ...]
    hosts: tuple[str, ...]
    delay: float = 0.5
    max_depth: int = 2
    page_hints: tuple[str, ...] = _PATH_HINTS


SITES: dict[str, Site] = {
    "iaa": Site(
        "iaa",
        "collect",
        (
            "https://actuaries.org/publications/papers/",
            "https://actuaries.org/publications/papers/?paged=2",
            "https://actuaries.org/paper/professional-considerations-for-actuaries-interacting-with-artificial-intelligence/",
        ),
        ("actuaries.org",),
        0.5,
    ),
    "iaba": Site(
        "iaba",
        "monitor_only",
        ("https://www.blackactuaries.org/sitemap_index.xml",),
        ("blackactuaries.org",),
        0.6,
    ),
    "soa": Site(
        "soa",
        "collect",
        (
            "https://www.soa.org/research/topics/artificial-intelligence-topic-landing/",
            "https://www.soa.org/research/research-topic-list/",
            "https://www.soa.org/research/soa-research/",
            "https://www.soa.org/resources/research-reports/2025/research-ai-bulletin/",
        ),
        ("soa.org",),
        0.5,
    ),
    "soa-ai-topic-landing": Site(
        "soa-ai-topic-landing",
        "collect",
        ("https://www.soa.org/research/topics/artificial-intelligence-topic-landing/",),
        ("soa.org",),
        2.0,
    ),
    "soa-ai-bulletin-focused": Site(
        "soa-ai-bulletin-focused",
        "collect",
        ("https://www.soa.org/resources/research-reports/2025/research-ai-bulletin/",),
        ("soa.org",),
        2.0,
        1,
    ),
    "cas": Site(
        "cas",
        "collect",
        (
            "https://www.casact.org/publications-research",
            "https://www.casact.org/sitemap.xml",
            "https://forum.casact.org/sitemap.xml",
            "https://variancejournal.org/sitemap.xml",
        ),
        ("casact.org", "variancejournal.org"),
        0.6,
    ),
    "aaa": Site(
        "aaa",
        "collect",
        (
            "https://actuary.org/sitemap_index.xml",
            "https://actuary.org/resource-center/",
            "https://actuary.org/policy-research/",
            "https://actuary.org/practice-area/health/",
        ),
        ("actuary.org",),
        0.6,
    ),
    "cia-ica": Site(
        "cia-ica",
        "collect",
        (
            "https://www.cia-ica.ca/sitemap_index.xml",
            "https://www.cia-ica.ca/publications",
            "https://www.cia-ica.ca/wp-json/wp/v2/publication?per_page=50&page=1",
        ),
        ("cia-ica.ca",),
        0.5,
    ),
    "cia-research-publications": Site(
        "cia-research-publications",
        "collect",
        (
            "https://www.cia-ica.ca/publications",
            "https://www.cia-ica.ca/wp-json/wp/v2/publication?per_page=50&page=1",
        ),
        ("cia-ica.ca",),
        0.5,
    ),
    "cca": Site(
        "cca",
        "collect",
        (
            "https://www.ccactuaries.org/professional-resources",
            "https://www.ccactuaries.org/professional-resources/archives",
            "https://www.ccactuaries.org/professional-resources/papers",
            "https://www.ccactuaries.org/professional-resources/technical-resource-library-main",
        ),
        ("ccactuaries.org",),
        0.5,
    ),
    "aae": Site(
        "aae",
        "collect",
        (
            "https://actuary.eu/sitemap.xml",
            "https://actuary.eu/papers/",
            "https://actuary.eu/publications/the-european-actuary/",
            "https://actuary.eu/the-european-actuary-issues/",
        ),
        ("actuary.eu",),
        0.6,
    ),
    "eaa": Site(
        "eaa",
        "monitor_only",
        ("https://actuarial-academy.com/en/sitemap.xml",),
        ("actuarial-academy.com",),
        0.6,
    ),
    "ifoa-uk": Site(
        "ifoa-uk",
        "collect",
        (
            "https://actuaries.org.uk/sitemap.xml",
            "https://actuaries.org.uk/insight/",
            "https://actuaries.org.uk/standards/",
            "https://actuaries.org.uk/about-us/",
        ),
        ("actuaries.org.uk",),
        0.6,
    ),
    "france": Site(
        "france",
        "collect",
        (
            "https://www.institutdesactuaires.com/se-documenter/memoires/memoires-d-actuariat-4651",
            "https://www.institutdesactuaires.com/se-documenter/supports-des-presentations-39",
        ),
        ("institutdesactuaires.com",),
        0.5,
    ),
    "dav": Site(
        "dav",
        "collect",
        ("https://aktuar.de/de/sitemap.xml", "https://aktuar.de/en/sitemap.xml"),
        ("aktuar.de",),
        0.6,
    ),
    "swiss-association-of-actuaries": Site(
        "swiss-association-of-actuaries",
        "collect",
        (
            "https://www.actuaries.ch/de/download",
            "https://www.actuaries.ch/fr/download",
            "https://www.actuaries.ch/en/download",
            "https://www.actuaries.ch/it/download",
            "https://www.actuaries.ch/de/download-archiv",
        ),
        ("actuaries.ch",),
        0.6,
    ),
    "italy": Site(
        "italy",
        "collect",
        (
            "https://ordineattuari.it/linee-guida/",
            "https://ordineattuari.it/linee-guida/documenti-tecnici/",
            "https://ordineattuari.it/linee-guida/position-paper/",
            "https://ordineattuari.it/pubblicazioni/",
        ),
        ("ordineattuari.it",),
        0.6,
    ),
    "istituto-italiano-degli-attuari": Site(
        "istituto-italiano-degli-attuari", "hold", (), ("italian-actuaries.org",), 0.5
    ),
    "spain": Site(
        "spain",
        "collect",
        (
            "https://actuarios.org/sitemap_index.xml",
            "https://actuarios.org/document-category/",
        ),
        ("actuarios.org",),
        0.6,
    ),
    "australia": Site(
        "australia",
        "collect",
        (
            "https://www.actuaries.asn.au/research-analysis/thought-leadership",
            "https://www.actuaries.asn.au/research-analysis",
            "https://www.actuaries.asn.au/research-analysis/public-policy-approach",
        ),
        ("actuaries.asn.au", "content.actuaries.asn.au"),
        0.6,
    ),
    "iai": Site(
        "iai",
        "collect",
        (
            "https://actuariesindia.org/magazine",
            "https://actuariesindia.org/annual-report",
            "https://actuariesindia.org/mortality-table",
            "https://actuariesindia.org/actuarial-practice-standards",
            "https://actuariesindia.org/regulations",
            "https://actuariesindia.org/library-services",
        ),
        ("actuariesindia.org",),
        1.0,
    ),
    "iaj": Site(
        "iaj",
        "collect",
        (
            "https://www.actuaries.jp/english/lib/index.html",
            "https://www.actuaries.jp/english/lib/journal/",
            "https://www.actuaries.jp/english/lib/y_ronbun/",
            "https://www.actuaries.jp/english/lib/IT_ronbun/",
        ),
        ("actuaries.jp",),
        0.7,
    ),
    "institute-of-actuaries-of": Site(
        "institute-of-actuaries-of",
        "collect",
        (
            "https://www.actuary.or.kr/main/main.do",
            "https://www.actuary.or.kr/work",
            "https://www.actuary.or.kr/regulation/report",
        ),
        ("actuary.or.kr",),
        0.6,
    ),
    "aict": Site(
        "aict",
        "collect",
        (
            "https://www.airc.org.tw/rule/101",
            "https://www.airc.org.tw/rule/201",
            "https://www.airc.org.tw/rule/202",
            "https://en.airc.org.tw/rule/101",
            "https://www.airc.org.tw/page/101",
            "https://www.airc.org.tw/re/pdf/",
        ),
        ("airc.org.tw",),
        0.6,
    ),
    "caa": Site("caa", "collect", ("https://www.e-caa.org.cn/",), ("e-caa.org.cn",), 0.5),
    "ashk": Site("ashk", "collect", ("https://www.actuaries.org.hk/",), ("actuaries.org.hk",), 0.6),
    "assa": Site(
        "assa",
        "collect",
        (
            "https://www.actuarialsociety.org.za/wp-sitemap.xml",
            "https://www.actuarialsociety.org.za/document-category/",
            "https://www.actuarialsociety.org.za/sa-actuarial-journal/",
            "https://www.actuarialsociety.org.za/public-benefits-research/",
        ),
        ("actuarialsociety.org.za",),
        0.5,
    ),
    "task": Site(
        "task",
        "collect",
        ("https://www.actuarieskenya.or.ke/resources",),
        ("actuarieskenya.or.ke",),
        0.5,
        1,
    ),
    "ama": Site(
        "ama",
        "collect",
        (
            "https://www.ama.ma/sitemap.xml",
            "https://www.ama.ma/en/",
            "https://www.ama.ma/en/documents/",
        ),
        ("ama.ma",),
        0.5,
    ),
    "iba-brazil": Site(
        "iba-brazil",
        "collect",
        (
            "https://atuarios.org.br/wp-sitemap.xml",
            "https://atuarios.org.br/wp-json/wp/v2/posts?per_page=50&page=1",
        ),
        ("atuarios.org.br",),
        0.5,
    ),
    "aca-colombia": Site(
        "aca-colombia",
        "collect",
        (
            "https://actuarios.org.co/wp-sitemap.xml",
            "https://actuarios.org.co/",
        ),
        ("actuarios.org.co",),
        0.6,
    ),
    "conac-mexico": Site(
        "conac-mexico",
        "collect",
        (
            "https://www.conacmexico.org.mx/sitemap.xml",
            "https://www.conacmexico.org.mx/publicaciones-y-recursos",
            "https://www.conacmexico.org.mx/normatividad",
            "https://www.conacmexico.org.mx/revista-vol-19-",
            "https://www.conacmexico.org.mx/revista-vol-18",
        ),
        ("conacmexico.org.mx",),
        0.5,
    ),
}


class RuntimeErrorSafe(RuntimeError):
    """A bounded, user-actionable local flow failure."""


class _TargetWorkerFailure(RuntimeErrorSafe):
    def __init__(self, target: str, reply: dict[str, Any], message: str) -> None:
        diagnostic = {
            "configured_target": target,
            "worker_target": reply.get("target"),
            "error": reply.get("error"),
        }
        details = reply.get("error_details")
        if details is not None:
            diagnostic["error_details"] = details
        super().__init__(f"{target} target worker failed: {message}")
        self.diagnostic_evidence = [
            json.dumps(diagnostic, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
        ]


class RequestBudgetExceeded(RuntimeErrorSafe):
    pass


class _BudgetCrawler(Crawler):
    def __init__(self, *args: Any, **kwargs: Any) -> None:
        super().__init__(*args, **kwargs)
        self.site_attempts = 0
        self.site_budget = MAX_SITE_REQUESTS
        self.caa_cookie: str | None = None
        self.algolia_key: str | None = None

    def begin_site(self, budget: int = MAX_SITE_REQUESTS) -> None:
        self.site_attempts = 0
        self.site_budget = min(MAX_SWISS_SITE_REQUESTS, max(1, int(budget)))

    def _open_pinned_http(
        self, url: str, resolution: Any, *, timeout: int, delay_seconds: float | None
    ):
        if self.site_attempts >= self.site_budget:
            raise RequestBudgetExceeded("site HTTP request budget reached before the next request")
        self.site_attempts += 1
        return super()._open_pinned_http(
            url, resolution, timeout=timeout, delay_seconds=delay_seconds
        )

    def _request_post(
        self, url: str, body: bytes, headers: dict[str, str], *, timeout: int, delay_seconds: float
    ) -> tuple[bytes, dict[str, str], str]:
        """POST through the crawler's DNS validation, pinning, and host pacing."""
        if self.site_attempts >= self.site_budget:
            raise RequestBudgetExceeded("site HTTP request budget reached before the next request")
        self.site_attempts += 1
        resolution = resolve_safe_http_url(url)
        parsed = urlsplit(url)
        scheme = parsed.scheme.lower()
        port = parsed.port or (443 if scheme == "https" else 80)
        target = urlunsplit(("", "", parsed.path or "/", parsed.query, ""))
        host_header = resolution.host
        if (scheme == "http" and port != 80) or (scheme == "https" and port != 443):
            host_header = f"{host_header}:{port}"
        last_error: Exception | None = None
        for address in self._preferred_addresses(resolution):
            conn = http.client.HTTPConnection(resolution.host, port=port, timeout=timeout)
            try:
                self._pace_request(url, delay_seconds)
                sock = socket.create_connection((str(address), port), timeout=timeout)
                if scheme == "https":
                    sock = ssl.create_default_context().wrap_socket(
                        sock, server_hostname=resolution.host
                    )
                conn.sock = sock
                request_headers = {"User-Agent": self.user_agent, "Host": host_header, **headers}
                conn.request("POST", target, body=body, headers=request_headers)
                response = conn.getresponse()
                status = int(response.status)
                response_headers = {
                    str(key).lower(): str(value) for key, value in response.getheaders()
                }
                if status in {301, 302, 303, 307, 308}:
                    raise RuntimeErrorSafe(
                        "POST endpoint redirected; refusing to forward form data"
                    )
                self._raise_for_status(url, status)
                return response.read(), response_headers, url
            except Exception as exc:
                conn.close()
                last_error = exc
        if last_error:
            raise last_error
        raise RuntimeErrorSafe("no validated address was available for the POST endpoint")


class _LinkParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.links: list[tuple[str, str]] = []
        self._href: str | None = None
        self._text: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag.lower() == "a":
            self._href = next(
                (value for key, value in attrs if key.lower() == "href" and value), None
            )
            self._text = []

    def handle_data(self, data: str) -> None:
        if self._href:
            self._text.append(data)

    def handle_endtag(self, tag: str) -> None:
        if tag.lower() == "a" and self._href:
            label = " ".join(" ".join(self._text).split())[:240]
            self.links.append((self._href, label))
            self._href, self._text = None, []


def _load_flow_module():
    spec = importlib.util.spec_from_file_location("local_actuarial_flow", FLOW_SCRIPT)
    if spec is None or spec.loader is None:
        raise RuntimeErrorSafe("the v2 local-flow receipt contract is unavailable")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def _atomic_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    content = json.dumps(value, ensure_ascii=False, sort_keys=True, indent=2) + "\n"
    fd, temp_name = tempfile.mkstemp(prefix=f".{path.name}.", suffix=".tmp", dir=path.parent)
    temp = Path(temp_name)
    try:
        with os.fdopen(fd, "w", encoding="utf-8", newline="\n") as handle:
            handle.write(content)
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temp, path)
    finally:
        temp.unlink(missing_ok=True)


@contextlib.contextmanager
def _exclusive_lock(path: Path) -> Iterator[None]:
    path.parent.mkdir(parents=True, exist_ok=True)
    handle = path.open("a+b")
    try:
        try:
            import msvcrt

            handle.seek(0)
            if handle.read(1) == b"":
                handle.seek(0)
                handle.write(b"0")
                handle.flush()
            handle.seek(0)
            msvcrt.locking(handle.fileno(), msvcrt.LK_NBLCK, 1)
        except ImportError:
            fcntl = importlib.import_module("fcntl")
            fcntl.flock(handle.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
        except OSError as exc:
            raise RuntimeErrorSafe("another local collection run holds the runtime lock") from exc
        yield
    finally:
        try:
            if os.name == "nt":
                import msvcrt

                handle.seek(0)
                msvcrt.locking(handle.fileno(), msvcrt.LK_UNLCK, 1)
            else:
                fcntl = importlib.import_module("fcntl")
                fcntl.flock(handle.fileno(), fcntl.LOCK_UN)
        except OSError:
            pass
        handle.close()


def _canonical_url(value: str) -> str:
    parsed = urlsplit(value.strip())
    if (
        parsed.scheme.lower() not in {"http", "https"}
        or not parsed.hostname
        or parsed.username
        or parsed.password
    ):
        raise RuntimeErrorSafe("discovered URL is not a safe HTTP(S) URL")
    host = parsed.hostname.lower().rstrip(".")
    if host in {
        "www.actuaries.org",
        "www.actuaries.org.uk",
        "www.actuarios.org",
        "www.actuary.org",
        "www.ordineattuari.it",
    }:
        host = host.removeprefix("www.")
    port = parsed.port
    netloc = (
        host
        if not port
        or (parsed.scheme == "https" and port == 443)
        or (parsed.scheme == "http" and port == 80)
        else f"{host}:{port}"
    )
    query = urlencode(sorted(parse_qsl(parsed.query, keep_blank_values=True)))
    return urlunsplit(
        (parsed.scheme.lower(), netloc, re.sub(r"/{2,}", "/", parsed.path or "/"), query, "")
    )


def _allowed_host(url: str, site: Site) -> bool:
    hostname = (urlsplit(url).hostname or "").lower().rstrip(".")
    if site.site_id == "australia" and hostname in {"content.actuaries.asn.au", ALGOLIA_HOST}:
        return True
    return any(hostname == allowed or hostname.endswith("." + allowed) for allowed in site.hosts)


def _looks_document(url: str, label: str = "") -> bool:
    path = urlsplit(url).path.lower()
    return Path(path).suffix in _FILE_SUFFIXES or any(
        token in path
        for token in (
            "/document/download",
            "/download/",
            "/app/uploads/",
            "/globalassets/",
            "/docs/",
            "/re/pdf/",
            "/re/files/",
        )
    )


def _looks_site_document(url: str, label: str, site: Site) -> bool:
    path = urlsplit(url).path.lower()
    if any(token in path for token in ("/private-files/", "/member-only/")):
        return False
    if _looks_document(url, label):
        return True
    parsed = urlsplit(url)
    if site.site_id in {"cia-ica", "cia-research-publications"}:
        query = dict(parse_qsl(parsed.query, keep_blank_values=True))
        return (
            parsed.path.lower() == "/app/themes/wicket/custom/dl_file.php"
            and bool(query.get("p"))
            and bool(query.get("fid"))
        )
    if site.site_id == "france" and parsed.path.lower() == "/global/gene/link.php":
        query = dict(parse_qsl(parsed.query, keep_blank_values=True))
        try:
            document_id = int(query.get("doc_id") or 0)
        except (TypeError, ValueError):
            return False
        return document_id > 1 and query.get("fg") == "1"
    return False


def _canonical_http_link(base_url: str, href: str) -> str | None:
    try:
        candidate = urljoin(base_url, html.unescape(href).strip())
        if urlsplit(candidate).scheme.lower() not in {"http", "https"}:
            return None
        return _canonical_url(candidate)
    except (RuntimeErrorSafe, ValueError):
        return None


def _relevant_page(url: str, label: str, site: Site) -> bool:
    probe = f"{url} {label}".lower()
    parsed = urlsplit(url)
    if (
        site.site_id == "cca"
        and (parsed.hostname or "").lower() == "my.ccactuaries.org"
        and parsed.path.lower() == "/security/sign-in"
    ):
        return False
    if any(token in probe for token in _EXCLUDE_HINTS):
        return False
    if site.site_id == "france" and re.search(
        r"mémoire|memoire|présentation|presentation|conférences|conferences|congrès|congres|actuariat",
        probe,
        re.I,
    ):
        return not bool(re.search(r"inscription|agenda général|événementiel", probe, re.I))
    return any(token in probe for token in site.page_hints)


def _initial_frontier(site: Site) -> list[dict[str, Any]]:
    if site.site_id == "caa":
        news = [
            {
                "url": "https://www.e-caa.org.cn/systemNews/getNewsInfo/index?newsType="
                + str(news_type),
                "kind": "json_post",
                "news_type": news_type,
                "page": 1,
                "page_size": 50,
                "referrer": "https://www.e-caa.org.cn/",
            }
            for news_type in (12, 13, 20, 21, 22, 25)
        ]
        datatype_parents = (
            "40284c814d8f8966014d8f8966c00000",
            "40284c814d8f8966014d8f89a5b70001",
            "40284c814d8f835b014d8f87b6030003",
        )
        return news + [
            {
                "url": "https://www.e-caa.org.cn/uploadDataType/getDataTypeByParent",
                "kind": "caa_children",
                "parent_id": parent_id,
                "referrer": "https://www.e-caa.org.cn/",
            }
            for parent_id in datatype_parents
        ]
    if site.site_id in {"cia-ica", "cia-research-publications"}:
        rest_url = "https://www.cia-ica.ca/wp-json/wp/v2/publication?per_page=100&page=1&_fields=id,slug,link,title,content,acf,publication-category,publication-type,publication-topic"
        page_url = "https://www.cia-ica.ca/publications/"
        return [
            {"url": page_url, "kind": "page", "depth": 0, "referrer": page_url},
            {"url": rest_url, "kind": "json_rest", "page": 1, "referrer": page_url},
        ]
    if site.site_id == "iba-brazil":
        return [
            {
                "url": "https://atuarios.org.br/wp-json/wp/v2/posts?per_page=100&page=1&_fields=id,slug,link,title,content,acf,categories,tags",
                "kind": "json_rest",
                "page": 1,
                "referrer": site.seeds[0],
            }
        ]
    if site.site_id == "australia":
        rows = [{"url": seed, "kind": "page", "depth": 0, "referrer": seed} for seed in site.seeds]
        rows.append(
            {
                "url": "https://www.actuaries.asn.au/",
                "kind": "algolia_bootstrap",
                "depth": 0,
                "referrer": "https://www.actuaries.asn.au/",
            }
        )
        return rows
    return [
        {
            "url": seed,
            "kind": (
                "sitemap"
                if "sitemap" in seed.lower()
                else "document" if _looks_document(seed) else "page"
            ),
            "depth": 0,
            "referrer": seed,
        }
        for seed in site.seeds
    ]


_SITE_EXCLUDE = {
    "iaa": re.compile(
        r"membership|election|meeting|award|event|newsletter|recruit|vacan|press-release|/news/|"
        r"org[_\s-]*chart|organizational[_\s-]*chart|organization[_\s-]*chart",
        re.I,
    ),
    "soa": re.compile(
        r"education|exam|curriculum|candidate|event|newsletter|governance|/link/[0-9a-f]{8,}", re.I
    ),
    "soa-ai-topic-landing": re.compile(
        r"education|exam|curriculum|candidate|event|/link/[0-9a-f]{8,}", re.I
    ),
    "soa-ai-bulletin-focused": re.compile(r"newsletter|exam|candidate|/link/[0-9a-f]{8,}", re.I),
    "cas": re.compile(
        r"admission|exam|membership|continuing-education|professional-education|/admin/|/user/",
        re.I,
    ),
    "aaa": re.compile(r"publication-issue|event|biography|glossary|newsletter|/wp-admin/", re.I),
    "cia-ica": re.compile(
        r"/private-files/|members[_-]?only|newsletter|exam|syllab|governance|admin", re.I
    ),
    "cia-research-publications": re.compile(
        r"/private-files/|members[_-]?only|newsletter|exam|syllab|governance|admin", re.I
    ),
    "cca": re.compile(r"license-fee|pension|asop-checklist|member-admin|/account/", re.I),
    "aae": re.compile(r"event|board|agenda|minutes|training|newsletter", re.I),
    "ifoa-uk": re.compile(
        r"exam|qualification|event|press-release|/news/|membership|governance|vacan", re.I
    ),
    "france": re.compile(r"event|agenda|news|training|membership|admin|member-only", re.I),
    "dav": re.compile(r"event|newsroom|exam|training|regularity|/admin/", re.I),
    "iba-brazil": re.compile(
        r"^/(?:eventos?|team|author|category|tag|blog)(?:/|$)|exame|admiss[aã]o|curso|congresso|elei[cç][aã]o|recrut|vaga|inscri",
        re.I,
    ),
    "aict": re.compile(r"/re/|/memberconf/|/member/|/board/|login", re.I),
    "swiss-association-of-actuaries": re.compile(
        r"mitglieder|mitgliedschaft|statuten|impressum|newsletter|account|register", re.I
    ),
    "australia": re.compile(
        r"consultation|professional-standard|media-release|newsletter|podcast|webinar|presentation-slide|video",
        re.I,
    ),
    "italy": re.compile(
        r"assembl|elezion|training|concorso|job|recruit|member-register|event|trasparenza", re.I
    ),
    "spain": re.compile(
        r"assemble|election|statute|conduct|CPD|training|job|news|event|member", re.I
    ),
    "iai": re.compile(
        r"advert|vacan|recruit|tender|\brfp\b|announcement|notice|member-admin", re.I
    ),
    "iaj": re.compile(r"collection|exam|directory|person|CERA", re.I),
    "institute-of-actuaries-of": re.compile(
        r"general-meeting|board|executive|member|about|application", re.I
    ),
    "caa": re.compile(r"exam|recruit|job|member-notice|meeting|education|admin", re.I),
    "ashk": re.compile(
        r"AGM|general-meeting|council|governance|minutes|membership|training|event|career|admin",
        re.I,
    ),
    "assa": re.compile(r"newsletter|brochure|annual-report|member|admin", re.I),
    "ama": re.compile(
        r"event|agenda|webinar|training|statute|governance|member|private|/fr/", re.I
    ),
    "task": re.compile(r"newsletter|brochure|apprenticeship|admin", re.I),
    "aca-colombia": re.compile(
        r"statute|conduct|membership|form|exam|syllabus|course|governance", re.I
    ),
    "conac-mexico": re.compile(
        r"assembly|governance|statute|ethic|exam|certification|education|event|award|fee", re.I
    ),
}


def _allowed_path(url: str, site: Site) -> bool:
    path = urlsplit(url).path
    if site.site_id == "aict" and path.lower().startswith(("/re/", "/tmp/")):
        # AICT explicitly disallows its entire public-document directory in robots.txt.
        return False
    rule = _SITE_EXCLUDE.get(site.site_id)
    return not (rule and rule.search(url))


def _site_candidate_allowed(
    site: Site, url: str, title: str = "", metadata: dict[str, Any] | None = None
) -> bool:
    if not _allowed_path(url, site):
        return False
    if (
        site.site_id == "iaa"
        and Path(urlsplit(url).path).name.casefold() in _IAA_ADMIN_DOCUMENT_NAMES
    ):
        return False
    probe = f"{url} {title}"
    rule = _SITE_EXCLUDE.get(site.site_id)
    if rule and rule.search(probe):
        return False
    if site.site_id == "iaa" and re.search(r"governance", probe, re.I):
        path = urlsplit(url).path.lower()
        ai_research_title = re.search(
            r"\b(?:ai|artificial intelligence)\b", title, re.I
        ) and re.search(r"framework|research|paper|study|report", title, re.I)
        if not path.startswith("/paper/") or not ai_research_title:
            return False
    if site.site_id == "swiss-association-of-actuaries":
        if metadata and metadata.get("category_id") in {5, 6}:
            return False
        category_id = (metadata or {}).get("category_id")
        if category_id == 11:
            return bool(
                re.search(
                    r"AI|KI|IA|machine learning|data science|neural|actuar|research|pricing",
                    title,
                    re.I,
                )
            )
        if category_id == 8:
            return bool(
                re.search(
                    r"AI|KI|IA|machine learning|data science|neural|actuar|research|pricing",
                    title,
                    re.I,
                )
            )
    if site.site_id == "australia" and metadata:
        content_types = metadata.get("contentTypes") or []
        return any(kind in content_types for kind in AUSTRALIA_RESEARCH_TYPES)
    return True


def _site_document_relevant(
    site: Site, url: str, title: str, referrer: str, metadata: dict[str, Any] | None = None
) -> bool:
    metadata = metadata or {}
    if metadata.get("operator_exact_url"):
        return True
    if site.site_id == "australia":
        return _site_candidate_allowed(site, url, title, metadata)
    evidence = " ".join(
        [url, title, referrer, *(str(value) for value in metadata.values())]
    ).lower()
    if _site_candidate_allowed(site, url, title, metadata) is False:
        return False
    return (
        _relevant_page(url, title, site)
        or _relevant_page(referrer, title, site)
        or bool(
            re.search(
                r"research|publication|report|paper|study|technical|standard|guideline|journal|document|bulletin|analysis|actuar|data.science|machine.learning|人工智能|研究|精算|论文|標準|标准|報告|报告|技術|技术",
                evidence,
                re.I,
            )
        )
    )


def _queue_identity(item: dict[str, Any]) -> str:
    if item.get("kind") == "algolia":
        return f"{item.get('url')}|{item.get('content_type')}|{item.get('page')}"
    if item.get("kind") == "json_post":
        return f"{item.get('url')}|{item.get('news_type')}|{item.get('page')}"
    if item.get("kind") == "caa_children":
        return f"{item.get('url')}|children|{item.get('parent_id')}"
    if item.get("kind") == "caa_shared":
        return f"{item.get('url')}|shared|{item.get('parent_id')}|{item.get('page')}"
    return str(item.get("url") or "")


_DOCUMENT_PRIORITY_SITES = {
    "cia-ica",
    "cia-research-publications",
    "cca",
    "france",
}


def _prioritize_pending_documents(site: Site, queue: list[dict[str, Any]], cursor: int) -> None:
    """Stable-partition the unconsumed queue suffix for audited sources only."""
    if site.site_id not in _DOCUMENT_PRIORITY_SITES or cursor < 0 or cursor >= len(queue):
        return

    def priority(item: dict[str, Any]) -> int:
        if not isinstance(item, dict):
            return 1
        if site.site_id == "france" and item.get("frontier_priority") == "france_memoirs_v2":
            return -1
        if item.get("kind") != "document":
            return 1
        url = str(item.get("url") or "")
        title = str(item.get("title") or "")
        referrer = str(item.get("referrer") or "")
        metadata = item.get("discovery_metadata")
        if not isinstance(metadata, dict):
            metadata = {}
        try:
            path = urlsplit(url).path.lower()
            if any(token in path for token in ("/private-files/", "/member-only/")):
                return 1
            public = (
                _allowed_host(url, site)
                and _allowed_path(url, site)
                and _looks_site_document(url, title, site)
                and _site_candidate_allowed(site, url, title, metadata)
                and _site_document_relevant(site, url, title, referrer, metadata)
            )
        except (RuntimeErrorSafe, ValueError):
            public = False
        return 0 if public else 1

    queue[cursor:] = sorted(queue[cursor:], key=priority)


def _frontier_repair_items(site: Site) -> tuple[set[str], list[dict[str, Any]]]:
    if site.site_id in {"cia-ica", "cia-research-publications"}:
        page = "https://www.cia-ica.ca/publications/"
        return set(), [{"url": page, "kind": "page", "depth": 0, "referrer": page}]
    if site.site_id == "cca":
        old = {
            "https://www.ccactuaries.org/papers",
            "https://www.ccactuaries.org/technical-resource-library-main",
        }
        replacements = (
            "https://www.ccactuaries.org/professional-resources/papers",
            "https://www.ccactuaries.org/professional-resources/technical-resource-library-main",
        )
        return old, [
            {"url": url, "kind": "page", "depth": 0, "referrer": url} for url in replacements
        ]
    if site.site_id == "france":
        obsolete = {
            "https://www.institutdesactuaires.com/global/gene/link.php?doc_id=1&fg=1",
            "https://www.institutdesactuaires.com/se-documenter/memoire-d-actuariat-38",
            "https://www.institutdesactuaires.com/se-documenter/memoires-38",
        }
        memoir_listing, presentations_listing = site.seeds
        return obsolete, [
            {
                "url": memoir_listing,
                "kind": "page",
                "depth": 0,
                "referrer": memoir_listing,
                "frontier_priority": "france_memoirs_v2",
            },
            {
                "url": presentations_listing,
                "kind": "page",
                "depth": 0,
                "referrer": presentations_listing,
            },
        ]
    return set(), []


def _repair_saved_frontier(site: Site, site_state: dict[str, Any]) -> None:
    """Patch only the observed obsolete seeds in persisted source frontiers."""
    marker_version = "v2" if site.site_id == "france" else "v1"
    marker = f"frontier_repair_{site.site_id}_{marker_version}"
    if site_state.get(marker):
        return
    queue = list(site_state.get("queue") or [])
    retry_queue = list(site_state.get("retry_queue") or [])
    cursor = min(max(0, int(site_state.get("cursor") or 0)), len(queue))
    visited = set(site_state.get("visited") or [])
    obsolete, replacements = _frontier_repair_items(site)
    obsolete_keys = set()
    for value in obsolete:
        try:
            obsolete_keys.add(_canonical_url(value))
        except RuntimeErrorSafe:
            continue

    kept: list[dict[str, Any]] = []
    removed_before_cursor = 0
    for index, row in enumerate(queue):
        if not isinstance(row, dict):
            if index < cursor:
                removed_before_cursor += 1
            continue
        raw_url = str(row.get("url") or "")
        try:
            scheme = urlsplit(raw_url).scheme.lower()
        except ValueError:
            scheme = ""
        try:
            canonical = _canonical_url(raw_url) if scheme in {"http", "https"} else None
        except (RuntimeErrorSafe, ValueError):
            canonical = None
        if canonical is None or canonical in obsolete_keys:
            if index < cursor:
                removed_before_cursor += 1
            continue
        kept.append(row)
    queue = kept
    cursor = max(0, cursor - removed_before_cursor)

    kept_retries = []
    for row in retry_queue:
        if not isinstance(row, dict):
            continue
        raw_url = str(row.get("url") or "")
        try:
            scheme = urlsplit(raw_url).scheme.lower()
        except ValueError:
            scheme = ""
        if scheme not in {"http", "https"}:
            continue
        try:
            if _canonical_url(raw_url) in obsolete_keys:
                continue
        except (RuntimeErrorSafe, ValueError):
            continue
        kept_retries.append(row)

    if site.site_id == "france":
        memoir_listing = site.seeds[0]
        memoir_identity = _queue_identity({"url": memoir_listing, "kind": "page"})
        if memoir_identity not in visited:
            for row in [*queue[cursor:], *kept_retries]:
                if isinstance(row, dict) and _queue_identity(row) == memoir_identity:
                    row["kind"] = "page"
                    row["depth"] = 0
                    row["frontier_priority"] = "france_memoirs_v2"

    existing = {
        _queue_identity(row) for row in [*queue, *kept_retries] if isinstance(row, dict)
    } | visited
    insertions = [row for row in replacements if _queue_identity(row) not in existing]
    if insertions:
        queue[cursor:cursor] = insertions
    site_state["queue"] = queue
    site_state["cursor"] = cursor
    site_state["retry_queue"] = kept_retries
    site_state[marker] = True


def _enqueue_url(
    site: Site,
    queue: list[dict[str, Any]],
    queued: set[str],
    url: str,
    *,
    title: str,
    referrer: str,
    metadata: dict[str, Any] | None = None,
    kind: str | None = None,
) -> None:
    candidate = _canonical_http_link(referrer, url)
    if candidate is None:
        return
    metadata = metadata or {}
    document = kind == "document" or (kind is None and _looks_site_document(candidate, title, site))
    if not _allowed_host(candidate, site) or not _site_candidate_allowed(
        site, candidate, title, metadata
    ):
        return
    if document and not _site_document_relevant(site, candidate, title, referrer, metadata):
        return
    row = {
        "url": candidate,
        "kind": "document" if document else "page",
        "depth": 1,
        "referrer": referrer,
        "title": title,
        "discovery_metadata": metadata,
    }
    identity = _queue_identity(row)
    if identity not in queued:
        queue.append(row)
        queued.add(identity)


def _parse_swiss_embedded_documents(document: str) -> list[dict[str, Any]]:
    match = re.search(r"\"items\"\s*:\s*\[", document)
    if not match:
        return []
    start = document.find("[", match.start())
    try:
        items, _end = json.JSONDecoder().raw_decode(document[start:])
    except json.JSONDecodeError:
        return []
    if not isinstance(items, list):
        return []
    output = []
    for item in items:
        if not isinstance(item, dict):
            continue
        title = str(item.get("title") or "").strip()
        link = str(item.get("downloadLinkDokument") or "").strip()
        category = item.get("documentCategoryId")
        try:
            category_id = int(category)
        except (TypeError, ValueError):
            category_id = None
        metadata = {
            "category_id": category_id,
            "document_detail_id": item.get("documentDetailId"),
            "document_type": item.get("documentType"),
            "language_id": item.get("languageId"),
            "published_hint": item.get("datum"),
        }
        if link and title:
            output.append({"url": link, "title": title, "metadata": metadata})
    return output


def _swiss_document_is_relevant(item: dict[str, Any]) -> bool:
    category = (item.get("metadata") or {}).get("category_id")
    if category in {2, 3, 4, 7, 9}:
        return True
    return category in {8, 11} and bool(
        re.search(
            r"AI|KI|IA|machine learning|data science|neural|actuar|research|pricing|reserving",
            item.get("title", ""),
            re.I,
        )
    )


def _parse_algolia_script(script: str) -> dict[str, str] | None:
    key_match = re.search(
        r"(?:apiKey|searchOnlyApiKey|searchApiKey|api_key)\s*[:=]\s*['\"]([^'\"]{20,})['\"]",
        script,
        re.I,
    )
    app_match = re.search(
        r"(?:appId|applicationId|applicationID)\s*[:=]\s*['\"]([A-Z0-9]{6,})['\"]", script, re.I
    )
    if not key_match or not app_match:
        return None
    index_match = re.search(r"indexName\s*[:=]\s*['\"]([^'\"]+)['\"]", script, re.I)
    index = index_match.group(1) if index_match else "prod_actuaries_search"
    return {"api_key": key_match.group(1), "app_id": app_match.group(1), "index": index}


def _wp_html_links(
    site: Site,
    html_body: str,
    referrer: str,
    queue: list[dict[str, Any]],
    queued: set[str],
    title: str,
) -> None:
    parser = _LinkParser()
    parser.feed(html_body)
    for href, label in parser.links:
        candidate = _canonical_http_link(referrer, href)
        if candidate is None:
            continue
        if _looks_site_document(candidate, label, site):
            _enqueue_url(
                site,
                queue,
                queued,
                candidate,
                title=label or title,
                referrer=referrer,
                kind="document",
            )


def _recursive_values(value: Any, key: str = "") -> Iterator[tuple[str, Any]]:
    if isinstance(value, dict):
        for child_key, child in value.items():
            yield from _recursive_values(child, str(child_key))
    elif isinstance(value, list):
        for child in value:
            yield from _recursive_values(child, key)
    else:
        yield key.lower(), value


def _wp_item_relevant(site: Site, row: dict[str, Any], title: str, url: str) -> bool:
    if not _site_candidate_allowed(site, url, title):
        return False
    probe = f"{url} {title} " + " ".join(
        str(v) for _, v in _recursive_values(row) if isinstance(v, str)
    )
    if site.site_id in {"cia-ica", "cia-research-publications"}:
        return bool(
            re.search(
                r"research|publication|study|report|technical|practice|standard|actuar|research|经验|研究|publication",
                probe,
                re.I,
            )
        )
    return bool(
        re.search(
            r"revista|white.?paper|caderno|publica|nota.?t[eé]cnica|estudo.?atuarial|posicionamento|manifesta|parecer|cpa[-_ ]?\d|consulta.?t[eé]cnica|metodologia|actuari|actuarial",
            probe,
            re.I,
        )
    )


def _enqueue_wordpress_results(
    site: Site,
    item: dict[str, Any],
    body: bytes,
    headers: dict[str, str],
    final_url: str,
    queue: list[dict[str, Any]],
    queued: set[str],
) -> None:
    try:
        rows = json.loads(body.decode("utf-8"))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise RuntimeErrorSafe("WordPress REST returned invalid JSON") from exc
    if not isinstance(rows, list):
        return
    for row in rows:
        if not isinstance(row, dict):
            continue
        title_value = row.get("title")
        title = str(
            (title_value.get("rendered") if isinstance(title_value, dict) else title_value) or ""
        ).strip()
        link = str(row.get("link") or "").strip()
        content = row.get("content")
        rendered = (
            str(content.get("rendered") or "") if isinstance(content, dict) else str(content or "")
        )
        if link and _wp_item_relevant(site, row, title, link):
            _enqueue_url(
                site,
                queue,
                queued,
                link,
                title=title,
                referrer=final_url,
                metadata={"wp_id": row.get("id")},
            )
            _wp_html_links(site, rendered, link, queue, queued, title)
            for key, value in _recursive_values(row):
                if isinstance(value, str) and (
                    "/app/uploads/" in value
                    or "/wp-content/uploads/" in value
                    or _looks_document(value)
                ):
                    candidate = urljoin(link, html.unescape(value))
                    _enqueue_url(
                        site, queue, queued, candidate, title=title, referrer=link, kind="document"
                    )
                elif (
                    key in {"file", "attachment", "file_id"}
                    and isinstance(value, int)
                    and value > 0
                ):
                    media_url = f"{urlsplit(link).scheme}://{urlsplit(link).netloc}/wp-json/wp/v2/media/{value}"
                    media_item = {
                        "url": _canonical_url(media_url),
                        "kind": "json_media",
                        "referrer": link,
                        "title": title,
                    }
                    identity = _queue_identity(media_item)
                    if identity not in queued:
                        queue.append(media_item)
                        queued.add(identity)
    try:
        page = int(item.get("page") or 1)
        total_pages = min(int(headers.get("x-wp-totalpages") or 1), 20)
    except (TypeError, ValueError):
        page, total_pages = 1, 1
    if page < total_pages:
        parsed = urlsplit(str(item["url"]))
        query = dict(parse_qsl(parsed.query, keep_blank_values=True))
        query["page"] = str(page + 1)
        next_item = {
            **item,
            "url": urlunsplit((parsed.scheme, parsed.netloc, parsed.path, urlencode(query), "")),
            "page": page + 1,
        }
        identity = _queue_identity(next_item)
        if identity not in queued:
            queue.append(next_item)
            queued.add(identity)


def _enqueue_caa_results(
    site: Site,
    item: dict[str, Any],
    body: bytes,
    final_url: str,
    queue: list[dict[str, Any]],
    queued: set[str],
) -> None:
    try:
        payload = json.loads(body.decode("utf-8"))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise RuntimeErrorSafe("CAA endpoint returned invalid JSON") from exc
    if item.get("kind") == "caa_children":
        rows = payload if isinstance(payload, list) else []
        for row in rows:
            if not isinstance(row, dict):
                continue
            child_id = row.get("dataTypeId")
            child_name = str(row.get("dataTypeName") or "").strip()
            if not child_id or not re.search(r"论文|规范|标准|研究|学术", child_name):
                continue
            next_item = {
                "url": "https://www.e-caa.org.cn/informationSharing/more",
                "kind": "caa_shared",
                "parent_id": str(child_id),
                "parent_name": child_name,
                "page": 1,
                "page_size": 50,
                "referrer": final_url,
            }
            identity = _queue_identity(next_item)
            if identity not in queued:
                queue.append(next_item)
                queued.add(identity)
        return
    if item.get("kind") == "caa_shared":
        rows = (
            payload
            if isinstance(payload, list)
            else payload.get("rows") or [] if isinstance(payload, dict) else []
        )
        for row in rows:
            if not isinstance(row, dict):
                continue
            title = str(row.get("fileName") or row.get("title") or "").strip()
            document_url = str(row.get("filePath") or row.get("url") or "").strip()
            if (
                not title
                or not document_url
                or not re.search(r"研究|论文|标准|規範|精算|actuar|research|standard", title, re.I)
            ):
                continue
            _enqueue_url(
                site,
                queue,
                queued,
                urljoin("https://www.e-caa.org.cn/", document_url),
                title=title,
                referrer=final_url,
                metadata={
                    "parent_name": item.get("parent_name"),
                    "parent_id": item.get("parent_id"),
                },
                kind="document",
            )
        page = int(item.get("page") or 1)
        if len(rows) >= int(item.get("page_size") or 50) and page < 5:
            next_item = {**item, "page": page + 1}
            identity = _queue_identity(next_item)
            if identity not in queued:
                queue.append(next_item)
                queued.add(identity)
        return
    rows = payload.get("rows") or [] if isinstance(payload, dict) else []
    for row in rows:
        if not isinstance(row, dict):
            continue
        title = str(
            row.get("informationTitle") or row.get("title") or row.get("informationIntro") or ""
        ).strip()
        url = str(row.get("informationUrl") or row.get("url") or "").strip()
        if not url or not title:
            continue
        metadata = {
            "news_type": item.get("news_type"),
            "date_hint": row.get("auditTime") or row.get("year"),
        }
        if not _site_candidate_allowed(site, urljoin(final_url, url), title, metadata):
            continue
        if not re.search(
            r"研究|经验|论文|标准|规范|实践|精算|数据|report|research|actuar|standard", title, re.I
        ):
            continue
        _enqueue_url(
            site,
            queue,
            queued,
            urljoin(final_url, url),
            title=title,
            referrer=final_url,
            metadata=metadata,
        )
    page = int(item.get("page") or 1)
    size = int(item.get("page_size") or 50)
    total = int(payload.get("records") or 0) if isinstance(payload, dict) else 0
    if page * size < total and page < 20:
        next_item = {**item, "page": page + 1}
        parsed = urlsplit(str(item["url"]))
        query = dict(parse_qsl(parsed.query, keep_blank_values=True))
        query["page"] = str(page + 1)
        next_item["url"] = urlunsplit(
            (parsed.scheme, parsed.netloc, parsed.path, urlencode(query), "")
        )
        identity = _queue_identity(next_item)
        if identity not in queued:
            queue.append(next_item)
            queued.add(identity)


def _enqueue_australia_results(
    site: Site,
    item: dict[str, Any],
    body: bytes,
    final_url: str,
    queue: list[dict[str, Any]],
    queued: set[str],
) -> None:
    try:
        payload = json.loads(body.decode("utf-8"))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise RuntimeErrorSafe("Australia Algolia endpoint returned invalid JSON") from exc
    if not isinstance(payload, dict):
        return
    for hit in payload.get("hits") or []:
        if not isinstance(hit, dict):
            continue
        access = hit.get("accessLevel") or []
        if access and "public" not in [str(value).lower() for value in access]:
            continue
        slug = str(hit.get("slug") or "").strip()
        if not slug:
            continue
        url = slug if slug.startswith("http") else urljoin("https://www.actuaries.asn.au/", slug)
        if "content.actuaries.asn.au/resources/" in url and Path(urlsplit(url).path).suffix == "":
            kind = "document"
        else:
            kind = None
        content_types = hit.get("contentTypes") or []
        content_types = [
            str(value[-1] if isinstance(value, list) and value else value)
            for value in content_types
        ]
        metadata = {
            "contentTypes": content_types,
            "practiceAreas": hit.get("practiceAreas") or [],
            "year": hit.get("year"),
            "ai_hint": bool(
                re.search(
                    r"artificial intelligence|machine learning|data science|AI",
                    str(hit.get("name") or ""),
                    re.I,
                )
            ),
        }
        if not _site_candidate_allowed(site, url, str(hit.get("name") or ""), metadata):
            continue
        _enqueue_url(
            site,
            queue,
            queued,
            url,
            title=str(hit.get("name") or Path(urlsplit(url).path).name),
            referrer=final_url,
            metadata=metadata,
            kind=kind,
        )
    page = int(item.get("page") or 0)
    total_pages = min(int(payload.get("nbPages") or 1), 10)
    if page + 1 < total_pages:
        next_item = {**item, "page": page + 1}
        identity = _queue_identity(next_item)
        if identity not in queued:
            queue.append(next_item)
            queued.add(identity)


def _extract_search_script_urls(html_body: str, page_url: str) -> list[str]:
    urls = []
    for match in re.finditer(r"<script\b[^>]*\bsrc=[\"']([^\"']+)[\"'][^>]*>", html_body, re.I):
        script = _canonical_url(urljoin(page_url, html.unescape(match.group(1))))
        if "search" in script.lower() or ("astro" in script.lower() and "main" in script.lower()):
            urls.append(script)
    return list(dict.fromkeys(urls))[:3]


def _detect_kind(raw: bytes, content_type: str, url: str) -> str:
    if raw.startswith(b"%PDF-"):
        return "pdf"
    if raw.startswith(b"PK\x03\x04"):
        ext = Path(urlsplit(url).path).suffix.lower()
        if ext in {".docx", ".pptx", ".xlsx"}:
            return ext[1:]
        ooxml_kind_by_mime = {
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
            "application/vnd.openxmlformats-officedocument.presentationml.presentation": "pptx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "xlsx",
        }
        mime_type = content_type.partition(";")[0].strip().lower()
        return ooxml_kind_by_mime.get(mime_type, "zip")
    if len(raw) > 8 and raw[0:8] == bytes.fromhex("d0cf11e0a1b11ae1"):
        return "office_ole"
    lowered_type = content_type.lower()
    if "html" in lowered_type or raw.lstrip().lower().startswith(
        (b"<!doctype html", b"<html", b"<?xml")
    ):
        return "html"
    if Path(urlsplit(url).path).suffix.lower() in {".txt", ".rtf"}:
        return "text"
    return "unknown"


def _extract_markdown(raw: bytes, content_type: str, url: str, data_root: Path) -> str:
    kind = _detect_kind(raw, content_type, url)
    if kind == "pdf":
        from io import BytesIO

        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise RuntimeErrorSafe("pypdf is required to extract verified PDF text") from exc
        reader = PdfReader(BytesIO(raw), strict=False)
        if reader.is_encrypted and not reader.decrypt(""):
            raise RuntimeErrorSafe("encrypted PDF requires a password")
        if len(reader.pages) > 600:
            raise RuntimeErrorSafe("PDF exceeds the bounded 600-page extraction limit")
        sections = []
        for index, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            if text:
                sections.append(f"## Page {index}\n\n{text}")
            if sum(map(len, sections)) > MAX_MARKDOWN_CHARS:
                raise RuntimeErrorSafe("PDF text exceeds the bounded Markdown extraction size")
        markdown = "\n\n".join(sections)
    elif kind == "html":
        collector = WebPageCollector(SimpleNamespace(data_root=data_root), str(data_root))
        markdown = collector._extract_text(raw.decode("utf-8", errors="replace"), url) or ""
    elif kind == "text":
        markdown = raw.decode("utf-8", errors="replace")
    elif kind == "docx":
        try:
            from io import BytesIO

            from docx import Document

            document = Document(BytesIO(raw))
            markdown = "\n\n".join(p.text for p in document.paragraphs if p.text.strip())
        except Exception as exc:
            raise RuntimeErrorSafe("DOCX text extraction failed") from exc
    elif kind == "pptx":
        try:
            from io import BytesIO

            from pptx import Presentation

            deck = Presentation(BytesIO(raw))
            markdown = "\n\n".join(
                "\n".join(
                    shape.text for shape in slide.shapes if getattr(shape, "has_text_frame", False)
                )
                for slide in deck.slides
            )
        except Exception as exc:
            raise RuntimeErrorSafe("PPTX text extraction failed") from exc
    elif kind == "xlsx":
        try:
            from io import BytesIO

            from openpyxl import load_workbook

            book = load_workbook(BytesIO(raw), read_only=True, data_only=True)
            markdown = "\n\n".join(
                f"## {sheet.title}\n"
                + "\n".join(
                    " | ".join(str(cell or "") for cell in row)
                    for row in sheet.iter_rows(values_only=True)
                )
                for sheet in book.worksheets
            )
        except Exception as exc:
            raise RuntimeErrorSafe("XLSX text extraction failed") from exc
    else:
        raise RuntimeErrorSafe(
            f"unsupported document bytes ({kind}); original remains in local CAS"
        )
    markdown = markdown.strip()
    if len(markdown) > MAX_MARKDOWN_CHARS:
        raise RuntimeErrorSafe("extracted Markdown exceeds the configured size limit")
    if len(markdown) < 100:
        raise RuntimeErrorSafe("source produced no usable text; recorded as an extraction gap")
    return markdown


def _local_cas_path(data_root: Path, digest: str) -> Path:
    return data_root / "cas" / "sha256" / digest[:2] / digest


def _persist_cas(data_root: Path, raw: bytes, digest: str) -> Path:
    if len(raw) > MAX_ARTICLE_BYTES or hashlib.sha256(raw).hexdigest() != digest:
        raise RuntimeErrorSafe("downloaded object failed the source digest or size check")
    from ai_actuarial.capacity import ensure_capacity

    ensure_capacity(path=str(data_root), operation="local_source_cas_capture")
    path = _local_cas_path(data_root, digest)
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        if hashlib.sha256(path.read_bytes()).hexdigest() != digest:
            raise RuntimeErrorSafe("existing local CAS object does not match its digest")
        return path
    fd, name = tempfile.mkstemp(prefix=".source-", dir=path.parent)
    temporary = Path(name)
    try:
        with os.fdopen(fd, "wb") as handle:
            handle.write(raw)
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)
    return path


def _load_config(path: Path) -> dict[str, Any]:
    try:
        config = json.loads(path.read_text(encoding="utf-8"))
    except FileNotFoundError as exc:
        raise RuntimeErrorSafe(
            f"local config is missing: {path}; copy config.example.json and set machine paths"
        ) from exc
    except (OSError, json.JSONDecodeError) as exc:
        raise RuntimeErrorSafe("local config could not be read as JSON") from exc
    if not isinstance(config, dict) or int(config.get("schema_version") or 0) != 1:
        raise RuntimeErrorSafe("local config schema_version must be 1")
    for key in ("data_root", "state_root", "hub_connection_file", "targets"):
        if key not in config:
            raise RuntimeErrorSafe(f"local config is missing {key}")
    return config


def _connection_url(path: Path) -> DatabaseConfig:
    try:
        text = path.read_text(encoding="utf-8").strip()
    except OSError as exc:
        raise RuntimeErrorSafe(
            "the configured private Hub connection file cannot be opened"
        ) from exc
    try:
        value = json.loads(text)
    except json.JSONDecodeError:
        value = text
    if isinstance(value, dict):
        value = next(
            (
                value.get(key)
                for key in ("database_url", "dsn", "url", "connection_string")
                if value.get(key)
            ),
            None,
        )
    if not isinstance(value, str) or not value.startswith(("postgresql://", "postgres://")):
        raise RuntimeErrorSafe("the private Hub connection file has no supported PostgreSQL URL")
    return DatabaseConfig(value)


def _state_path(config: dict[str, Any]) -> Path:
    return Path(config["state_root"]).expanduser().resolve() / "runtime-state.json"


def _load_state(path: Path) -> dict[str, Any]:
    if not path.exists():
        return {
            "schema_version": 1,
            "sites": {},
            "articles": {},
            "target_acks": {},
            "backlog_cursor": 0,
        }
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise RuntimeErrorSafe("local runtime state is corrupt; no state was changed") from exc
    if value.get("schema_version") != 1 or not all(
        isinstance(value.get(key), dict) for key in ("sites", "articles", "target_acks")
    ):
        raise RuntimeErrorSafe("local runtime state does not match schema v1")
    return value


def _file_key(url: str, raw_sha: str) -> str:
    return hashlib.sha256(f"{url}\0{raw_sha}".encode("utf-8")).hexdigest()


def _article_identity(article: dict[str, Any]) -> str:
    return _file_key(str(article.get("url") or ""), str(article.get("raw_sha256") or ""))


def _load_category_keys() -> list[str]:
    from ai_actuarial.content_hub import APPROVED_CATEGORY_KEYS

    return sorted(APPROVED_CATEGORY_KEYS)


def _model_schema() -> dict[str, Any]:
    strings = {"type": "string", "minLength": 1}
    span_id = {"type": "string", "pattern": EVIDENCE_SPAN_ID_PATTERN}
    package = {
        "type": "object",
        "additionalProperties": False,
        "required": ["title", "summary", "category", "keywords"],
        "properties": {
            "title": strings,
            "summary": strings,
            "category": {"type": "string", "enum": _load_category_keys()},
            "keywords": {"type": "array", "minItems": 1, "maxItems": 16, "items": strings},
        },
    }
    date_fact = {
        "anyOf": [
            {"type": "null"},
            {
                "type": "object",
                "additionalProperties": False,
                "required": ["value", "precision", "semantic", "evidence", "location"],
                "properties": {
                    "value": {"type": "string"},
                    "precision": {"type": "string", "enum": ["year", "month", "day"]},
                    "semantic": {
                        "type": "string",
                        "enum": ["publication", "revision", "document_version"],
                    },
                    "evidence": span_id,
                    "location": strings,
                },
            },
        ]
    }
    return {
        "type": "object",
        "additionalProperties": False,
        "required": [
            "primary_category",
            "bilingual_metadata",
            "categories",
            "date_facts",
            "date_review_status",
            "date_review_reason",
            "date_review_unresolved",
        ],
        "properties": {
            "primary_category": {"type": "string", "enum": _load_category_keys()},
            "bilingual_metadata": {
                "type": "object",
                "additionalProperties": False,
                "required": ["zh", "en"],
                "properties": {"zh": package, "en": package},
            },
            "categories": {
                "type": "array",
                "minItems": 1,
                "maxItems": 4,
                "items": {
                    "type": "object",
                    "additionalProperties": False,
                    "required": ["key", "evidence"],
                    "properties": {
                        "key": {"type": "string", "enum": _load_category_keys()},
                        "evidence": {
                            "type": "array",
                            "minItems": 1,
                            "maxItems": 3,
                            "items": span_id,
                        },
                    },
                },
            },
            "date_facts": {
                "type": "object",
                "additionalProperties": False,
                "required": ["publication_date", "revision_date", "document_version_date"],
                "properties": {
                    key: date_fact
                    for key in ("publication_date", "revision_date", "document_version_date")
                },
            },
            "date_review_status": {
                "type": "string",
                "enum": ["complete", "candidate_limit_gap", "unsupported_format_gap"],
            },
            "date_review_reason": {"anyOf": [{"type": "null"}, strings]},
            "date_review_unresolved": {
                "type": "array",
                "items": {
                    "type": "object",
                    "additionalProperties": False,
                    "required": ["field", "evidence"],
                    "properties": {
                        "field": {
                            "type": "string",
                            "enum": ["publication_date", "revision_date", "document_version_date"],
                        },
                        "evidence": span_id,
                    },
                },
            },
        },
    }


def _build_evidence_span_catalog(markdown: str) -> list[dict[str, Any]]:
    """Create stable IDs for overlapping, verbatim source spans covering the Markdown."""
    if not isinstance(markdown, str) or not markdown:
        raise RuntimeErrorSafe("model evidence source Markdown is empty")
    spans: list[dict[str, Any]] = []
    start = 0
    while start < len(markdown):
        end = min(start + MAX_EVIDENCE_SPAN_CHARS, len(markdown))
        text = markdown[start:end]
        identity = hashlib.sha256(f"{start}\0{end}\0{text}".encode("utf-8")).hexdigest()[:16]
        spans.append(
            {
                "id": f"span-{identity}",
                "text": text,
                "start": start,
                "end": end,
            }
        )
        if end == len(markdown):
            break
        start = end - EVIDENCE_SPAN_OVERLAP_CHARS
    return spans


def _resolve_model_evidence_span_ids(
    model_result: dict[str, Any], span_catalog: list[dict[str, Any]]
) -> tuple[dict[str, Any], list[str]]:
    """Replace model-selected span IDs with their exact source Markdown substrings."""
    span_text = {
        row.get("id"): row.get("text")
        for row in span_catalog
        if isinstance(row, dict)
        and isinstance(row.get("id"), str)
        and isinstance(row.get("text"), str)
    }
    errors: list[str] = []

    def resolve(path: str, span_id: Any) -> Any:
        quote = span_text.get(span_id) if isinstance(span_id, str) else None
        if quote is None:
            errors.append(path)
            return span_id
        return quote

    result = dict(model_result)
    categories = model_result.get("categories")
    if isinstance(categories, list):
        resolved_categories = []
        for index, category in enumerate(categories):
            if not isinstance(category, dict):
                resolved_categories.append(category)
                continue
            copied = dict(category)
            evidence = category.get("evidence")
            if isinstance(evidence, list):
                if not evidence:
                    errors.append(f"categories[{index}].evidence")
                copied["evidence"] = [
                    resolve(f"categories[{index}].evidence[{quote_index}]", span_id)
                    for quote_index, span_id in enumerate(evidence)
                ]
            resolved_categories.append(copied)
        result["categories"] = resolved_categories

    date_facts = model_result.get("date_facts")
    if isinstance(date_facts, dict):
        resolved_dates = {}
        for field, fact in date_facts.items():
            if isinstance(fact, dict) and "evidence" in fact:
                copied = dict(fact)
                copied["evidence"] = resolve(f"date_facts.{field}.evidence", fact.get("evidence"))
                resolved_dates[field] = copied
            else:
                resolved_dates[field] = fact
        result["date_facts"] = resolved_dates

    unresolved = model_result.get("date_review_unresolved")
    if isinstance(unresolved, list):
        resolved_unresolved = []
        for index, candidate in enumerate(unresolved):
            if isinstance(candidate, dict) and "evidence" in candidate:
                copied = dict(candidate)
                copied["evidence"] = resolve(
                    f"date_review_unresolved[{index}].evidence",
                    candidate.get("evidence"),
                )
                resolved_unresolved.append(copied)
            else:
                resolved_unresolved.append(candidate)
        result["date_review_unresolved"] = resolved_unresolved
    return result, list(dict.fromkeys(errors))


def _model_evidence_errors(markdown: str, model_result: dict[str, Any]) -> list[str]:
    """Return evidence paths that are not short, exact source substrings."""
    errors: list[str] = []

    def check(path: str, quote: Any) -> None:
        if not isinstance(quote, str) or not quote.strip():
            errors.append(path)
        elif len(quote) > MAX_EVIDENCE_QUOTE_CHARS or quote not in markdown:
            errors.append(path)

    categories = model_result.get("categories")
    if isinstance(categories, list):
        for category_index, category in enumerate(categories):
            if not isinstance(category, dict) or not isinstance(category.get("evidence"), list):
                errors.append(f"categories[{category_index}].evidence")
                continue
            quotes = category["evidence"]
            if not quotes:
                errors.append(f"categories[{category_index}].evidence")
                continue
            for quote_index, quote in enumerate(quotes):
                check(f"categories[{category_index}].evidence[{quote_index}]", quote)
    else:
        errors.append("categories")

    date_fields = ("publication_date", "revision_date", "document_version_date")
    date_facts = model_result.get("date_facts")
    if isinstance(date_facts, dict):
        for field in date_fields:
            fact = date_facts.get(field)
            if fact is not None:
                path = f"date_facts.{field}.evidence"
                check(path, fact.get("evidence") if isinstance(fact, dict) else None)
    else:
        errors.append("date_facts")

    unresolved = model_result.get("date_review_unresolved")
    if isinstance(unresolved, list):
        for index, candidate in enumerate(unresolved):
            check(
                f"date_review_unresolved[{index}].evidence",
                candidate.get("evidence") if isinstance(candidate, dict) else None,
            )
    else:
        errors.append("date_review_unresolved")
    return errors


def _run_model(
    markdown: str,
    title: str,
    state_root: Path,
    version_key: str,
    codex_executable: str | None,
    *,
    evidence_feedback: list[str] | None = None,
    evidence_spans: list[dict[str, Any]] | None = None,
) -> dict[str, Any]:
    executable = codex_executable or shutil.which("codex.exe") or shutil.which("codex")
    if not executable:
        raise RuntimeErrorSafe("logged-in local Codex CLI was not found")
    attempt = state_root / "model-runs" / version_key
    attempt.mkdir(parents=True, exist_ok=True)
    schema_path = attempt / "output-schema.json"
    phase = "correction" if evidence_feedback else "initial"
    stamp = dt.datetime.now(dt.timezone.utc).strftime("%Y%m%dT%H%M%S%fZ")
    output_path = attempt / f"raw-{phase}-{stamp}.json"
    suffix = 1
    while output_path.exists():
        output_path = attempt / f"raw-{phase}-{stamp}-{suffix}.json"
        suffix += 1
    evidence_spans = (
        evidence_spans if evidence_spans is not None else _build_evidence_span_catalog(markdown)
    )
    catalog_path = attempt / "evidence-spans.json"
    catalog_path.write_text(
        json.dumps(evidence_spans, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    catalog_lines = "\n".join(
        json.dumps(
            {"id": row["id"], "text": row["text"]},
            ensure_ascii=False,
            separators=(",", ":"),
        )
        for row in evidence_spans
    )
    schema_path.write_text(json.dumps(_model_schema(), ensure_ascii=False), encoding="utf-8")
    prompt = (
        "Enrich this source-bound actuarial article using the complete source span catalog below. "
        "Return only JSON matching the supplied schema. Use only the approved canonical categories. "
        "Set the same primary category in zh/en. Categories are ordered, unique, and limited to four. "
        "For every category evidence array, choose one to three span IDs. For each non-null date fact and "
        "each unresolved-date entry, choose one span ID. The IDs map to exact, non-empty, contiguous "
        f"original Markdown substrings no longer than {MAX_EVIDENCE_QUOTE_CHARS} characters. The catalog "
        "uses overlapping spans and covers the full original Markdown. Return IDs only in evidence fields; "
        "never copy, normalize, paraphrase, or invent quote text or IDs. Use the span text and nearby spans "
        "as source context when deciding which evidence supports each field. Preserve all three date "
        "semantics separately, with ISO value and year/month/day precision; do not treat HTTP Last-Modified "
        "as a document date. Unknown dates are null; incomplete review requires a reason. Do not infer a date.\n\n"
        f"Source title: {title}\n\nExact source evidence spans as JSONL:\n{catalog_lines}"
    )
    if evidence_feedback:
        prompt += (
            "\n\nEvidence span correction required. The previous JSON had unknown or invalid IDs at these fields: "
            + ", ".join(evidence_feedback)
            + ". Choose only IDs that appear in the source span catalog above. Category evidence fields use "
            "an array of IDs; each date evidence field uses one ID. Do not return quote text, invent IDs, or "
            "remove a required evidence field. Return the complete corrected JSON."
        )
    command = [
        str(executable),
        "--ask-for-approval",
        "never",
        "--model",
        "gpt-6-luna",
        "--cd",
        str(state_root),
        "--sandbox",
        "read-only",
        "--no-daemon",
        "exec",
        "--ephemeral",
        "--ignore-user-config",
        "--skip-git-repo-check",
        "--output-schema",
        str(schema_path),
        "--output-last-message",
        str(output_path),
        "--color",
        "never",
        "-c",
        "model_reasoning_effort=high",
        "-",
    ]
    allowed_env = {
        key: value
        for key, value in os.environ.items()
        if key.upper()
        in {
            "PATH",
            "PATHEXT",
            "SYSTEMROOT",
            "WINDIR",
            "USERPROFILE",
            "APPDATA",
            "LOCALAPPDATA",
            "CODEX_HOME",
            "TEMP",
            "TMP",
            "LANG",
            "LC_ALL",
        }
    }
    try:
        completed = subprocess.run(
            command, input=prompt, text=True, capture_output=True, timeout=2400, env=allowed_env
        )
    except subprocess.TimeoutExpired as exc:
        raise RuntimeErrorSafe(
            "local Codex enrichment timed out; version remains queued for retry"
        ) from exc
    if completed.returncode != 0 or not output_path.is_file():
        raise RuntimeErrorSafe("local Codex enrichment did not produce a structured result")
    try:
        value = json.loads(output_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise RuntimeErrorSafe("local Codex enrichment output was not valid JSON") from exc
    if not isinstance(value, dict):
        raise RuntimeErrorSafe("local Codex enrichment output root must be an object")
    return value


def _valid_existing_enrichment(existing: dict[str, Any] | None) -> bool:
    if not existing or not isinstance(existing.get("annotation"), dict):
        return False
    bilingual = existing.get("bilingual_metadata")
    if not isinstance(bilingual, dict) or set(bilingual) != {"zh", "en"}:
        return False
    return all(
        isinstance(bilingual[lang].get("title"), str)
        and bilingual[lang]["title"].strip()
        and isinstance(bilingual[lang].get("summary"), str)
        and bilingual[lang]["summary"].strip()
        and bilingual[lang].get("category")
        == existing["annotation"].get("categories", [{}])[0].get("key")
        and isinstance(bilingual[lang].get("keywords"), list)
        and bilingual[lang]["keywords"]
        for lang in ("zh", "en")
    )


def _capture(
    site: Site,
    state: dict[str, Any],
    crawler: _BudgetCrawler,
    data_root: Path,
    site_state: dict[str, Any],
    *,
    initial_frontier: list[dict[str, Any]] | None = None,
    request_limit: int = MAX_SITE_REQUESTS,
) -> dict[str, Any]:
    """Enumerate one resumable source frontier, capturing verified files/pages."""
    if site_state.get("completed_frontier"):
        site_state["cycle"] = int(site_state.get("cycle") or 0) + 1
        site_state["completed_frontier"] = False
        site_state["queue"] = initial_frontier or _initial_frontier(site)
        site_state["visited"] = []
        site_state["cursor"] = 0
        site_state["retry_queue"] = []
    if not site_state.get("queue"):
        site_state["queue"] = initial_frontier or _initial_frontier(site)
        site_state["visited"] = []
        site_state["cursor"] = 0
        site_state["retry_queue"] = []
    maximum = (
        MAX_SWISS_SITE_REQUESTS
        if site.site_id == "swiss-association-of-actuaries"
        else MAX_SITE_REQUESTS
    )
    request_limit = int(request_limit)
    if not 1 <= request_limit <= maximum:
        raise RuntimeErrorSafe(f"request limit for {site.site_id} must be between 1 and {maximum}")
    if initial_frontier is None:
        _repair_saved_frontier(site, site_state)
    # A previous partial/budget result describes that attempt only.
    site_state.pop("partial_reason", None)
    crawler.begin_site(request_limit)
    downloads: list[dict[str, Any]] = []
    errors: list[str] = []
    successful_requests = 0
    pages_visited = 0
    links_discovered = 0
    page_ids: set[str] = set()
    link_ids: set[str] = set()

    def record_page_read(url: str) -> None:
        nonlocal pages_visited
        pages_visited += 1
        page_ids.add(hashlib.sha256(_canonical_url(url).encode("utf-8")).hexdigest())

    queue = site_state["queue"]
    visited = set(site_state.get("visited") or [])
    cursor = int(site_state.get("cursor") or 0)
    queue.extend(site_state.pop("retry_queue", []))
    queued_urls = {_queue_identity(row) for row in queue if isinstance(row, dict)}
    _prioritize_pending_documents(site, queue, cursor)

    if site.site_id == "caa" and any(
        str(row.get("kind") or "").startswith("caa_") or row.get("kind") == "json_post"
        for row in queue[cursor:]
    ):
        # CAA issues a session cookie on its public home page before accepting AJAX POSTs.
        try:
            _body, headers, final_url = crawler._request(
                site.seeds[0], timeout=30, delay_seconds=site.delay
            )
            record_page_read(final_url)
            set_cookie = headers.get("set-cookie") or ""
            crawler.caa_cookie = set_cookie.split(";", 1)[0] if "=" in set_cookie else None
        except Exception as exc:
            errors.append(f"CAA session bootstrap: {type(exc).__name__}: {exc}")
    if site.site_id == "australia" and any(row.get("kind") == "algolia" for row in queue[cursor:]):
        script_url = site_state.get("algolia_script_url") or next(
            (
                row.get("script_url")
                for row in queue[cursor:]
                if row.get("kind") == "algolia" and row.get("script_url")
            ),
            None,
        )
        if script_url:
            try:
                script_body, _headers, final_url = crawler._request(
                    str(script_url), timeout=30, delay_seconds=site.delay
                )
                record_page_read(final_url)
                parsed = _parse_algolia_script(script_body.decode("utf-8", errors="replace"))
                if parsed:
                    crawler.algolia_key = parsed["api_key"]
                else:
                    errors.append("Australia Algolia bootstrap: public search key was not found")
            except Exception as exc:
                errors.append(f"Australia Algolia bootstrap: {type(exc).__name__}: {exc}")

    while cursor < len(queue):
        item = queue[cursor]
        url = _canonical_url(str(item.get("url") or ""))
        work_key = _queue_identity(item)
        if not _allowed_host(url, site) or not _allowed_path(url, site):
            item["outcome"] = "rejected_off_host"
            cursor += 1
            continue
        if work_key in visited:
            cursor += 1
            continue
        queued_before = set(queued_urls)
        try:
            if crawler.site_attempts >= crawler.site_budget:
                raise RequestBudgetExceeded(
                    f"site reached its {crawler.site_budget}-request limit; cursor retained"
                )
            if item.get("kind") == "document":
                from ai_actuarial.capacity import ensure_capacity

                ensure_capacity(path=str(data_root), operation="local_document_download")
                staging, headers, final_url, digest, size = crawler._download_file(
                    url, data_root / "downloads", delay_seconds=site.delay
                )
                raw = staging.read_bytes()
                staging.unlink(missing_ok=True)
                if (
                    size > MAX_ARTICLE_BYTES
                    or len(raw) != size
                    or hashlib.sha256(raw).hexdigest() != digest
                ):
                    raise RuntimeErrorSafe(
                        "download exceeded size limit or failed streamed digest check"
                    )
                if not _allowed_host(final_url, site) or not _site_candidate_allowed(
                    site, final_url, str(item.get("title") or ""), item.get("discovery_metadata")
                ):
                    raise RuntimeErrorSafe(
                        "document redirect left the approved source host or matched an exclusion"
                    )
                kind = _detect_kind(raw, headers.get("content-type", ""), final_url)
                if kind not in {"pdf", "docx", "pptx", "xlsx", "office_ole", "text"}:
                    raise RuntimeErrorSafe("document link returned a non-document response")
                artifact = _persist_cas(data_root, raw, digest)
                record = {
                    "site_id": site.site_id,
                    "url": _canonical_url(final_url),
                    "source_page_url": str(item.get("referrer") or ""),
                    "title": str(
                        item.get("title") or Path(urlsplit(final_url).path).name or site.site_id
                    ),
                    "raw_sha256": digest,
                    "byte_size": size,
                    "content_type": headers.get("content-type", "application/octet-stream"),
                    "last_modified": headers.get("last-modified"),
                    "etag": headers.get("etag"),
                    "artifact_path": str(artifact),
                    "detected_kind": kind,
                    "source_page": str(item.get("referrer") or ""),
                    "discovery_metadata": item.get("discovery_metadata") or {},
                }
                key = _file_key(record["url"], digest)
                prior = state["articles"].get(key)
                if prior:
                    prior.setdefault("discovery_sites", [])
                    if site.site_id not in prior["discovery_sites"]:
                        prior["discovery_sites"].append(site.site_id)
                    downloads.append(prior)
                else:
                    state["articles"][key] = {
                        **record,
                        "stage": "downloaded",
                        "discovery_sites": [site.site_id],
                    }
                    downloads.append(state["articles"][key])
                visited.add(work_key)
                item.pop("last_error", None)
                item.pop("attempts", None)
                cursor += 1
                _prioritize_pending_documents(site, queue, cursor)
                successful_requests += 1
                _atomic_json(_state_path_for_data(state), state)
                continue

            if item.get("kind") in {"json_post", "caa_children", "caa_shared"}:
                if item.get("kind") == "caa_children":
                    form_fields = {"parentId": str(item["parent_id"])}
                elif item.get("kind") == "caa_shared":
                    form_fields = {
                        "pageNo": int(item.get("page") or 1),
                        "pageSize": int(item.get("page_size") or 50),
                        "parentDataType": str(item["parent_id"]),
                        "fileName": "",
                        "upLoaderType": "",
                    }
                else:
                    form_fields = {
                        "pageNo": int(item.get("page") or 1),
                        "pageSize": int(item.get("page_size") or 50),
                    }
                form = urlencode(form_fields).encode("ascii")
                request_headers = {
                    "X-Requested-With": "XMLHttpRequest",
                    "Referer": "https://www.e-caa.org.cn/",
                    "Content-Type": "application/x-www-form-urlencoded",
                }
                if crawler.caa_cookie:
                    request_headers["Cookie"] = crawler.caa_cookie
                body, headers, final_url = crawler._request_post(
                    url, form, request_headers, timeout=30, delay_seconds=site.delay
                )
            elif item.get("kind") == "algolia":
                if not crawler.algolia_key:
                    raise RuntimeErrorSafe(
                        "public Algolia search key could not be read from the site bundle"
                    )
                query = {
                    "query": "",
                    "hitsPerPage": 100,
                    "page": int(item.get("page") or 0),
                    "facetFilters": [[f"contentTypes:{item.get('content_type')}"]],
                    "attributesToRetrieve": [
                        "name",
                        "slug",
                        "contentTypes",
                        "accessLevel",
                        "practiceAreas",
                        "tags",
                        "year",
                    ],
                }
                request_headers = {
                    "X-Algolia-API-Key": crawler.algolia_key,
                    "X-Algolia-Application-Id": str(item.get("app_id") or "54YLV1QRFB"),
                    "Content-Type": "application/json",
                }
                body, headers, final_url = crawler._request_post(
                    url,
                    json.dumps(query).encode("utf-8"),
                    request_headers,
                    timeout=35,
                    delay_seconds=site.delay,
                )
            else:
                body, headers, final_url = crawler._request(
                    url, timeout=40, delay_seconds=site.delay
                )
            if not _allowed_host(final_url, site):
                raise RuntimeErrorSafe("source page redirect left the approved host list")
            successful_requests += 1
            if item.get("kind") != "document":
                record_page_read(final_url)
            content_type = headers.get("content-type", "")
            if item.get("kind") in {"json_post", "caa_children", "caa_shared"}:
                _enqueue_caa_results(site, item, body, final_url, queue, queued_urls)
            elif item.get("kind") == "json_rest":
                _enqueue_wordpress_results(site, item, body, headers, final_url, queue, queued_urls)
            elif item.get("kind") == "algolia":
                _enqueue_australia_results(site, item, body, final_url, queue, queued_urls)
            elif item.get("kind") == "json_media":
                try:
                    media = json.loads(body.decode("utf-8"))
                except (UnicodeDecodeError, json.JSONDecodeError) as exc:
                    raise RuntimeErrorSafe(
                        "WordPress media endpoint returned invalid JSON"
                    ) from exc
                source_url = media.get("source_url") if isinstance(media, dict) else None
                title_value = media.get("title") if isinstance(media, dict) else None
                media_title = (
                    title_value.get("rendered") if isinstance(title_value, dict) else title_value
                )
                if source_url:
                    _enqueue_url(
                        site,
                        queue,
                        queued_urls,
                        str(source_url),
                        title=str(media_title or item.get("title") or ""),
                        referrer=str(item.get("referrer") or final_url),
                        kind="document",
                    )
            elif item.get("kind") == "algolia_bootstrap":
                scripts = _extract_search_script_urls(
                    body.decode("utf-8", errors="replace"), final_url
                )
                for script_url in scripts:
                    script_item = {
                        "url": script_url,
                        "kind": "algolia_config",
                        "depth": 1,
                        "referrer": final_url,
                    }
                    identity = _queue_identity(script_item)
                    if identity not in queued_urls:
                        queue.append(script_item)
                        queued_urls.add(identity)
            elif (
                item.get("kind") == "sitemap"
                or "xml" in content_type
                or b"<urlset" in body[:512].lower()
                or b"<sitemapindex" in body[:512].lower()
            ):
                try:
                    root = ET.fromstring(body)
                    locs = [
                        el.text.strip()
                        for el in root.iter()
                        if el.tag.rsplit("}", 1)[-1].lower() == "loc"
                        and el.text
                        and el.text.strip()
                    ]
                except ET.ParseError as exc:
                    raise RuntimeErrorSafe("sitemap XML did not parse") from exc
                for loc in locs:
                    candidate = _canonical_http_link(final_url, loc)
                    if candidate is None:
                        continue
                    if (
                        not _allowed_host(candidate, site)
                        or not _allowed_path(candidate, site)
                        or candidate in queued_urls
                        or candidate in visited
                    ):
                        continue
                    is_document = _looks_site_document(candidate, "", site)
                    is_map = "sitemap" in urlsplit(candidate).path.lower() and not is_document
                    if is_map or is_document or _relevant_page(candidate, "", site):
                        queue.append(
                            {
                                "url": candidate,
                                "kind": (
                                    "sitemap" if is_map else "document" if is_document else "page"
                                ),
                                "depth": 1,
                                "referrer": final_url,
                            }
                        )
                        queued_urls.add(candidate)
            else:
                parser = _LinkParser()
                parser.feed(body.decode("utf-8", errors="replace"))
                if site.site_id == "swiss-association-of-actuaries":
                    for document in _parse_swiss_embedded_documents(
                        body.decode("utf-8", errors="replace")
                    ):
                        candidate = _canonical_http_link(final_url, document["url"])
                        if candidate is None:
                            continue
                        if (
                            _allowed_host(candidate, site)
                            and candidate not in queued_urls
                            and _swiss_document_is_relevant(document)
                        ):
                            queue.append(
                                {
                                    "url": candidate,
                                    "kind": "document",
                                    "depth": 1,
                                    "referrer": final_url,
                                    "title": document["title"],
                                }
                            )
                            queued_urls.add(candidate)
                if item.get("kind") == "algolia_config":
                    parsed = _parse_algolia_script(body.decode("utf-8", errors="replace"))
                    if not parsed:
                        raise RuntimeErrorSafe(
                            "Australia search bundle lacks a readable public Algolia configuration"
                        )
                    crawler.algolia_key = parsed["api_key"]
                    site_state["algolia_script_url"] = url
                    for content_type_name in AUSTRALIA_RESEARCH_TYPES:
                        endpoint = f"https://{parsed['app_id']}-dsn.algolia.net/1/indexes/{parsed['index']}/query"
                        query_item = {
                            "url": endpoint,
                            "kind": "algolia",
                            "content_type": content_type_name,
                            "page": 0,
                            "app_id": parsed["app_id"],
                            "script_url": url,
                            "referrer": final_url,
                        }
                        identity = _queue_identity(query_item)
                        if identity not in queued_urls:
                            queue.append(query_item)
                            queued_urls.add(identity)
                for href, label in parser.links:
                    candidate = _canonical_http_link(final_url, href)
                    if candidate is None:
                        continue
                    if (
                        not _allowed_host(candidate, site)
                        or not _allowed_path(candidate, site)
                        or candidate in queued_urls
                        or candidate in visited
                    ):
                        continue
                    if any(
                        token in urlsplit(candidate).path.lower()
                        for token in ("/private-files/", "/member-only/")
                    ):
                        continue
                    if _looks_site_document(candidate, label, site):
                        if site.site_id == "iaa" and not _site_candidate_allowed(
                            site, candidate, label
                        ):
                            continue
                        queue.append(
                            {
                                "url": candidate,
                                "kind": "document",
                                "depth": int(item.get("depth") or 0) + 1,
                                "referrer": final_url,
                                "title": label,
                            }
                        )
                        queued_urls.add(candidate)
                    elif int(item.get("depth") or 0) < site.max_depth and _relevant_page(
                        candidate, label, site
                    ):
                        queue.append(
                            {
                                "url": candidate,
                                "kind": "page",
                                "depth": int(item.get("depth") or 0) + 1,
                                "referrer": final_url,
                                "title": label,
                            }
                        )
                        queued_urls.add(candidate)
                # Direct publication pages are useful source records; nav/list pages are not.
                if int(item.get("depth") or 0) > 0 and _relevant_page(
                    url, str(item.get("title") or ""), site
                ):
                    kind = _detect_kind(body, content_type, final_url)
                    if kind == "html":
                        markdown = WebPageCollector(
                            SimpleNamespace(data_root=data_root), str(data_root)
                        )._extract_text(body.decode("utf-8", errors="replace"), final_url)
                        if markdown and len(markdown.strip()) >= 100:
                            raw_sha = hashlib.sha256(body).hexdigest()
                            artifact = _persist_cas(data_root, body, raw_sha)
                            record = {
                                "site_id": site.site_id,
                                "url": _canonical_url(final_url),
                                "source_page_url": str(item.get("referrer") or ""),
                                "title": str(
                                    item.get("title")
                                    or urlsplit(final_url).path.rsplit("/", 1)[-1]
                                    or site.site_id
                                ),
                                "raw_sha256": raw_sha,
                                "byte_size": len(body),
                                "content_type": content_type,
                                "last_modified": headers.get("last-modified"),
                                "etag": headers.get("etag"),
                                "artifact_path": str(artifact),
                                "detected_kind": "html",
                                "preset_markdown": markdown,
                            }
                            key = _file_key(record["url"], raw_sha)
                            if key not in state["articles"]:
                                state["articles"][key] = {
                                    **record,
                                    "stage": "downloaded",
                                    "discovery_sites": [site.site_id],
                                }
                                downloads.append(state["articles"][key])
                            else:
                                downloads.append(state["articles"][key])
            # Parsing and adapter work must finish before this item becomes visited.
            # Failed JSON/XML parsing remains eligible for the durable retry queue.
            if item.get("kind") != "document":
                discovered = queued_urls - queued_before
                links_discovered += len(discovered)
                link_ids.update(
                    hashlib.sha256(identity.encode("utf-8")).hexdigest() for identity in discovered
                )
            visited.add(work_key)
            cursor += 1
            _prioritize_pending_documents(site, queue, cursor)
            site_state["cursor"] = cursor
            site_state["visited"] = sorted(visited)
            _atomic_json(_state_path_for_data(state), state)
        except RequestBudgetExceeded as exc:
            site_state["cursor"] = cursor
            site_state["visited"] = sorted(visited)
            site_state["partial_reason"] = str(exc)
            break
        except Exception as exc:
            errors.append(f"{url}: {type(exc).__name__}: {exc}")
            if isinstance(item, dict):
                item["last_error"] = str(exc)[:500]
                item["attempts"] = int(item.get("attempts") or 0) + 1
            if item.get("attempts", 0) >= 3:
                visited.add(work_key)
            else:
                retry = list(site_state.get("retry_queue") or [])
                retry.append(item)
                site_state["retry_queue"] = retry
            cursor += 1
            _prioritize_pending_documents(site, queue, cursor)
            site_state["cursor"] = cursor
            site_state["visited"] = sorted(visited)
            _atomic_json(_state_path_for_data(state), state)

    site_state["cursor"] = cursor
    site_state["visited"] = sorted(visited)
    site_state["queue"] = queue
    site_state["last_request_count"] = crawler.site_attempts
    site_state["last_successful_requests"] = successful_requests
    if cursor >= len(queue) and not site_state.get("retry_queue"):
        site_state["completed_frontier"] = True
    file_types = {"html": 0, "pdf": 0, "other": 0}
    file_ids = set()
    for article in downloads:
        kind = article.get("detected_kind")
        file_types["html" if kind == "html" else "pdf" if kind == "pdf" else "other"] += 1
        file_ids.add(_article_identity(article))
    return {
        "downloads": downloads,
        "errors": errors,
        "request_count": crawler.site_attempts,
        "request_limit": crawler.site_budget,
        "successful_requests": successful_requests,
        "budget_exhausted": bool(site_state.get("partial_reason")),
        "frontier_remaining": max(0, len(queue) - cursor),
        "telemetry": {
            "request_count": crawler.site_attempts,
            "pages_visited": pages_visited,
            "links_discovered": links_discovered,
            "page_ids": sorted(page_ids),
            "link_ids": sorted(link_ids),
            "files_by_type": file_types,
            "file_ids": sorted(file_ids),
        },
    }


# State writes from _capture use the active path registered for the process.
_ACTIVE_STATE_PATH: Path | None = None


def _state_path_for_data(_state: dict[str, Any]) -> Path:
    if _ACTIVE_STATE_PATH is None:
        raise RuntimeErrorSafe("runtime state path has not been initialized")
    return _ACTIVE_STATE_PATH


def _convert_article(article: dict[str, Any], data_root: Path) -> str:
    preset = article.get("preset_markdown")
    if isinstance(preset, str) and preset.strip():
        markdown = preset.strip()
    else:
        raw = Path(article["artifact_path"]).read_bytes()
        if hashlib.sha256(raw).hexdigest() != article["raw_sha256"]:
            raise RuntimeErrorSafe("source CAS readback failed SHA-256 verification")
        markdown = _extract_markdown(
            raw, article.get("content_type") or "", article["url"], data_root
        )
    if not markdown or not markdown.strip():
        raise RuntimeErrorSafe("source Markdown is empty")
    return markdown


def _verified_article_raw_bytes(article: dict[str, Any]) -> bytes:
    raw = Path(article["artifact_path"]).read_bytes()
    if hashlib.sha256(raw).hexdigest() != article["raw_sha256"]:
        raise RuntimeErrorSafe("source CAS readback failed SHA-256 verification")
    return raw


def _ingest_live_capture(
    article: dict[str, Any],
    config: dict[str, Any],
    hub: DatabaseConfig,
    raw: bytes,
    markdown: str | None,
) -> dict[str, Any]:
    return ingest_live_article(
        canonical_url=article["url"],
        source_site=article["site_id"],
        source_page_url=article.get("source_page_url") or article["url"],
        title=article.get("title") or article["url"],
        raw_bytes=raw,
        markdown=markdown,
        content_type=article.get("content_type") or "application/octet-stream",
        data_root=config["data_root"],
        database_url=hub,
        metadata={
            "last_modified": article.get("last_modified"),
            "etag": article.get("etag"),
            "detected_kind": article.get("detected_kind"),
        },
    )


def _prepare_hub_article(
    article: dict[str, Any],
    state: dict[str, Any],
    config: dict[str, Any],
    hub: DatabaseConfig,
    preview: bool,
    *,
    on_hub_version: Callable[[bool, dict[str, Any]], None] | None = None,
    on_conversion: Callable[[bool], None] | None = None,
) -> tuple[dict[str, Any], bool, bool]:
    def convert_article() -> str:
        try:
            markdown = _convert_article(article, Path(config["data_root"]).expanduser().resolve())
        except Exception:
            if on_conversion:
                on_conversion(False)
            raise
        if on_conversion:
            on_conversion(True)
        return markdown

    raw_sha = str(article["raw_sha256"])
    existing = content_hub_read_version(article["url"], raw_sha, hub)
    if existing and existing.get("excluded_reason"):
        raise RuntimeErrorSafe(f"Hub excludes this URL as {existing['excluded_reason']}")
    if not existing:
        exclusion = content_hub_url_exclusion(article["url"], hub)
        if exclusion:
            raise RuntimeErrorSafe(f"Hub excludes this URL as {exclusion['reason']}")
    newly_enriched = False
    new_version = existing is None
    hub_version_reported = False
    if preview:
        if (
            existing
            and isinstance(existing.get("markdown"), str)
            and existing.get("markdown_text_variant_id")
        ):
            markdown = existing["markdown"]
        else:
            try:
                markdown = convert_article()
            except Exception as exc:
                detected_kind = str(article.get("detected_kind") or "unknown")
                raise RuntimeErrorSafe(
                    f"extraction gap ({detected_kind}): {type(exc).__name__}: {exc}"
                ) from exc
        if not existing:
            return (
                {
                    **article,
                    "markdown": markdown,
                    "content_version_id": None,
                    "markdown_text_variant_id": None,
                    "markdown_sha256": hashlib.sha256(markdown.encode("utf-8")).hexdigest(),
                    "annotation": None,
                    "bilingual_metadata": {},
                },
                True,
                False,
            )
        if not existing.get("markdown_text_variant_id"):
            return (
                {
                    **article,
                    "markdown": markdown,
                    "content_version_id": existing["content_version_id"],
                    "markdown_text_variant_id": None,
                    "markdown_sha256": hashlib.sha256(markdown.encode("utf-8")).hexdigest(),
                    "annotation": None,
                    "bilingual_metadata": {},
                },
                new_version,
                False,
            )
    elif not existing:
        raw_bytes = _verified_article_raw_bytes(article)
        capture_result = _ingest_live_capture(article, config, hub, raw_bytes, None)
        existing = content_hub_read_version(article["url"], raw_sha, hub)
        if (
            not existing
            or existing["content_version_id"] != capture_result["content_version_id"]
            or existing.get("raw_sha256") != raw_sha
            or not existing.get("blob_available")
            or existing.get("markdown_text_variant_id") is not None
        ):
            raise RuntimeErrorSafe(
                "Hub raw-only version readback did not confirm the original capture"
            )
        new_version = not bool(
            capture_result.get("already_versioned") or capture_result.get("already_imported")
        )
        if on_hub_version:
            on_hub_version(
                new_version,
                {
                    "content_version_id": existing["content_version_id"],
                    "markdown_text_variant_id": None,
                    "markdown_sha256": None,
                },
            )
            hub_version_reported = True
        try:
            markdown = convert_article()
        except Exception as exc:
            detected_kind = str(article.get("detected_kind") or "unknown")
            raise RuntimeErrorSafe(
                f"extraction gap ({detected_kind}): {type(exc).__name__}: {exc}"
            ) from exc
    elif isinstance(existing.get("markdown"), str) and existing.get("markdown_text_variant_id"):
        markdown = existing["markdown"]
    else:
        try:
            markdown = convert_article()
        except Exception as exc:
            detected_kind = str(article.get("detected_kind") or "unknown")
            raise RuntimeErrorSafe(
                f"extraction gap ({detected_kind}): {type(exc).__name__}: {exc}"
            ) from exc

    if not preview and not existing.get("markdown_text_variant_id"):
        raw_bytes = _verified_article_raw_bytes(article)
        markdown_capture = _ingest_live_capture(article, config, hub, raw_bytes, markdown)
        if markdown_capture.get("content_version_id") != existing["content_version_id"]:
            raise RuntimeErrorSafe(
                "Hub Markdown capture changed the immutable raw version identity"
            )
        existing = content_hub_read_version(article["url"], raw_sha, hub)
        if (
            not existing
            or existing.get("content_version_id") != markdown_capture["content_version_id"]
            or not existing.get("markdown_text_variant_id")
            or not isinstance(existing.get("markdown"), str)
        ):
            raise RuntimeErrorSafe("Hub Markdown exact readback did not complete the raw version")
        markdown = existing["markdown"]

    if not preview and not hub_version_reported and on_hub_version:
        on_hub_version(
            new_version,
            {
                "content_version_id": existing["content_version_id"],
                "markdown_text_variant_id": existing.get("markdown_text_variant_id"),
                "markdown_sha256": existing.get("markdown_sha256"),
                "hub_markdown_sha256": existing.get("hub_markdown_sha256"),
                "hub_markdown_sha256_kind": existing.get("hub_markdown_sha256_kind"),
            },
        )
        hub_version_reported = True
    if (
        not existing
        or not existing.get("markdown_text_variant_id")
        or not isinstance(markdown, str)
    ):
        raise RuntimeErrorSafe("Hub version has no reusable bound Markdown text variant")
    md_sha = str(
        existing.get("markdown_sha256") or hashlib.sha256(markdown.encode("utf-8")).hexdigest()
    )
    if hashlib.sha256(markdown.encode("utf-8")).hexdigest() != md_sha:
        raise RuntimeErrorSafe("Hub Markdown version/hash binding failed")
    if not _valid_existing_enrichment(existing):
        if preview:
            return (
                {
                    **article,
                    "markdown": markdown,
                    "content_version_id": existing["content_version_id"],
                    "markdown_text_variant_id": existing["markdown_text_variant_id"],
                    "markdown_sha256": md_sha,
                    "hub_markdown_sha256": existing.get("hub_markdown_sha256"),
                    "hub_markdown_sha256_kind": existing.get("hub_markdown_sha256_kind"),
                    "annotation": None,
                    "bilingual_metadata": {},
                },
                new_version,
                False,
            )
        evidence_spans = _build_evidence_span_catalog(markdown)
        raw_model_result = _run_model(
            markdown,
            article.get("title") or article["url"],
            Path(config["state_root"]).expanduser().resolve(),
            f"{existing['content_version_id']}_{md_sha[:16]}",
            config.get("codex_executable"),
            evidence_spans=evidence_spans,
        )

        def resolve_model_result(raw_result: dict[str, Any]) -> tuple[dict[str, Any], list[str]]:
            resolved, unresolved_ids = _resolve_model_evidence_span_ids(raw_result, evidence_spans)
            errors = unresolved_ids + _model_evidence_errors(markdown, resolved)
            return resolved, list(dict.fromkeys(errors))

        model_result, invalid_evidence = resolve_model_result(raw_model_result)
        if invalid_evidence:
            raw_model_result = _run_model(
                markdown,
                article.get("title") or article["url"],
                Path(config["state_root"]).expanduser().resolve(),
                f"{existing['content_version_id']}_{md_sha[:16]}",
                config.get("codex_executable"),
                evidence_feedback=invalid_evidence,
                evidence_spans=evidence_spans,
            )
            model_result, invalid_evidence = resolve_model_result(raw_model_result)
            if invalid_evidence:
                raise RuntimeErrorSafe(
                    "model evidence failed the exact Markdown quote check after one correction attempt: "
                    + ", ".join(invalid_evidence)
                )
        if model_result.get("primary_category") not in _load_category_keys():
            raise RuntimeErrorSafe("model returned a noncanonical primary category")
        saved = save_live_enrichment(
            content_version_id=existing["content_version_id"],
            markdown_text_variant_id=existing["markdown_text_variant_id"],
            markdown_sha256=md_sha,
            primary_category=model_result["primary_category"],
            bilingual_metadata=model_result["bilingual_metadata"],
            categories=model_result["categories"],
            date_facts=model_result["date_facts"],
            date_review_status=model_result["date_review_status"],
            date_review_reason=model_result.get("date_review_reason"),
            date_review_unresolved=model_result["date_review_unresolved"],
            recipe=RECIPE,
            run_id=f"{dt.datetime.now(dt.timezone.utc).strftime('%Y%m%dT%H%M%SZ')}-{existing['content_version_id']}",
            database_url=hub,
        )
        if not saved.get("success"):
            raise RuntimeErrorSafe(
                "Hub enrichment writer did not confirm the metadata and annotation"
            )
        existing = content_hub_read_version(article["url"], raw_sha, hub)
        if not existing or not _valid_existing_enrichment(existing):
            raise RuntimeErrorSafe("Hub enrichment version/hash readback failed")
        newly_enriched = True
    annotation = existing["annotation"]
    categories = annotation.get("categories") or []
    category_keys = [item.get("key") for item in categories if isinstance(item, dict)]
    source_primary_category = str(annotation.get("source_primary_category") or "")
    return (
        {
            **article,
            "markdown": existing["markdown"],
            "content_version_id": existing["content_version_id"],
            "markdown_text_variant_id": existing["markdown_text_variant_id"],
            "markdown_sha256": existing["markdown_sha256"],
            "hub_markdown_sha256": existing.get("hub_markdown_sha256"),
            "hub_markdown_sha256_kind": existing.get("hub_markdown_sha256_kind"),
            "annotation": annotation,
            "bilingual_metadata": existing["bilingual_metadata"],
            "categories": categories,
            "date_facts": annotation.get("date_facts") or {},
            "date_review_status": annotation.get("date_review_status"),
            "date_review_reason": annotation.get("date_review_reason"),
            "date_review_unresolved": annotation.get("date_review_unresolved") or [],
            "source_primary_category": source_primary_category,
            "source_signature": annotation.get("source_signature"),
            "result_sha256": annotation.get("result_sha256"),
            "ai_candidate": bool(
                category_keys and category_keys[0] == "AI" and source_primary_category == "AI"
            ),
            "legacy_hub_annotation": not newly_enriched,
        },
        new_version,
        newly_enriched,
    )


def _target_ack_key(target: str, item: dict[str, Any]) -> str:
    return f"{target}:{item['content_version_id']}:{item['raw_sha256']}"


def _target_needs_item(state: dict[str, Any], target: str, item: dict[str, Any]) -> bool:
    ack = state["target_acks"].get(_target_ack_key(target, item))
    return not (
        isinstance(ack, dict)
        and ack.get("status") == "acknowledged"
        and ack.get("metadata_fingerprint") == _target_metadata_fingerprint(item)
    )


def _target_metadata_fingerprint(item: dict[str, Any]) -> str:
    value = {
        key: item.get(key)
        for key in (
            "content_version_id",
            "raw_sha256",
            "markdown_text_variant_id",
            "markdown_sha256",
            "hub_markdown_sha256",
            "hub_markdown_sha256_kind",
            "source_site",
            "bilingual_metadata",
            "categories",
            "source_primary_category",
            "date_facts",
            "date_review_status",
            "date_review_reason",
            "date_review_unresolved",
            "source_signature",
            "result_sha256",
        )
    }
    publication_date = (item.get("date_facts") or {}).get("publication_date")
    projected_date = publication_date.get("value") if isinstance(publication_date, dict) else None
    # Keep the historical fingerprint for unknown dates; a verified date must
    # invalidate old acknowledgements so target published_time is backfilled.
    if projected_date is not None:
        value["published_time"] = projected_date
    return hashlib.sha256(
        json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
    ).hexdigest()


def _invoke_target_worker(
    config: dict[str, Any], target: str, payload: dict[str, Any]
) -> dict[str, Any]:
    target_cfg = config["targets"].get(target)
    if not isinstance(target_cfg, dict) or not target_cfg.get("host"):
        raise RuntimeErrorSafe(f"target {target} has no configured SSH host")
    worker_code = TARGET_WORKER.read_bytes()
    b64_code = base64.b64encode(worker_code).decode("ascii")
    one_shot = (
        b64_code + "\n" + json.dumps(payload, ensure_ascii=False, separators=(",", ":")) + "\n"
    )
    container = str(target_cfg.get("container") or "ai-api")
    remote_python = 'import sys,base64;exec(compile(base64.b64decode(sys.stdin.readline()),"<local-target-worker>","exec"))'
    remote_tokens = (["sudo", "-n"] if target_cfg.get("sudo") is True else []) + [
        "docker",
        "exec",
        "-i",
        container,
        "python",
        "-c",
        remote_python,
    ]
    remote = shlex.join(remote_tokens)
    try:
        result = subprocess.run(
            [
                "ssh",
                "-o",
                "BatchMode=yes",
                "-o",
                "ConnectTimeout=20",
                "-o",
                "ServerAliveInterval=30",
                "-o",
                "ServerAliveCountMax=3",
                str(target_cfg["host"]),
                remote,
            ],
            input=one_shot,
            text=True,
            capture_output=True,
            timeout=int(target_cfg.get("timeout_seconds") or 21600),
        )
    except (OSError, subprocess.TimeoutExpired) as exc:
        raise RuntimeErrorSafe(
            f"{target} target worker could not complete: {type(exc).__name__}"
        ) from exc
    lines = [line for line in result.stdout.splitlines() if line.strip().startswith("{")]
    if not lines:
        message = f"{target} target worker returned no JSON result (exit {result.returncode})"
        details: dict[str, Any] = {"exit_code": result.returncode}
        stderr = result.stderr.strip()
        if stderr:
            details["stderr"] = stderr[-2000:]
        stdout = result.stdout.strip()
        if stdout:
            details["stdout_tail"] = stdout[-1000:]
        raise _TargetWorkerFailure(
            target,
            {"target": None, "error": message, "error_details": details},
            message,
        )
    try:
        reply = json.loads(lines[-1])
    except json.JSONDecodeError as exc:
        message = f"{target} target worker returned malformed JSON"
        details = {"parse_error": str(exc), "exit_code": result.returncode}
        stderr = result.stderr.strip()
        if stderr:
            details["stderr"] = stderr[-2000:]
        details["stdout_tail"] = result.stdout[-1000:]
        raise _TargetWorkerFailure(
            target,
            {"target": None, "error": message, "error_details": details},
            message,
        ) from exc
    if isinstance(reply, dict) and reply.get("error"):
        raise _TargetWorkerFailure(target, reply, str(reply["error"]))
    if (
        result.returncode not in (0, 2)
        or not isinstance(reply, dict)
        or reply.get("target") != target
    ):
        raise RuntimeErrorSafe(f"{target} target worker did not confirm its own target identity")
    return reply


def _stream_target(
    config: dict[str, Any],
    target: str,
    items: list[dict[str, Any]],
    ready_pending: bool,
    *,
    preview: bool = False,
) -> dict[str, Any]:
    payload_items = []
    for item in items:
        raw = Path(item["artifact_path"]).read_bytes()
        if hashlib.sha256(raw).hexdigest() != item["raw_sha256"]:
            raise RuntimeErrorSafe("target source bytes changed after Hub validation")
        payload_items.append(
            {
                key: item.get(key)
                for key in (
                    "url",
                    "raw_sha256",
                    "markdown",
                    "markdown_sha256",
                    "hub_markdown_sha256",
                    "hub_markdown_sha256_kind",
                    "content_version_id",
                    "markdown_text_variant_id",
                    "bilingual_metadata",
                    "categories",
                    "date_facts",
                    "title",
                    "source_site",
                    "source_page_url",
                    "content_type",
                    "last_modified",
                    "etag",
                    "source_primary_category",
                    "date_review_status",
                    "date_review_reason",
                    "date_review_unresolved",
                    "source_signature",
                    "result_sha256",
                    "legacy_hub_annotation",
                )
            }
            | {"raw_bytes_b64": base64.b64encode(raw).decode("ascii")}
        )
    return _invoke_target_worker(
        config,
        target,
        {
            "schema_version": 1,
            "target": target,
            "items": payload_items,
            "ready_pending": bool(ready_pending),
            "preview_only": bool(preview),
            "approved_categories": _load_category_keys(),
        },
    )


def _stream_target_cleanup(
    config: dict[str, Any],
    target: str,
    records: list[dict[str, Any]],
    *,
    dry_run: bool,
) -> dict[str, Any]:
    return _invoke_target_worker(
        config,
        target,
        {
            "schema_version": 1,
            "operation": "retire_non_ai_primary",
            "target": target,
            "dry_run": bool(dry_run),
            "records": records,
            "approved_categories": _load_category_keys(),
        },
    )


def _stream_target_metadata_backfill(
    config: dict[str, Any],
    target: str,
    records: list[dict[str, Any]],
    *,
    dry_run: bool,
) -> dict[str, Any]:
    if not isinstance(records, list) or not 1 <= len(records) <= MAX_RECONCILE_PAGE:
        raise RuntimeErrorSafe("metadata backfill requires one to 100 exact records")
    hub = _connection_url(Path(config["hub_connection_file"]).expanduser())
    for record in records:
        if not isinstance(record, dict):
            raise RuntimeErrorSafe("metadata backfill record must be an object")
        url = record.get("url")
        raw_sha = record.get("raw_sha256")
        if (
            not isinstance(url, str)
            or not url.startswith(("http://", "https://"))
            or not isinstance(raw_sha, str)
            or not re.fullmatch(r"[0-9a-f]{64}", raw_sha)
        ):
            raise RuntimeErrorSafe("metadata backfill requires exact URL/raw-SHA identities")
        current = content_hub_read_version(url, raw_sha, hub)
        annotation = current.get("annotation") if current else None
        if not current or current.get("excluded_reason") or not isinstance(annotation, dict):
            raise RuntimeErrorSafe("metadata manifest has no current exact Hub annotation")
        authoritative = {
            "url": current.get("url"),
            "raw_sha256": current.get("raw_sha256"),
            "content_version_id": current.get("content_version_id"),
            "markdown_text_variant_id": current.get("markdown_text_variant_id"),
            "markdown_sha256": current.get("markdown_sha256"),
            "hub_markdown_sha256": current.get("hub_markdown_sha256"),
            "hub_markdown_sha256_kind": current.get("hub_markdown_sha256_kind"),
            "source_primary_category": annotation.get("source_primary_category"),
            "categories": annotation.get("categories") or [],
            "date_facts": annotation.get("date_facts") or {},
            "date_review_status": annotation.get("date_review_status"),
            "date_review_reason": annotation.get("date_review_reason"),
            "date_review_unresolved": annotation.get("date_review_unresolved") or [],
            "source_signature": annotation.get("source_signature"),
            "result_sha256": annotation.get("result_sha256"),
        }
        if any(record.get(field) != value for field, value in authoritative.items()):
            raise RuntimeErrorSafe("metadata manifest no longer matches its fresh exact Hub read")
    return _invoke_target_worker(
        config,
        target,
        {
            "schema_version": 1,
            "operation": "backfill_ai_primary_metadata",
            "target": target,
            "dry_run": bool(dry_run),
            "records": records,
            "approved_categories": _load_category_keys(),
        },
    )


def _process_articles(
    site_id: str,
    state: dict[str, Any],
    config: dict[str, Any],
    hub: DatabaseConfig,
    *,
    preview: bool,
    article_limit: int = MAX_ENRICHMENTS_PER_SITE_RUN,
    include_url: str | None = None,
) -> dict[str, Any]:
    site_articles = [
        row for row in state["articles"].values() if site_id in (row.get("discovery_sites") or [])
    ]
    site_incomplete_count = sum(row.get("stage") != "complete" for row in site_articles)
    if include_url:
        articles = [
            row
            for row in site_articles
            if row.get("url") == include_url and row.get("stage") != "complete"
        ]
    else:
        cursor_key = state.setdefault("article_cursors", {}).get(site_id)
        if cursor_key:
            prior_index = next(
                (
                    index
                    for index, row in enumerate(site_articles)
                    if _article_identity(row) == cursor_key
                ),
                None,
            )
            if prior_index is not None:
                # Completed entries remain useful cursor anchors across runs.
                site_articles = site_articles[prior_index + 1 :] + site_articles[: prior_index + 1]
        articles = [row for row in site_articles if row.get("stage") != "complete"]
    article_limit = min(MAX_ENRICHMENTS_PER_SITE_RUN, max(1, int(article_limit)))
    scanned_articles = articles[:MAX_ARTICLE_SCAN_PER_SITE_RUN]
    if preview:
        selected_articles = scanned_articles[:article_limit]
        reuse_articles: list[dict[str, Any]] = []
        enrichment_articles = selected_articles
    else:
        reuse_articles = []
        enrichment_articles = []
        for article in scanned_articles:
            existing = None
            if article.get("stage") == "hub_verified" and isinstance(
                article.get("enrichment"), dict
            ):
                existing = article["enrichment"]
            else:
                try:
                    existing = content_hub_read_version(
                        str(article.get("url") or ""),
                        str(article.get("raw_sha256") or ""),
                        hub,
                    )
                except Exception:
                    # The normal per-article path records the concrete failure.
                    existing = None
            if _valid_existing_enrichment(existing):
                reuse_articles.append(article)
            else:
                enrichment_articles.append(article)
        # Existing Hub reuse and target-only retries have a separate finite cap.
        # They cannot consume the model-generation allowance for new enrichment.
        reuse_articles = reuse_articles[:MAX_CHEAP_ARTICLES_PER_SITE_RUN]
        enrichment_articles = enrichment_articles[:article_limit]
        selected_articles = enrichment_articles + reuse_articles
    selected_identities = {_article_identity(article) for article in selected_articles}
    # Stop before the first deferred row so it leads the next bounded scan.
    first_deferred_index = next(
        (
            index
            for index, article in enumerate(scanned_articles)
            if _article_identity(article) not in selected_identities
        ),
        len(scanned_articles),
    )
    cursor_anchor = scanned_articles[first_deferred_index - 1] if first_deferred_index else None
    summary = {
        "candidate_count": len(selected_articles),
        "pending_article_count": max(0, site_incomplete_count - len(selected_articles)),
        "enrichment_candidate_count": len(enrichment_articles),
        "reuse_candidate_count": len(reuse_articles),
        "new_versions": 0,
        "duplicates": 0,
        "enriched": 0,
        "ai_candidates": 0,
        "database": {"new_version_ids": [], "duplicate_version_ids": []},
        "enrichment": {"succeeded_count": 0, "succeeded_ids": []},
        "conversion": {
            "succeeded_count": 0,
            "failed_count": 0,
            "succeeded_ids": [],
            "failed_ids": [],
        },
        "ai_first_candidate_ids": [],
        "article_errors": [],
        "target": {
            name: {
                "acknowledged": 0,
                "pending": 0,
                "failed": 0,
                "failures": [],
                "ready_status": "unknown",
                "ready_error": None,
                "ready_evidence": [],
                "acknowledged_ids": [],
                "ready_ids": [],
                "invalidated_ready_ids": [],
            }
            for name in ("china", "overseas")
        },
    }
    prepared: list[dict[str, Any]] = []
    for article in selected_articles:
        version_counted = False
        # Local v2 state can contain the former any-category-AI flag. The exact
        # Hub read in _prepare_hub_article below is the current eligibility source.
        article["ai_candidate"] = False

        def count_hub_version(is_new: bool, version: dict[str, Any]) -> None:
            nonlocal version_counted
            if version_counted:
                raise RuntimeErrorSafe("Hub version stage was reported more than once")
            summary["new_versions"] += int(is_new)
            summary["duplicates"] += int(not is_new)
            version_id = version.get("content_version_id")
            if (
                isinstance(version_id, (str, int))
                and not isinstance(version_id, bool)
                and str(version_id)
            ):
                version_id = str(version_id)
                field = "new_version_ids" if is_new else "duplicate_version_ids"
                if version_id not in summary["database"][field]:
                    summary["database"][field].append(version_id)
            article["hub_version"] = {
                **version,
                "status": "new" if is_new else "existing",
                "verified_at": dt.datetime.now(dt.timezone.utc).isoformat(),
            }
            article["stage"] = "hub_ingested"
            version_counted = True
            _atomic_json(_state_path_for_data(state), state)

        def count_conversion(succeeded: bool) -> None:
            key = "succeeded" if succeeded else "failed"
            summary["conversion"][f"{key}_count"] += 1
            identity = _article_identity(article)
            if identity not in summary["conversion"][f"{key}_ids"]:
                summary["conversion"][f"{key}_ids"].append(identity)

        try:
            full, is_new, enriched = _prepare_hub_article(
                article,
                state,
                config,
                hub,
                preview,
                on_hub_version=count_hub_version,
                on_conversion=count_conversion,
            )
            if not version_counted:
                summary["new_versions"] += int(is_new)
                summary["duplicates"] += int(not is_new)
                version_id = full.get("content_version_id")
                if (
                    isinstance(version_id, (str, int))
                    and not isinstance(version_id, bool)
                    and str(version_id)
                ):
                    version_id = str(version_id)
                    field = "new_version_ids" if is_new else "duplicate_version_ids"
                    if version_id not in summary["database"][field]:
                        summary["database"][field].append(version_id)
            summary["enriched"] += int(enriched)
            if enriched:
                summary["enrichment"]["succeeded_count"] += 1
                version_id = full.get("content_version_id")
                if (
                    isinstance(version_id, (str, int))
                    and not isinstance(version_id, bool)
                    and str(version_id)
                ):
                    version_id = str(version_id)
                    if version_id not in summary["enrichment"]["succeeded_ids"]:
                        summary["enrichment"]["succeeded_ids"].append(version_id)
            if preview:
                article["stage"] = "preview"
                article["preview"] = {
                    "content_version_id": full.get("content_version_id"),
                    "ai_candidate": full.get("ai_candidate"),
                }
                prepared.append(full)
                continue
            if not full.get("source_site") and article.get("site_id") in SITES:
                full["source_site"] = article["site_id"]
            article.update(
                {
                    "stage": "hub_verified",
                    "content_version_id": full["content_version_id"],
                    "markdown_text_variant_id": full["markdown_text_variant_id"],
                    "markdown_sha256": full["markdown_sha256"],
                    "hub_markdown_sha256": full.get("hub_markdown_sha256"),
                    "hub_markdown_sha256_kind": full.get("hub_markdown_sha256_kind"),
                    "ai_candidate": full["ai_candidate"],
                    "enrichment": full,
                }
            )
            if full["ai_candidate"]:
                summary["ai_candidates"] += 1
                version_id = full.get("content_version_id")
                if (
                    isinstance(version_id, (str, int))
                    and not isinstance(version_id, bool)
                    and str(version_id)
                ):
                    version_id = str(version_id)
                    if version_id not in summary["ai_first_candidate_ids"]:
                        summary["ai_first_candidate_ids"].append(version_id)
                prepared.append(full)
            else:
                categories = full.get("categories") or []
                primary = (
                    categories[0].get("key")
                    if categories and isinstance(categories[0], dict)
                    else "unknown"
                )
                reason = (
                    f"Hub source primary category {full.get('source_primary_category')!r} "
                    "conflicts with categories[0]"
                    if primary == "AI"
                    else f"current Hub primary category is {primary!r}"
                )
                _suppress_non_primary_retry(
                    state,
                    full,
                    reason,
                )
                article["stage"] = "complete"
            _atomic_json(_state_path_for_data(state), state)
        except Exception as exc:
            summary["article_errors"].append(f"{article.get('url')}: {type(exc).__name__}: {exc}")
            article["last_error"] = str(exc)[:500]
            article["stage"] = "pending_retry"
            _atomic_json(_state_path_for_data(state), state)
    if preview:
        return summary
    for target in ("china", "overseas"):
        needs = [item for item in prepared if _target_needs_item(state, target, item)]
        ready_pending = any(
            bool(state["target_acks"].get(_target_ack_key(target, item), {}).get("ready_pending"))
            for item in prepared
        )
        if not needs and not ready_pending:
            if not prepared and summary["article_errors"]:
                # A raw-only capture cannot confirm target membership or Ready.
                # Keep downstream state unknown so an extraction gap cannot
                # close an earlier Ready failure for this site.
                summary["target"][target]["ready_status"] = "unknown"
                continue
            ready_confirmed = True
            for item in prepared:
                state_key = _target_ack_key(target, item)
                ack = state["target_acks"].get(state_key) or {}
                if ack.get("status") != "acknowledged" or ack.get(
                    "metadata_fingerprint"
                ) != _target_metadata_fingerprint(item):
                    continue
                summary["target"][target]["acknowledged"] += 1
                summary["target"][target]["acknowledged_ids"].append(state_key)
                if ack.get("ready_pending") is False:
                    summary["target"][target]["ready_ids"].append(state_key)
                else:
                    ready_confirmed = False
            summary["target"][target]["ready_status"] = (
                "succeeded" if ready_confirmed else "unknown"
            )
            continue
        for item in needs:
            state_key = _target_ack_key(target, item)
            state["target_acks"][state_key] = {
                "status": "pending",
                "metadata_fingerprint": _target_metadata_fingerprint(item),
                "ready_pending": True,
            }
        _atomic_json(_state_path_for_data(state), state)
        try:
            reply = _stream_target(config, target, needs, ready_pending)
            ack_by_version = {
                (row.get("content_version_id"), row.get("raw_sha256")): row
                for row in reply.get("acknowledgements") or []
            }
            for item in prepared:
                ack = ack_by_version.get((item["content_version_id"], item["raw_sha256"]))
                state_key = _target_ack_key(target, item)
                if (
                    ack
                    and ack.get("status") in {"written", "already_current"}
                    and ack.get("metadata_fingerprint") == _target_metadata_fingerprint(item)
                ):
                    state["target_acks"][state_key] = {
                        "status": "acknowledged",
                        "metadata_fingerprint": ack["metadata_fingerprint"],
                        "ready_pending": reply.get("ready_status") != "succeeded",
                        "ack": ack,
                    }
                    summary["target"][target]["acknowledged"] += 1
                elif (
                    reply.get("ready_status") == "succeeded"
                    and isinstance(state["target_acks"].get(state_key), dict)
                    and state["target_acks"][state_key].get("status") == "acknowledged"
                    and state["target_acks"][state_key].get("metadata_fingerprint")
                    == _target_metadata_fingerprint(item)
                ):
                    state["target_acks"][state_key]["ready_pending"] = False
                    summary["target"][target]["acknowledged"] += 1
                elif (
                    reply.get("ready_status") != "succeeded"
                    and isinstance(state["target_acks"].get(state_key), dict)
                    and state["target_acks"][state_key].get("status") == "acknowledged"
                    and state["target_acks"][state_key].get("metadata_fingerprint")
                    == _target_metadata_fingerprint(item)
                ):
                    # A Ready-only retry cannot revoke the earlier per-item
                    # readback acknowledgement. Keep that ack and report the
                    # failed Ready stage separately.
                    state["target_acks"][state_key]["ready_pending"] = True
                    summary["target"][target]["acknowledged"] += 1
                else:
                    state["target_acks"].setdefault(
                        state_key,
                        {
                            "status": "pending",
                            "metadata_fingerprint": _target_metadata_fingerprint(item),
                            "ready_pending": True,
                        },
                    )
                    summary["target"][target]["pending"] += 1
            target_failures = reply.get("failures") or []
            summary["target"][target]["failed"] = len(target_failures)
            summary["target"][target]["failures"] = [
                {"url": failure["url"], "error": failure["error"]}
                for failure in target_failures
                if isinstance(failure, dict)
                and isinstance(failure.get("url"), str)
                and failure["url"].strip()
                and isinstance(failure.get("error"), str)
                and failure["error"].strip()
            ]
            summary["target"][target]["ready_status"] = reply.get("ready_status") or "unknown"
            summary["target"][target]["ready_error"] = _target_ready_error_text(
                reply.get("ready_error")
            )
            summary["target"][target]["ready_evidence"] = _target_ready_evidence(
                reply.get("ready_evidence")
            )
            if reply.get("ready_status") == "succeeded":
                for item in prepared:
                    state_key = _target_ack_key(target, item)
                    if state_key in state["target_acks"]:
                        state["target_acks"][state_key]["ready_pending"] = False
        except Exception as exc:
            summary["target"][target]["failed"] = max(1, len(needs))
            summary["target"][target]["pending"] = len(prepared)
            summary["target"][target]["ready_status"] = "failed"
            summary["target"][target]["ready_error"] = f"{type(exc).__name__}: {exc}"
            summary["target"][target]["ready_evidence"] = list(
                getattr(exc, "diagnostic_evidence", []) or []
            )
            for item in prepared:
                state_key = _target_ack_key(target, item)
                state["target_acks"].setdefault(
                    state_key,
                    {
                        "status": "pending",
                        "metadata_fingerprint": _target_metadata_fingerprint(item),
                        "ready_pending": True,
                    },
                )
        _atomic_json(_state_path_for_data(state), state)
        for item in prepared:
            state_key = _target_ack_key(target, item)
            ack = state["target_acks"].get(state_key) or {}
            if ack.get("status") == "acknowledged" and ack.get(
                "metadata_fingerprint"
            ) == _target_metadata_fingerprint(item):
                if state_key not in summary["target"][target]["acknowledged_ids"]:
                    summary["target"][target]["acknowledged_ids"].append(state_key)
                if (
                    ack.get("ready_pending") is False
                    and state_key not in summary["target"][target]["ready_ids"]
                ):
                    summary["target"][target]["ready_ids"].append(state_key)
            elif ack.get("status") == "acknowledged" and ack.get("ready_pending"):
                if state_key not in summary["target"][target]["invalidated_ready_ids"]:
                    summary["target"][target]["invalidated_ready_ids"].append(state_key)
    for article in selected_articles:
        if article.get("ai_candidate"):
            flags = [
                state["target_acks"].get(_target_ack_key(target, article), {}).get("status")
                == "acknowledged"
                and state["target_acks"]
                .get(_target_ack_key(target, article), {})
                .get("ready_pending")
                is False
                for target in ("china", "overseas")
            ]
            if all(flags):
                article["stage"] = "complete"
        elif article.get("stage") == "hub_verified":
            article["stage"] = "complete"
    _atomic_json(_state_path_for_data(state), state)
    if not preview and not include_url and cursor_anchor is not None:
        state.setdefault("article_cursors", {})[site_id] = _article_identity(cursor_anchor)
        _atomic_json(_state_path_for_data(state), state)
    return summary


def _receipt_stage(
    status: str,
    *,
    count: int | None = None,
    pending_count: int | None = None,
    error: str | None = None,
    evidence: list[str] | None = None,
    request_budget: dict[str, Any] | None = None,
) -> dict[str, Any]:
    row: dict[str, Any] = {"status": status, "evidence": evidence or []}
    if count is not None:
        row["count"] = count
    if pending_count is not None:
        row["pending_count"] = pending_count
    if error:
        row["error"] = error
    if request_budget is not None:
        row["request_budget"] = request_budget
    return row


def _target_ready_error_text(value: Any) -> str | None:
    if isinstance(value, dict):
        parts = [
            f"{kb_id}: {message}"
            for kb_id, message in value.items()
            if isinstance(message, str) and message.strip()
        ]
        return "; ".join(parts) or None
    return value if isinstance(value, str) and value.strip() else None


def _target_ready_evidence(value: Any) -> list[str]:
    if isinstance(value, dict):
        return [
            f"{kb_id}: {item}"
            for kb_id, items in value.items()
            if isinstance(items, list)
            for item in items
            if isinstance(item, str) and item.strip()
        ]
    if isinstance(value, list):
        return [item for item in value if isinstance(item, str) and item.strip()]
    return []


def _ack_stage(value: dict[str, Any]) -> dict[str, Any]:
    result = {key: int(value.get(key) or 0) for key in ("acknowledged", "pending", "failed")}
    result["status"] = (
        "failed" if value.get("failed") else "partial" if value.get("pending") else "succeeded"
    )
    failures = [
        f"{failure['url']}: {failure['error']}"
        for failure in value.get("failures") or []
        if isinstance(failure, dict)
        and isinstance(failure.get("url"), str)
        and failure["url"].strip()
        and isinstance(failure.get("error"), str)
        and failure["error"].strip()
    ]
    result["evidence"] = failures
    result["ready_status"] = value.get("ready_status") or "unknown"
    result["ready_evidence"] = _target_ready_evidence(value.get("ready_evidence"))
    if result["status"] == "failed":
        result["error"] = (
            "target item failure: " + "; ".join(failures)
            if failures
            else str(value.get("ready_error") or "target item acknowledgement failed")
        )
    if result["ready_status"] == "failed":
        result["ready_error"] = str(value.get("ready_error") or "native Ready/index build failed")
    return result


def _write_receipt(
    receipts_dir: Path,
    date_text: str,
    site_id: str,
    stages: dict[str, Any],
    evidence: list[str],
    reason: str | None,
    telemetry: dict[str, Any] | None = None,
) -> dict[str, Any]:
    flow = _load_flow_module()
    destination = receipts_dir / date_text / f"{site_id}.json"
    old = None
    if destination.is_file():
        old = flow.normalize_receipt(
            json.loads(destination.read_text(encoding="utf-8")), destination
        )
    merged = dict(old.get("stages", {}) if old else {})
    for stage, value in stages.items():
        if stage in flow.COUNT_STAGES:
            previous = merged.get(stage, {})
            count = value.get("count")
            old_count = previous.get("count")
            value = {
                **value,
                "count": (
                    (count + old_count)
                    if count is not None and old_count is not None
                    else count if count is not None else old_count
                ),
            }
            if "pending_count" not in value and "pending_count" in previous:
                value["pending_count"] = previous["pending_count"]
            if stage == "downloaded" and "request_budget" not in value:
                if "request_budget" in previous:
                    value["request_budget"] = previous["request_budget"]
        elif stage in flow.ACK_STAGES:
            previous = merged.get(stage, {})
            for name in ("acknowledged", "pending", "failed"):
                current, prior = value.get(name), previous.get(name)
                value[name] = (
                    (current + prior)
                    if current is not None and prior is not None
                    else current if current is not None else prior
                )
        merged[stage] = value
    merged_telemetry = (
        dict(telemetry or {})
        if old is None
        else _merge_attempt_telemetry(old.get("telemetry"), telemetry)
    )
    statuses = [row.get("status", "unknown") for row in merged.values()]
    failed = [
        value
        for value in merged.values()
        if value.get("status") in {"failed", "partial"}
        or value.get("ready_status") in {"failed", "partial"}
    ]
    ai_count = (merged.get("ai_candidates") or {}).get("count") or 0
    if ai_count and any(
        value.get("ready_status") == "not_required"
        for name, value in merged.items()
        if name in flow.ACK_STAGES
    ):
        failed.append({"status": "partial", "error": "native Ready was marked not_required"})
    run_status = (
        "failed"
        if failed and all(value.get("status") == "failed" for value in failed)
        else "partial" if failed or any(value == "unknown" for value in statuses) else "succeeded"
    )
    receipt_reason = reason if run_status != "succeeded" else None
    if run_status != "succeeded" and not (
        isinstance(receipt_reason, str) and receipt_reason.strip()
    ):
        receipt_reason = next(
            (
                candidate
                for value in merged.values()
                for candidate in (
                    value.get("error") if value.get("status") in {"failed", "partial"} else None,
                    (
                        value.get("ready_error")
                        if value.get("ready_status") in {"failed", "partial"}
                        else None
                    ),
                )
                if isinstance(candidate, str) and candidate.strip()
            ),
            (
                "native Ready was marked not_required"
                if any(
                    value.get("ready_status") == "not_required"
                    for name, value in merged.items()
                    if name in flow.ACK_STAGES
                )
                else "one or more stages reported failure or partial completion"
            ),
        )
    all_evidence = list(dict.fromkeys((old or {}).get("evidence", []) + evidence))
    record = {
        "schema_version": 2,
        "receipt_id": f"{date_text}/{site_id}",
        "date_local": date_text,
        "site_id": site_id,
        "run_status": run_status,
        "reason": receipt_reason,
        "evidence": all_evidence,
        "attempt_count": int((old or {}).get("attempt_count") or 0) + 1,
        "stages": merged,
    }
    if merged_telemetry:
        record["telemetry"] = merged_telemetry
    normalized = flow.normalize_receipt(record, destination)
    _atomic_json(destination, normalized)
    return normalized


def _merge_attempt_telemetry(
    previous: dict[str, Any] | None, current: dict[str, Any] | None
) -> dict[str, Any]:
    """Keep cumulative fields only when every attempt supplied that field."""
    if previous is None or current is None:
        return {}

    merged: dict[str, Any] = {}

    def merge_sum(
        target: dict[str, Any], old: dict[str, Any], new: dict[str, Any], name: str
    ) -> None:
        if name in old and name in new:
            target[name] = old[name] + new[name]

    def merge_ids(
        target: dict[str, Any], old: dict[str, Any], new: dict[str, Any], name: str
    ) -> None:
        if name in old and name in new:
            target[name] = sorted(set(old[name]) | set(new[name]))

    merge_sum(merged, previous, current, "elapsed_seconds")
    old_source = previous.get("source") or {}
    new_source = current.get("source") or {}
    if isinstance(old_source, dict) and isinstance(new_source, dict):
        source: dict[str, Any] = {}
        for name in ("request_count", "pages_visited", "links_discovered"):
            merge_sum(source, old_source, new_source, name)
        if "files_by_type" in old_source and "files_by_type" in new_source:
            source["files_by_type"] = {
                kind: old_source["files_by_type"][kind] + new_source["files_by_type"][kind]
                for kind in ("html", "pdf", "other")
            }
        for name in ("file_ids", "page_ids", "link_ids"):
            merge_ids(source, old_source, new_source, name)
        if source:
            merged["source"] = source

    for group, counters, identities in (
        ("database", (), ("new_version_ids", "duplicate_version_ids")),
        ("enrichment", ("succeeded_count",), ("succeeded_ids",)),
        ("conversion", ("succeeded_count", "failed_count"), ("succeeded_ids", "failed_ids")),
    ):
        old_group = previous.get(group) or {}
        new_group = current.get(group) or {}
        if isinstance(old_group, dict) and isinstance(new_group, dict):
            value: dict[str, Any] = {}
            for name in counters:
                merge_sum(value, old_group, new_group, name)
            for name in identities:
                merge_ids(value, old_group, new_group, name)
            if value:
                merged[group] = value

    merge_sum(merged, previous, current, "ai_first_candidate_count")
    merge_ids(merged, previous, current, "ai_first_candidate_ids")

    old_targets = previous.get("targets") or {}
    new_targets = current.get("targets") or {}
    if isinstance(old_targets, dict) and isinstance(new_targets, dict):
        targets: dict[str, Any] = {}
        for target in old_targets.keys() & new_targets.keys():
            old_target = old_targets[target]
            new_target = new_targets[target]
            if not isinstance(old_target, dict) or not isinstance(new_target, dict):
                continue
            value: dict[str, Any] = {}
            merge_ids(value, old_target, new_target, "acknowledged_ids")
            if "ready_status" in new_target:
                value["ready_status"] = new_target["ready_status"]
            if (
                "ready_ids" in old_target
                and "ready_ids" in new_target
                and "invalidated_ready_ids" in new_target
            ):
                ready_ids = set(old_target["ready_ids"])
                ready_ids.difference_update(new_target["invalidated_ready_ids"])
                ready_ids.update(new_target["ready_ids"])
                value["ready_ids"] = sorted(ready_ids)
            if (
                "invalidated_ready_ids" in old_target
                and "invalidated_ready_ids" in new_target
                and "ready_ids" in new_target
            ):
                invalidated = set(old_target["invalidated_ready_ids"])
                invalidated.update(new_target["invalidated_ready_ids"])
                invalidated.difference_update(new_target["ready_ids"])
                value["invalidated_ready_ids"] = sorted(invalidated)
            if value:
                targets[target] = value
        if targets:
            merged["targets"] = targets
    return merged


def run_monitor(
    site_id: str, config: dict[str, Any], state: dict[str, Any], crawler: _BudgetCrawler
) -> dict[str, Any]:
    site = SITES[site_id]
    crawler.begin_site(4)
    url = site.seeds[0]
    try:
        body, headers, final_url = crawler._request(url, timeout=40, delay_seconds=site.delay)
        success = bool(body) and bool(headers)
        return {
            "status": "succeeded" if success else "failed",
            "error": None if success else "monitor response was empty",
            "request_count": crawler.site_attempts,
            "page_id": (
                hashlib.sha256(_canonical_url(final_url).encode("utf-8")).hexdigest()
                if success
                else None
            ),
            "evidence": [
                f"requested {url}; response bytes={len(body)}; final_host={urlsplit(final_url).hostname}"
            ],
        }
    except Exception as exc:
        return {
            "status": "failed",
            "error": f"{type(exc).__name__}: {exc}",
            "request_count": crawler.site_attempts,
            "page_id": None,
            "evidence": [url],
        }


def _new_crawler(data_root: Path) -> _BudgetCrawler:
    return _BudgetCrawler(
        SimpleNamespace(data_root=data_root),
        str(data_root / "download-staging"),
        "AI-Actuarial-Local-Flow/1.0",
        default_delay_seconds=0.5,
    )


def _run_site(
    site_id: str,
    date_text: str,
    config: dict[str, Any],
    state: dict[str, Any],
    hub: DatabaseConfig,
    *,
    preview: bool = False,
    exact_url: str | None = None,
    article_limit: int = MAX_ENRICHMENTS_PER_SITE_RUN,
    request_limit: int = MAX_SITE_REQUESTS,
    source_page_url: str | None = None,
) -> dict[str, Any]:
    if site_id not in SITES:
        raise RuntimeErrorSafe(f"unknown monitored site: {site_id}")
    site = SITES[site_id]
    site_started = time.perf_counter()
    if site.action == "hold":
        raise RuntimeErrorSafe("held source cannot be executed")
    request_maximum = (
        MAX_SWISS_SITE_REQUESTS
        if site_id == "swiss-association-of-actuaries"
        else MAX_SITE_REQUESTS
    )
    if not 1 <= int(request_limit) <= request_maximum:
        raise RuntimeErrorSafe(
            f"--request-limit for {site_id} must be between 1 and {request_maximum}"
        )
    receipts_dir = Path(
        config.get("receipts_root") or (Path(config["state_root"]) / "receipts")
    ).resolve()
    evidence: list[str] = []
    if site.action == "monitor_only":
        monitor = run_monitor(
            site_id, config, state, _new_crawler(Path(config["data_root"]).resolve())
        )
        stage = {"monitor": monitor}
        telemetry = {
            "elapsed_seconds": time.perf_counter() - site_started,
            "source": {
                "request_count": monitor.get("request_count", 0),
                "pages_visited": int(monitor.get("page_id") is not None),
                "links_discovered": 0,
                "files_by_type": {"html": 0, "pdf": 0, "other": 0},
                "file_ids": [],
                "page_ids": [monitor["page_id"]] if monitor.get("page_id") else [],
                "link_ids": [],
            },
        }
        if preview:
            return {
                "site_id": site_id,
                "date_local": date_text,
                "preview": True,
                "stages": stage,
                "telemetry": telemetry,
            }
        return _write_receipt(
            receipts_dir,
            date_text,
            site_id,
            stage,
            monitor.get("evidence") or [],
            monitor.get("error"),
            telemetry,
        )
    site_state = state["sites"].setdefault(site_id, {})
    capture_state = site_state
    selected_url = None
    initial_frontier = None
    if exact_url:
        selected_url = _canonical_url(exact_url)
        if not _allowed_host(selected_url, site) or not _site_candidate_allowed(
            site, selected_url, Path(urlsplit(selected_url).path).name, {"operator_exact_url": True}
        ):
            raise RuntimeErrorSafe(
                "exact URL is outside this site's approved source and include/exclude rules"
            )
        replay_key = hashlib.sha256(selected_url.encode("utf-8")).hexdigest()
        replay_states = site_state.setdefault("exact_replays", {})
        capture_state = replay_states.setdefault(replay_key, {"url": selected_url})
        exact_site = Site(
            site.site_id, site.action, (selected_url,), site.hosts, site.delay, 0, site.page_hints
        )
        initial_frontier = [
            {
                "url": selected_url,
                "kind": "document",
                "depth": 0,
                "referrer": _canonical_url(source_page_url) if source_page_url else selected_url,
                "title": Path(urlsplit(selected_url).path).name,
                "discovery_metadata": {"operator_exact_url": True},
            }
        ]
    else:
        exact_site = site
    crawler = _new_crawler(Path(config["data_root"]).resolve())
    captured = _capture(
        exact_site,
        state,
        crawler,
        Path(config["data_root"]).resolve(),
        capture_state,
        initial_frontier=initial_frontier,
        request_limit=request_limit,
    )
    process_url = selected_url
    if selected_url and captured["downloads"]:
        # A validated same-host redirect is stored under the canonical final URL.
        # Process precisely that captured identity so exact replay cannot report a
        # successful download while leaving its article unprocessed.
        process_url = _canonical_url(str(captured["downloads"][0]["url"]))
    processed = _process_articles(
        site_id,
        state,
        config,
        hub,
        preview=preview,
        article_limit=article_limit,
        include_url=process_url,
    )
    source_errors = captured["errors"]
    downloaded = captured["downloads"]
    budget_pause = {
        "used": int(captured["request_count"]),
        "cap": int(captured.get("request_limit") or request_limit),
        "paused": bool(captured["budget_exhausted"]),
        "frontier_remaining": int(captured["frontier_remaining"]),
    }
    if source_errors:
        source_error = "; ".join(source_errors[:3])
        if captured["budget_exhausted"]:
            source_error += "; request budget reached with discovery cursor retained"
        source_status = "partial" if downloaded or processed["candidate_count"] else "failed"
    elif captured["budget_exhausted"]:
        source_error = "site reached its bounded HTTP request budget; discovery cursor was retained"
        source_status = "partial"
    elif captured["request_count"] == 0:
        source_error = "source made no verified HTTP requests"
        source_status = "failed"
    elif not downloaded and not processed["candidate_count"] and not processed["article_errors"]:
        source_error = (
            "bounded source discovery produced no verified article or document candidates"
        )
        source_status = "failed"
    else:
        source_error = None
        source_status = "succeeded"
    if captured["request_count"]:
        evidence.append(
            f"HTTP requests={captured['request_count']}/{budget_pause['cap']}; frontier_remaining={captured['frontier_remaining']}; downloaded={len(downloaded)}"
        )
    stages = {
        "downloaded": _receipt_stage(
            source_status,
            count=len(downloaded),
            error=(source_error if source_errors or not captured["budget_exhausted"] else None),
            evidence=evidence,
            request_budget=budget_pause,
        ),
        "new_versions": _receipt_stage(
            "partial" if processed["article_errors"] else "succeeded",
            count=processed["new_versions"],
            error=(
                "; ".join(processed["article_errors"][:3]) if processed["article_errors"] else None
            ),
            pending_count=processed.get("pending_article_count", 0),
        ),
        "duplicates": _receipt_stage("succeeded", count=processed["duplicates"]),
        "enriched": _receipt_stage(
            "partial" if processed["article_errors"] else "succeeded",
            count=processed["enriched"],
            error=(
                "; ".join(processed["article_errors"][:3]) if processed["article_errors"] else None
            ),
            pending_count=processed.get("pending_article_count", 0),
        ),
        "ai_candidates": _receipt_stage(
            "partial" if processed["article_errors"] else "succeeded",
            count=processed["ai_candidates"],
            error=(
                "; ".join(processed["article_errors"][:3]) if processed["article_errors"] else None
            ),
        ),
        "china_ack": _ack_stage(processed["target"]["china"]),
        "overseas_ack": _ack_stage(processed["target"]["overseas"]),
    }
    telemetry = {
        "elapsed_seconds": time.perf_counter() - site_started,
        "source": captured.get("telemetry") or {"request_count": captured["request_count"]},
        "targets": {
            target: {
                **(
                    {"acknowledged_ids": processed["target"][target]["acknowledged_ids"]}
                    if "acknowledged_ids" in processed["target"][target]
                    else {}
                ),
                **(
                    {"ready_ids": processed["target"][target]["ready_ids"]}
                    if "ready_ids" in processed["target"][target]
                    else {}
                ),
                **(
                    {"invalidated_ready_ids": processed["target"][target]["invalidated_ready_ids"]}
                    if "invalidated_ready_ids" in processed["target"][target]
                    else {}
                ),
                "ready_status": processed["target"][target].get("ready_status", "unknown"),
            }
            for target in ("china", "overseas")
        },
    }
    if "database" in processed:
        telemetry["database"] = processed["database"]
    if "enrichment" in processed:
        telemetry["enrichment"] = processed["enrichment"]
    if "conversion" in processed:
        telemetry["conversion"] = processed["conversion"]
    if "ai_first_candidate_ids" in processed:
        telemetry["ai_first_candidate_count"] = processed["ai_candidates"]
        telemetry["ai_first_candidate_ids"] = processed["ai_first_candidate_ids"]
    if preview:
        evidence.append(
            "preview only: no Hub, model, target or receipt writes were performed; local CAS and restart state may be updated"
        )
        # Preserve the plan for callers without claiming a production receipt.
        return {
            "date_local": date_text,
            "site_id": site_id,
            "preview": True,
            "stages": stages,
            "target_candidates": processed["ai_candidates"],
            "http_requests": captured["request_count"],
            "request_limit": budget_pause["cap"],
            "pending_article_count": processed.get("pending_article_count", 0),
            "telemetry": telemetry,
        }
    return _write_receipt(
        receipts_dir, date_text, site_id, stages, evidence, source_error, telemetry
    )


def _restore_hub_blob(version: dict[str, Any], raw: bytes, hub: DatabaseConfig) -> None:
    source_site = version.get("source_site")
    if not isinstance(source_site, str) or not re.fullmatch(
        r"[a-z0-9][a-z0-9._-]{0,127}", source_site
    ):
        raise RuntimeErrorSafe(
            "historical source bytes cannot be restored without known source-site provenance"
        )
    restored = ingest_live_article(
        canonical_url=str(version["url"]),
        source_site=source_site,
        source_page_url=str(version.get("source_page_url") or version["url"]),
        title=str(version["bilingual_metadata"]["en"]["title"]),
        raw_bytes=raw,
        markdown=str(version["markdown"]),
        content_type=str(version.get("content_type") or "application/octet-stream"),
        data_root=Path(version.get("local_data_root") or "."),
        database_url=hub,
        metadata={
            "recovered_historical_version_id": int(version["content_version_id"]),
            "recovered_from_hub_ai_backlog": True,
        },
    )
    if int(restored["content_version_id"]) != int(version["content_version_id"]):
        raise RuntimeErrorSafe(
            "verified source bytes did not restore the expected Hub URL/raw-SHA version"
        )
    current = content_hub_read_version(str(version["url"]), str(version["raw_sha256"]), hub)
    if not current or int(current["content_version_id"]) != int(version["content_version_id"]):
        raise RuntimeErrorSafe("restored Hub version failed exact URL/raw-SHA readback")


def _refresh_pending_target_item(
    item: dict[str, Any], hub: DatabaseConfig
) -> tuple[dict[str, Any] | None, str | None]:
    """Re-read durable retries so old any-category-AI snapshots cannot publish."""
    current = content_hub_read_version(
        str(item.get("url") or ""), str(item.get("raw_sha256") or ""), hub
    )
    if not current or current.get("excluded_reason"):
        return None, "exact Hub version is missing, excluded, or no longer enriched"
    if int(current.get("content_version_id") or 0) != int(item.get("content_version_id") or 0):
        return None, "exact Hub version ID changed"
    if not _valid_existing_enrichment(current):
        return None, "exact Hub enrichment is no longer valid"
    annotation = current["annotation"]
    categories = annotation.get("categories") or []
    keys = [entry.get("key") for entry in categories if isinstance(entry, dict)]
    if not keys or keys[0] != "AI" or annotation.get("source_primary_category") != "AI":
        primary = keys[0] if keys else "unknown"
        if primary == "AI":
            return None, (
                f"Hub source primary category {annotation.get('source_primary_category')!r} "
                "conflicts with categories[0]"
            )
        return None, f"current Hub primary category is {primary!r}"
    refreshed = {
        **item,
        "url": current["url"],
        "raw_sha256": current["raw_sha256"],
        "content_version_id": current["content_version_id"],
        "markdown_text_variant_id": current["markdown_text_variant_id"],
        "markdown_sha256": current["markdown_sha256"],
        "hub_markdown_sha256": current["hub_markdown_sha256"],
        "hub_markdown_sha256_kind": current["hub_markdown_sha256_kind"],
        "markdown": current["markdown"],
        "bilingual_metadata": current["bilingual_metadata"],
        "categories": categories,
        "source_primary_category": annotation["source_primary_category"],
        "date_facts": annotation.get("date_facts") or {},
        "date_review_status": annotation.get("date_review_status"),
        "date_review_reason": annotation.get("date_review_reason"),
        "date_review_unresolved": annotation.get("date_review_unresolved") or [],
        "source_signature": annotation.get("source_signature"),
        "result_sha256": annotation.get("result_sha256"),
        "legacy_hub_annotation": True,
        "blob_availability": "present" if current.get("blob_available") else "missing",
    }
    return refreshed, None


def _suppress_non_primary_retry(state: dict[str, Any], item: dict[str, Any], reason: str) -> None:
    version_id = str(item.get("content_version_id") or "")
    if version_id:
        state.setdefault("pending_target_reconciliation", {}).pop(version_id, None)
        state.setdefault("pending_source_recovery", {}).pop(version_id, None)
        state.setdefault("ineligible_target_reconciliation", {})[version_id] = {
            "url": item.get("url"),
            "raw_sha256": item.get("raw_sha256"),
            "reason": reason,
            "checked_at": dt.datetime.now(dt.timezone.utc).isoformat(),
        }
        for target in ("china", "overseas"):
            state.setdefault("target_acks", {}).pop(_target_ack_key(target, item), None)


def _recover_one_source(
    state: dict[str, Any], data_root: Path, hub: DatabaseConfig
) -> dict[str, Any] | None:
    pending = state.setdefault("pending_source_recovery", {})
    if not pending:
        return None
    key = sorted(pending, key=lambda value: int(value))[0]
    version = pending[key]
    try:
        refreshed, ineligible_reason = _refresh_pending_target_item(version, hub)
        if refreshed is None:
            del pending[key]
            _suppress_non_primary_retry(state, version, ineligible_reason or "not AI-first")
            _atomic_json(_state_path_for_data(state), state)
            return None
        version = refreshed
        url = str(version["url"])
        crawler = _new_crawler(data_root)
        crawler.begin_site(1)
        staging, headers, final_url, digest, size = crawler._download_file(
            url, data_root / "recovery-staging", delay_seconds=0.5
        )
        raw = staging.read_bytes()
        staging.unlink(missing_ok=True)
        if _canonical_url(final_url) != _canonical_url(url):
            raise RuntimeErrorSafe(
                "source recovery redirected away from the Hub's exact URL identity"
            )
        if (
            digest != version["raw_sha256"]
            or hashlib.sha256(raw).hexdigest() != version["raw_sha256"]
        ):
            raise RuntimeErrorSafe(
                "current source bytes do not match the missing historical Hub SHA; version remains pending"
            )
        artifact = _persist_cas(data_root, raw, digest)
        _restore_hub_blob({**version, "local_data_root": str(data_root)}, raw, hub)
        item = {
            **version,
            "artifact_path": str(artifact),
            "byte_size": size,
            "content_type": headers.get("content-type") or version.get("content_type"),
        }
        state.setdefault("pending_target_reconciliation", {})[key] = item
        del pending[key]
        return item
    except Exception as exc:
        version["recovery_attempts"] = int(version.get("recovery_attempts") or 0) + 1
        version["recovery_error"] = f"{type(exc).__name__}: {exc}"[:500]
        return None


def reconcile_backlog(
    config: dict[str, Any], state: dict[str, Any], hub: DatabaseConfig, *, limit: int = 10
) -> dict[str, Any]:
    limit = min(MAX_RECONCILE_PAGE - 1, max(1, int(limit)))
    cursor = int(state.get("backlog_cursor") or 0)
    data_root = Path(config["data_root"]).resolve()
    recovered = _recover_one_source(state, data_root, hub)
    pending_targets = state.setdefault("pending_target_reconciliation", {})
    # Reserve room for both a fresh Hub page and a source recovery. The remote
    # worker rejects batches above 100 and must never receive an overflow page.
    pending_cap = MAX_RECONCILE_PAGE - 2
    pending_items = [
        pending_targets[key]
        for key in sorted(pending_targets, key=lambda value: int(value))[:pending_cap]
    ]
    refreshed_pending_items: list[dict[str, Any]] = []
    for item in pending_items:
        refreshed, ineligible_reason = _refresh_pending_target_item(item, hub)
        if refreshed is None:
            _suppress_non_primary_retry(
                state, item, ineligible_reason or "current Hub version is not AI-first"
            )
            continue
        pending_targets[str(refreshed["content_version_id"])] = refreshed
        refreshed_pending_items.append(refreshed)
    pending_items = refreshed_pending_items
    page_limit = min(limit, MAX_RECONCILE_PAGE - len(pending_items) - int(recovered is not None))
    page_limit = max(1, page_limit)
    page = content_hub_list_ai_versions(after_version_id=cursor, limit=page_limit, database_url=hub)
    page_versions = page.get("versions") or []
    fresh_version_ids = {int(row["content_version_id"]) for row in page_versions}
    candidates = {int(row["content_version_id"]): row for row in page_versions}
    if recovered:
        candidates[int(recovered["content_version_id"])] = recovered
        fresh_version_ids.discard(int(recovered["content_version_id"]))
    for item in pending_items:
        candidates[int(item["content_version_id"])] = item
        fresh_version_ids.discard(int(item["content_version_id"]))
    if not candidates:
        _atomic_json(_state_path_for_data(state), state)
        return {
            "processed": 0,
            "cursor": cursor,
            "has_more": page.get("has_more"),
            "eligible_ai_count": page.get("eligible_ai_count"),
            "remaining_estimate": max(
                0,
                int(page.get("eligible_ai_count") or 0)
                - len(state.get("reconciled_version_ids") or []),
            ),
            "missing_blob_ai_count": page.get("missing_blob_ai_count"),
            "pending_source_recovery": len(state.get("pending_source_recovery") or {}),
            "pending_target_reconciliation": len(pending_targets),
            "deferred_ambiguous_urls": len(state.get("deferred_ambiguous") or {}),
        }
    candidate_versions = list(candidates.values())
    groups = content_hub_ai_url_versions(
        list(dict.fromkeys(str(row["url"]) for row in candidate_versions)), database_url=hub
    )
    state.setdefault("deferred_ambiguous", {})
    ambiguous_groups = {url: ids for url, ids in groups.items() if len(ids) > 1}
    state["deferred_ambiguous"].update(ambiguous_groups)
    ambiguous_ids = {version_id for ids in ambiguous_groups.values() for version_id in ids}
    state["deferred_ambiguous_ids"] = sorted(
        set(state.get("deferred_ambiguous_ids", [])) | ambiguous_ids
    )
    eligible_versions = [
        row for row in candidate_versions if int(row["content_version_id"]) not in ambiguous_ids
    ]

    pending_by_target: dict[str, list[dict[str, Any]]] = {"china": [], "overseas": []}
    keys_by_version: dict[int, dict[str, Any]] = {}
    newly_missing: list[dict[str, Any]] = []
    for version in eligible_versions:
        version_id = int(version["content_version_id"])
        if version_id in fresh_version_ids:
            refreshed, ineligible_reason = _refresh_pending_target_item(version, hub)
            if refreshed is None:
                _suppress_non_primary_retry(
                    state,
                    version,
                    ineligible_reason or "current Hub version is no longer AI-first",
                )
                continue
            version = refreshed
        artifact = Path(
            version.get("artifact_path") or _local_cas_path(data_root, str(version["raw_sha256"]))
        )
        if not artifact.is_file():
            recovery_key = str(version_id)
            state.setdefault("pending_source_recovery", {}).setdefault(recovery_key, version)
            newly_missing.append(version)
            continue
        raw = artifact.read_bytes()
        if hashlib.sha256(raw).hexdigest() != version["raw_sha256"]:
            raise RuntimeErrorSafe("AI backlog CAS object failed its exact SHA binding")
        if version.get("blob_availability") != "present":
            _restore_hub_blob({**version, "local_data_root": str(data_root)}, raw, hub)
        item = {
            **version,
            "artifact_path": str(artifact),
            "title": version.get("title") or version["bilingual_metadata"]["en"]["title"],
            "source_page_url": version.get("source_page_url") or version.get("url"),
            "content_type": version.get("content_type") or "application/octet-stream",
        }
        keys_by_version[version_id] = item
        state.setdefault("pending_target_reconciliation", {})[str(version_id)] = item
        for target in pending_by_target:
            if _target_needs_item(state, target, item):
                pending_by_target[target].append(item)

    ready_errors = {}
    for target, items in pending_by_target.items():
        pending_ready = any(
            bool(state["target_acks"].get(_target_ack_key(target, item), {}).get("ready_pending"))
            for item in keys_by_version.values()
        )
        if not items and not pending_ready:
            continue
        for item in items:
            key = _target_ack_key(target, item)
            state["target_acks"].setdefault(
                key,
                {
                    "status": "pending",
                    "metadata_fingerprint": _target_metadata_fingerprint(item),
                    "ready_pending": True,
                },
            )
        _atomic_json(_state_path_for_data(state), state)
        try:
            reply = _stream_target(config, target, items, pending_ready)
            by_version = {
                (row.get("content_version_id"), row.get("raw_sha256")): row
                for row in reply.get("acknowledgements") or []
            }
            for item in keys_by_version.values():
                ack = by_version.get((item["content_version_id"], item["raw_sha256"]))
                key = _target_ack_key(target, item)
                if (
                    ack
                    and ack.get("status") in {"written", "already_current"}
                    and ack.get("metadata_fingerprint") == _target_metadata_fingerprint(item)
                ):
                    state["target_acks"][key] = {
                        "status": "acknowledged",
                        "metadata_fingerprint": ack["metadata_fingerprint"],
                        "ready_pending": reply.get("ready_status") != "succeeded",
                        "ack": ack,
                    }
                elif reply.get("ready_status") == "succeeded":
                    # A previous run may have committed the item's row but
                    # failed while building the complete native index/Ready.
                    # The ready-only retry has no per-item acknowledgement;
                    # clear its durable pending flag after the full readback.
                    prior = state["target_acks"].get(key)
                    if (
                        isinstance(prior, dict)
                        and prior.get("status") == "acknowledged"
                        and prior.get("metadata_fingerprint") == _target_metadata_fingerprint(item)
                    ):
                        prior["ready_pending"] = False
            if reply.get("ready_status") != "succeeded":
                ready_errors[target] = str(
                    reply.get("ready_error") or "target Ready remains pending"
                )
        except Exception as exc:
            ready_errors[target] = f"{type(exc).__name__}: {exc}"
        _atomic_json(_state_path_for_data(state), state)
    for item in keys_by_version.values():
        if all(
            state["target_acks"].get(_target_ack_key(target, item), {}).get("status")
            == "acknowledged"
            and not state["target_acks"].get(_target_ack_key(target, item), {}).get("ready_pending")
            for target in ("china", "overseas")
        ):
            state["pending_target_reconciliation"].pop(str(item["content_version_id"]), None)
    if ready_errors:
        _atomic_json(_state_path_for_data(state), state)
        return {
            "processed": 0,
            "cursor": cursor,
            "has_more": True,
            "ready_errors": ready_errors,
            "eligible_ai_count": page.get("eligible_ai_count"),
            "missing_blob_ai_count": page.get("missing_blob_ai_count"),
            "pending_source_recovery": len(state.get("pending_source_recovery") or {}),
            "pending_target_reconciliation": len(state["pending_target_reconciliation"]),
        }
    if page_versions:
        state["backlog_cursor"] = int(page.get("next_after_version_id") or cursor)
        state.setdefault("reconciled_version_ids", []).extend(
            int(item["content_version_id"])
            for item in eligible_versions
            if int(item["content_version_id"]) in keys_by_version
        )
        state["reconciled_version_ids"] = sorted(set(state["reconciled_version_ids"]))
    _atomic_json(_state_path_for_data(state), state)
    return {
        "processed": len(keys_by_version),
        "cursor": int(state.get("backlog_cursor") or cursor),
        "has_more": page.get("has_more"),
        "eligible_ai_count": page.get("eligible_ai_count"),
        "missing_blob_ai_count": page.get("missing_blob_ai_count"),
        "remaining_estimate": max(
            0,
            int(page.get("eligible_ai_count") or 0)
            - len(state.get("reconciled_version_ids") or []),
        ),
        "pending_source_recovery": len(state.get("pending_source_recovery") or {}),
        "pending_target_reconciliation": len(state["pending_target_reconciliation"]),
        "source_recovery_failures": sum(
            1
            for row in (state.get("pending_source_recovery") or {}).values()
            if row.get("recovery_error")
        ),
        "deferred_ambiguous_urls": len(state.get("deferred_ambiguous") or {}),
        "deferred_ambiguous_version_count": len(state.get("deferred_ambiguous_ids") or []),
    }


def _load_receipts(receipts: Path) -> dict[tuple[str, str], dict[str, Any]]:
    return _load_flow_module().read_receipts(receipts)


def preflight(config: dict[str, Any], *, remote: bool = False) -> dict[str, Any]:
    report: dict[str, Any] = {
        "schema_version": 1,
        "read_only": True,
        "checks": {},
        "production_writes": False,
    }
    state_root, data_root = (
        Path(config["state_root"]).expanduser().resolve(),
        Path(config["data_root"]).expanduser().resolve(),
    )

    def writable_location(path: Path) -> bool:
        candidate = path
        while not candidate.exists() and candidate != candidate.parent:
            candidate = candidate.parent
        return candidate.is_dir() and os.access(candidate, os.W_OK)

    report["checks"]["state_root_writable"] = writable_location(state_root)
    report["checks"]["data_root_writable"] = writable_location(data_root)
    report["checks"]["site_registry_count"] = len(SITES)
    report["checks"]["planned_collect_count"] = sum(
        site.action == "collect" for site in SITES.values()
    )
    report["checks"]["monitor_only_count"] = sum(
        site.action == "monitor_only" for site in SITES.values()
    )
    report["checks"]["hold_count"] = sum(site.action == "hold" for site in SITES.values())
    hub = _connection_url(Path(config["hub_connection_file"]).expanduser())
    report["hub_schema"] = content_hub_schema_plan(hub)
    report["checks"]["hub_schema_current"] = not bool(report["hub_schema"].get("pending"))
    if report["checks"]["hub_schema_current"]:
        report["hub_status"] = content_hub_status(hub)
        backlog = content_hub_list_ai_versions(after_version_id=0, limit=1, database_url=hub)
        report["ai_backlog"] = {
            "eligible_versions": backlog.get("eligible_ai_count"),
            "missing_original_versions": backlog.get("missing_blob_ai_count"),
            "first_page_count": len(backlog.get("versions") or []),
            "bounded_page_limit": MAX_RECONCILE_PAGE,
        }
    else:
        report["hub_status"] = {
            "available": False,
            "reason": "apply the pending Hub migration before live operations",
        }
        report["ai_backlog"] = {
            "available": False,
            "reason": "requires the live-capture Hub schema",
        }
    report["targets"] = {}
    if remote:
        for target in ("china", "overseas"):
            target_result = _stream_target(config, target, [], False, preview=True)
            report["targets"][target] = target_result
            capacity = target_result.get("capacity") or {}
            report["checks"][f"{target}_target_capacity_clear"] = not bool(
                capacity.get("blocked", True)
            )
            kb_ids = set((target_result.get("knowledge_bases") or {}).keys())
            expected_ids = {"ai", "all"} if target == "china" else {"ai"}
            report["checks"][f"{target}_target_profiles_valid"] = kb_ids == expected_ids
    report["checks"]["codex_cli_found"] = bool(
        config.get("codex_executable") or shutil.which("codex.exe") or shutil.which("codex")
    )
    report["checks"]["all_32_registry_actions_match_offline_contract"] = set(SITES) == set(
        _load_flow_module().SITES
    )
    return report


def _daily_run(
    date_text: str, config: dict[str, Any], *, preview: bool = False, limit_reconcile: int = 5
) -> dict[str, Any]:
    run_started = time.perf_counter()
    date_value = dt.date.fromisoformat(date_text)
    receipts_dir = Path(
        config.get("receipts_root") or (Path(config["state_root"]) / "receipts")
    ).resolve()
    receipts = _load_receipts(receipts_dir)
    plan = _load_flow_module().make_plan(date_value, receipts)
    state_path = _state_path(config)
    global _ACTIVE_STATE_PATH
    _ACTIVE_STATE_PATH = state_path
    with _exclusive_lock(Path(config["state_root"]).resolve() / "runtime.lock"):
        state = _load_state(state_path)
        hub = _connection_url(Path(config["hub_connection_file"]).expanduser())
        results = []
        run_count = 0
        for plan_item in plan["items"]:
            site = plan_item["site_id"]
            if plan_item["action"] == "hold":
                continue
            if run_count >= MAX_DAILY_SITE_RUNS:
                break
            if preview:
                results.append({"site_id": site, "action": plan_item["action"], "preview": True})
            else:
                site_started = time.perf_counter()
                try:
                    results.append(_run_site(site, date_text, config, state, hub))
                except Exception as exc:
                    results.append(
                        {
                            "site_id": site,
                            "error": f"{type(exc).__name__}: {exc}",
                            "elapsed_seconds": time.perf_counter() - site_started,
                        }
                    )
            run_count += 1
            _atomic_json(state_path, state)
        backlog = None
        if not preview and limit_reconcile > 0:
            try:
                backlog = reconcile_backlog(config, state, hub, limit=limit_reconcile)
            except Exception as exc:
                backlog = {"error": f"{type(exc).__name__}: {exc}", "remaining_estimate": None}
        return {
            "schema_version": 1,
            "date_local": date_text,
            "plan": plan,
            "runs": results,
            "backlog": backlog,
            "receipts_dir": str(receipts_dir),
            "run_count": run_count,
            "preview": preview,
            "elapsed_seconds": time.perf_counter() - run_started,
        }


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Bounded local collection, Content Hub enrichment and native two-target publication."
    )
    sub = parser.add_subparsers(dest="command", required=True)
    pre = sub.add_parser(
        "preflight", help="read-only local, Hub and optional target contract checks"
    )
    pre.add_argument(
        "--config", type=Path, default=Path("reports/local-automation-state/config.local.json")
    )
    pre.add_argument(
        "--remote", action="store_true", help="also run read-only target container/profile probes"
    )
    pre.add_argument(
        "--json", action="store_true", help="emit machine-readable JSON (the default output format)"
    )
    migrate = sub.add_parser(
        "apply-hub-schema",
        help="apply the additive live-capture Hub migration after approved review",
    )
    migrate.add_argument(
        "--config", type=Path, default=Path("reports/local-automation-state/config.local.json")
    )
    migrate.add_argument(
        "--json", action="store_true", help="emit machine-readable JSON (the default output format)"
    )
    for command in ("plan", "run-daily"):
        item = sub.add_parser(command, help=f"{command} the PR #427 six-site daily schedule")
        item.add_argument("--date", default=dt.datetime.now(TIME_ZONE).date().isoformat())
        item.add_argument("--receipts", type=Path, default=None)
        item.add_argument(
            "--config", type=Path, default=Path("reports/local-automation-state/config.local.json")
        )
        item.add_argument("--preview", action="store_true")
        item.add_argument("--reconcile-limit", type=int, default=5)
        item.add_argument(
            "--json",
            action="store_true",
            help="emit machine-readable JSON (the default output format)",
        )
    site = sub.add_parser(
        "run-site", help="run one bounded site discovery and its Hub/AI/target stages"
    )
    site.add_argument("--site", required=True, choices=tuple(SITES))
    site.add_argument("--date", default=dt.datetime.now(TIME_ZONE).date().isoformat())
    site.add_argument(
        "--config", type=Path, default=Path("reports/local-automation-state/config.local.json")
    )
    site.add_argument(
        "--preview",
        action="store_true",
        help="fetch and plan without writing Hub, model, target or receipt data",
    )
    site.add_argument(
        "--exact-url", help="capture one validated official-site URL, useful for low-volume replay"
    )
    site.add_argument(
        "--source-page-url", help="official listing/detail page that linked the exact URL"
    )
    site.add_argument(
        "--article-limit",
        type=int,
        default=MAX_ENRICHMENTS_PER_SITE_RUN,
        help="maximum pending articles to enrich/publish in this site run (1-8)",
    )
    site.add_argument(
        "--request-limit",
        type=int,
        default=MAX_SITE_REQUESTS,
        help=("HTTP request cap for this run (1-30; Switzerland may explicitly use 1-60)"),
    )
    site.add_argument(
        "--json", action="store_true", help="emit machine-readable JSON (the default output format)"
    )
    backlog = sub.add_parser(
        "reconcile-targets", help="resume bounded AI-primary target acknowledgements"
    )
    backlog.add_argument("--limit", type=int, default=10)
    backlog.add_argument(
        "--config", type=Path, default=Path("reports/local-automation-state/config.local.json")
    )
    backlog.add_argument(
        "--json", action="store_true", help="emit machine-readable JSON (the default output format)"
    )
    cleanup = sub.add_parser(
        "cleanup-targets",
        help="audit or retire exact Hub-verified non-AI-primary files on one serving target",
    )
    cleanup.add_argument("--manifest", type=Path, required=True)
    cleanup.add_argument("--target", choices=("china", "overseas"), required=True)
    cleanup.add_argument(
        "--config", type=Path, default=Path("reports/local-automation-state/config.local.json")
    )
    cleanup.add_argument(
        "--apply", action="store_true", help="apply exact file retirement after the default dry-run"
    )
    cleanup.add_argument(
        "--json", action="store_true", help="emit machine-readable JSON (the default output format)"
    )
    backfill = sub.add_parser(
        "backfill-target-metadata",
        help="audit or backfill exact existing AI-primary Hub metadata and article dates",
    )
    backfill.add_argument("--manifest", type=Path, required=True)
    backfill.add_argument("--target", choices=("china", "overseas"), required=True)
    backfill.add_argument(
        "--config", type=Path, default=Path("reports/local-automation-state/config.local.json")
    )
    backfill.add_argument(
        "--apply", action="store_true", help="apply metadata-only updates after the default dry-run"
    )
    backfill.add_argument(
        "--json", action="store_true", help="emit machine-readable JSON (the default output format)"
    )
    return parser


def main(argv: list[str] | None = None) -> int:
    global _ACTIVE_STATE_PATH
    args = build_parser().parse_args(argv)
    try:
        if args.command == "preflight":
            config = _load_config(args.config.resolve())
            output = preflight(config, remote=args.remote)
        elif args.command == "apply-hub-schema":
            config = _load_config(args.config.resolve())
            _ACTIVE_STATE_PATH = _state_path(config)
            with _exclusive_lock(Path(config["state_root"]).resolve() / "runtime.lock"):
                hub = _connection_url(Path(config["hub_connection_file"]).expanduser())
                output = apply_content_hub_schema(hub)
        elif args.command == "plan":
            config = _load_config(args.config.resolve())
            receipts = Path(
                args.receipts
                or config.get("receipts_root")
                or (Path(config["state_root"]) / "receipts")
            ).resolve()
            output = _load_flow_module().make_plan(
                dt.date.fromisoformat(args.date), _load_receipts(receipts)
            )
        elif args.command == "run-daily":
            config = _load_config(args.config.resolve())
            output = _daily_run(
                args.date, config, preview=args.preview, limit_reconcile=args.reconcile_limit
            )
        elif args.command == "run-site":
            if not 1 <= args.article_limit <= MAX_ENRICHMENTS_PER_SITE_RUN:
                raise RuntimeErrorSafe(
                    f"--article-limit must be between 1 and {MAX_ENRICHMENTS_PER_SITE_RUN}"
                )
            maximum_requests = (
                MAX_SWISS_SITE_REQUESTS
                if args.site == "swiss-association-of-actuaries"
                else MAX_SITE_REQUESTS
            )
            if not 1 <= args.request_limit <= maximum_requests:
                raise RuntimeErrorSafe(
                    f"--request-limit for {args.site} must be between 1 and {maximum_requests}"
                )
            config = _load_config(args.config.resolve())
            _ACTIVE_STATE_PATH = _state_path(config)
            with _exclusive_lock(Path(config["state_root"]).resolve() / "runtime.lock"):
                state = _load_state(_ACTIVE_STATE_PATH)
                hub = _connection_url(Path(config["hub_connection_file"]).expanduser())
                output = _run_site(
                    args.site,
                    args.date,
                    config,
                    state,
                    hub,
                    preview=args.preview,
                    exact_url=args.exact_url,
                    article_limit=args.article_limit,
                    request_limit=args.request_limit,
                    source_page_url=args.source_page_url,
                )
                if not args.preview:
                    _atomic_json(_ACTIVE_STATE_PATH, state)
        elif args.command == "cleanup-targets":
            config = _load_config(args.config.resolve())
            manifest = json.loads(args.manifest.resolve().read_text(encoding="utf-8"))
            if (
                not isinstance(manifest, dict)
                or manifest.get("schema_version") != 1
                or not isinstance(manifest.get("records"), list)
            ):
                raise RuntimeErrorSafe("cleanup manifest must use schema_version 1 with records")
            output = _stream_target_cleanup(
                config,
                args.target,
                manifest["records"],
                dry_run=not args.apply,
            )
        elif args.command == "backfill-target-metadata":
            config = _load_config(args.config.resolve())
            manifest = json.loads(args.manifest.resolve().read_text(encoding="utf-8"))
            if (
                not isinstance(manifest, dict)
                or manifest.get("schema_version") != 1
                or not isinstance(manifest.get("records"), list)
            ):
                raise RuntimeErrorSafe(
                    "metadata backfill manifest must use schema_version 1 with records"
                )
            output = _stream_target_metadata_backfill(
                config,
                args.target,
                manifest["records"],
                dry_run=not args.apply,
            )
        else:
            config = _load_config(args.config.resolve())
            _ACTIVE_STATE_PATH = _state_path(config)
            with _exclusive_lock(Path(config["state_root"]).resolve() / "runtime.lock"):
                state = _load_state(_ACTIVE_STATE_PATH)
                output = reconcile_backlog(
                    config,
                    state,
                    _connection_url(Path(config["hub_connection_file"]).expanduser()),
                    limit=args.limit,
                )
        print(json.dumps(output, ensure_ascii=False, sort_keys=True, indent=2))
        if args.command == "preflight":
            required = [
                key
                for key in output["checks"]
                if key.endswith(
                    ("_writable", "_found", "_current", "_capacity_clear", "_profiles_valid")
                )
                or key.startswith("all_32")
            ]
            return 0 if required and all(output["checks"][key] for key in required) else 2
        if args.command == "run-site":
            if args.preview:
                statuses = {
                    stage.get("status")
                    for stage in (output.get("stages") or {}).values()
                    if isinstance(stage, dict)
                }
                if "failed" in statuses:
                    return 2
                if "partial" in statuses:
                    return 3
                return 0
            return 0 if output.get("run_status") == "succeeded" else 2
        if args.command == "run-daily":
            failed = any(
                row.get("error") or row.get("run_status") not in (None, "succeeded")
                for row in output.get("runs", [])
            )
            backlog = output.get("backlog") or {}
            failed = failed or bool(backlog.get("error") or backlog.get("ready_errors"))
            return 2 if failed else 0
        if args.command == "cleanup-targets":
            return 2 if output.get("failures") or output.get("ready_status") == "failed" else 0
        if args.command == "backfill-target-metadata":
            return 2 if output.get("failures") else 0
        if args.command == "reconcile-targets":
            return 2 if output.get("ready_errors") or output.get("error") else 0
        return 0
    except Exception as exc:
        print(
            json.dumps(
                {"schema_version": 1, "error": f"{type(exc).__name__}: {exc}"}, ensure_ascii=False
            )
        )
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
