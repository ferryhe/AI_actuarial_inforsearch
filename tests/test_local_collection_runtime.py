from __future__ import annotations

import base64
import hashlib
import json
import shlex
import subprocess
import sys
from pathlib import Path
from typing import Any

import pytest

import ai_actuarial.capacity as capacity
import ai_actuarial.content_hub as content_hub
import ai_actuarial.local_collection_runtime as runtime


def test_registry_and_audited_source_filters_cover_all_32_sites() -> None:
    assert len(runtime.SITES) == 32
    assert sum(site.action == "collect" for site in runtime.SITES.values()) == 29
    assert sum(site.action == "monitor_only" for site in runtime.SITES.values()) == 2
    assert sum(site.action == "hold" for site in runtime.SITES.values()) == 1
    assert runtime.SITES["iaa"].seeds[0] == "https://actuaries.org/publications/papers/"
    aict = runtime.SITES["aict"]
    assert not runtime._allowed_path("https://www.airc.org.tw/re/pdf/old-rule.pdf", aict)
    assert not runtime._site_candidate_allowed(
        aict, "https://www.airc.org.tw/memberconf/meeting.pdf"
    )
    assert runtime._allowed_path("https://www.airc.org.tw/rule/101", aict)


def test_iaa_filter_excludes_organization_chart_but_keeps_ai_governance_research() -> None:
    iaa = runtime.SITES["iaa"]
    chart_url = "https://actuaries.org/app/uploads/2026/04/IAA_Org_Chart_2026.pdf"
    chart_title = "Organizational Chart"
    assert not runtime._site_candidate_allowed(iaa, chart_url, chart_title, {})
    assert not runtime._site_document_relevant(
        iaa, chart_url, chart_title, "https://actuaries.org/publications/papers/", {}
    )

    governance_url = "https://actuaries.org/paper/artificial-intelligence-governance-framework/"
    governance_title = "Artificial Intelligence Governance Framework"
    assert runtime._site_candidate_allowed(iaa, governance_url, governance_title, {})
    assert runtime._site_document_relevant(
        iaa, governance_url, governance_title, "https://actuaries.org/publications/papers/", {}
    )


def test_iaa_observed_administrative_pdfs_are_filtered_but_research_is_kept() -> None:
    iaa = runtime.SITES["iaa"]
    administrative_urls = (
        "https://actuaries.org/app/uploads/2025/04/Statutes.pdf",
        "https://actuaries.org/app/uploads/2025/04/Statutes_FR.pdf",
        "https://actuaries.org/app/uploads/2025/04/2012-Subsidiarity-Principle-EN.pdf",
        "https://actuaries.org/app/uploads/2025/04/2012-Subsidiarity-Principle-FR.pdf",
        "https://actuaries.org/app/uploads/2025/04/InternalRegulations.pdf",
    )
    for url in administrative_urls:
        assert not runtime._site_candidate_allowed(iaa, url, "", {})
        assert not runtime._site_document_relevant(
            iaa, url, "", "https://actuaries.org/guidelines-and-policies/", {}
        )

    research_url = "https://actuaries.org/paper/artificial-intelligence-governance-framework/"
    assert runtime._site_candidate_allowed(
        iaa, research_url, "Artificial Intelligence Governance Framework", {}
    )


def test_unknown_historical_source_is_not_rewritten_as_a_live_site(
    monkeypatch, tmp_path: Path
) -> None:
    monkeypatch.setattr(
        runtime, "ingest_live_article", lambda **_kwargs: pytest.fail("must not invent source")
    )
    with pytest.raises(runtime.RuntimeErrorSafe, match="without known source-site provenance"):
        runtime._restore_hub_blob(
            {
                "url": "https://example.test/historical.pdf",
                "source_site": None,
                "content_version_id": 18,
                "bilingual_metadata": {"en": {"title": "Historical paper"}},
            },
            b"verified source bytes",
            "unused-hub",
        )


def test_realistic_json_and_embedded_site_adapters_apply_include_rules() -> None:
    cia = runtime.SITES["cia-research-publications"]
    queue: list[dict] = []
    queued: set[str] = set()
    wordpress = [
        {
            "id": 44,
            "link": "https://www.cia-ica.ca/publications/ai-pricing-report",
            "title": {"rendered": "AI pricing research report"},
            "content": {
                "rendered": '<a href="/app/uploads/2026/10/ai-pricing-report.pdf">Download paper</a>'
            },
            "acf": {
                "attachment": {
                    "url": "https://www.cia-ica.ca/app/uploads/2026/10/ai-pricing-report.pdf"
                }
            },
        }
    ]
    runtime._enqueue_wordpress_results(
        cia,
        {"url": "https://www.cia-ica.ca/wp-json/wp/v2/publication?page=1", "page": 1},
        json.dumps(wordpress).encode(),
        {"x-wp-totalpages": "1"},
        "https://www.cia-ica.ca/wp-json/wp/v2/publication?page=1",
        queue,
        queued,
    )
    assert any(
        row["kind"] == "document" and row["url"].endswith("ai-pricing-report.pdf") for row in queue
    )
    assert any(row["kind"] == "page" and row["url"].endswith("ai-pricing-report") for row in queue)

    caa = runtime.SITES["caa"]
    queue, queued = [], set()
    runtime._enqueue_caa_results(
        caa,
        {
            "kind": "json_post",
            "news_type": 12,
            "page": 1,
            "page_size": 50,
            "url": "https://www.e-caa.org.cn/systemNews/getNewsInfo/index?newsType=12",
        },
        json.dumps(
            {
                "rows": [
                    {
                        "informationTitle": "人工智能研究论文",
                        "informationUrl": "/upload/research/ai-paper.pdf",
                    }
                ],
                "records": 1,
            }
        ).encode(),
        "https://www.e-caa.org.cn/systemNews/getNewsInfo/index?newsType=12",
        queue,
        queued,
    )
    assert any(row["kind"] == "document" and row["url"].endswith("ai-paper.pdf") for row in queue)

    australia = runtime.SITES["australia"]
    queue, queued = [], set()
    hits = {
        "hits": [
            {
                "name": "Public AI research report",
                "slug": "/research-analysis/ai-report",
                "accessLevel": ["public"],
                "contentTypes": ["Publication > Report"],
                "practiceAreas": ["AI"],
            },
            {
                "name": "Restricted AI paper",
                "slug": "/research-analysis/restricted",
                "accessLevel": ["members"],
                "contentTypes": ["Publication > Report"],
            },
            {
                "name": "AI webinar",
                "slug": "/events/ai-webinar",
                "accessLevel": ["public"],
                "contentTypes": ["Event"],
            },
        ],
        "nbPages": 1,
    }
    runtime._enqueue_australia_results(
        australia,
        {"page": 0, "content_type": "Publication > Report", "url": "https://search.test/query"},
        json.dumps(hits).encode(),
        "https://search.test/query",
        queue,
        queued,
    )
    assert [row["title"] for row in queue] == ["Public AI research report"]
    parsed = runtime._parse_algolia_script(
        "const appId = '54YLV1QRFB'; const apiKey='abcdefghijklmnopqrstuvwxyz123456'; const indexName='prod_actuaries_search';"
    )
    assert parsed == {
        "api_key": "abcdefghijklmnopqrstuvwxyz123456",
        "app_id": "54YLV1QRFB",
        "index": "prod_actuaries_search",
    }

    swiss_document = json.dumps(
        {
            "items": [
                {
                    "title": "AI pricing research",
                    "downloadLinkDokument": "/media/ai.pdf",
                    "documentCategoryId": 8,
                },
                {
                    "title": "Membership statutes",
                    "downloadLinkDokument": "/media/statutes.pdf",
                    "documentCategoryId": 5,
                },
                {
                    "title": "Professional standard",
                    "downloadLinkDokument": "/media/standard.pdf",
                    "documentCategoryId": 2,
                },
            ]
        }
    )
    records = runtime._parse_swiss_embedded_documents(
        f"<script>window.download = {swiss_document};</script>"
    )
    assert [item["title"] for item in records if runtime._swiss_document_is_relevant(item)] == [
        "AI pricing research",
        "Professional standard",
    ]


class _FixtureCrawler:
    def __init__(self, listing: bytes, documents: list[bytes]) -> None:
        self.listing = listing
        self.documents = iter(documents)
        self.site_attempts = 0
        self.site_budget = 30
        self.requests: list[str] = []

    def begin_site(self, budget: int) -> None:
        self.site_attempts = 0
        self.site_budget = min(runtime.MAX_SITE_REQUESTS, budget)

    def _request(self, url: str, **_kwargs):
        self.site_attempts += 1
        self.requests.append(url)
        return self.listing, {"content-type": "text/html"}, url

    def _download_file(self, url: str, staging_dir: Path, **_kwargs):
        self.site_attempts += 1
        raw = next(self.documents)
        staging_dir.mkdir(parents=True, exist_ok=True)
        path = staging_dir / f"{hashlib.sha256(url.encode()).hexdigest()}.pdf"
        path.write_bytes(raw)
        return (
            path,
            {"content-type": "application/pdf"},
            url,
            hashlib.sha256(raw).hexdigest(),
            len(raw),
        )


def test_iaa_administrative_pdfs_are_rejected_before_frontier_enqueue(
    tmp_path: Path, monkeypatch
) -> None:
    monkeypatch.setattr(runtime, "_ACTIVE_STATE_PATH", tmp_path / "runtime-state.json")
    admin_urls = (
        "/app/uploads/2025/04/Statutes.pdf",
        "/app/uploads/2025/04/Statutes_FR.pdf",
        "/app/uploads/2025/04/2012-Subsidiarity-Principle-EN.pdf",
        "/app/uploads/2025/04/2012-Subsidiarity-Principle-FR.pdf",
        "/app/uploads/2025/04/InternalRegulations.pdf",
    )
    links = "".join(f'<a href="{url}">Administrative document</a>' for url in admin_urls)
    links += (
        '<a href="/paper/artificial-intelligence-governance-framework/">'
        "Artificial Intelligence Governance Framework</a>"
        '<a href="/app/uploads/2026/10/AI-Research-Report.pdf">AI research report</a>'
    )
    listing = f"<html><body>{links}</body></html>".encode()
    site = runtime.SITES["iaa"]
    seed = "https://actuaries.org/publications/papers/"
    site_state = {
        "queue": [{"url": seed, "kind": "page", "depth": 0, "referrer": seed}],
        "cursor": 0,
        "visited": [],
    }
    state = {"articles": {}, "sites": {"iaa": site_state}}

    class OneRequestCrawler(_FixtureCrawler):
        def begin_site(self, _budget: int) -> None:
            self.site_attempts = 0
            self.site_budget = 1

    result = runtime._capture(
        site, state, OneRequestCrawler(listing, []), tmp_path / "data", site_state
    )

    queued_urls = {row["url"] for row in site_state["queue"]}
    assert result["request_count"] == 1
    assert result["frontier_remaining"] == 2
    assert all(f"https://actuaries.org{url}" not in queued_urls for url in admin_urls)
    assert (
        "https://actuaries.org/paper/artificial-intelligence-governance-framework/" in queued_urls
    )
    assert "https://actuaries.org/app/uploads/2026/10/AI-Research-Report.pdf" in queued_urls


@pytest.mark.parametrize(
    ("kind", "content_type"),
    [
        (
            "docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ),
        (
            "pptx",
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        ),
        ("xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    ],
)
def test_extensionless_ooxml_download_uses_mime_for_capture_and_extraction(
    kind: str, content_type: str, tmp_path: Path, monkeypatch
) -> None:
    from io import BytesIO

    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation

    monkeypatch.setattr(runtime, "_ACTIVE_STATE_PATH", tmp_path / "runtime-state.json")
    monkeypatch.setattr(capacity, "ensure_capacity", lambda **_kwargs: None)
    marker = "Supported actuarial research content from an extensionless Office download. " * 4
    output = BytesIO()
    if kind == "docx":
        document = Document()
        document.add_paragraph(marker)
        document.save(output)
    elif kind == "pptx":
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(0, 0, 5_000_000, 1_000_000).text_frame.text = marker
        presentation.save(output)
    else:
        workbook = Workbook()
        workbook.active.append(["Actuarial research", marker])
        workbook.save(output)
    raw = output.getvalue()
    digest = hashlib.sha256(raw).hexdigest()
    url = "https://documents.example/document/download?id=123"
    referrer = "https://documents.example/research/"
    site = runtime.Site("ooxml-fixture", "collect", (url,), ("documents.example",))

    class DownloadCrawler:
        site_attempts = 0
        site_budget = runtime.MAX_SITE_REQUESTS

        def begin_site(self, budget: int) -> None:
            self.site_attempts = 0
            self.site_budget = budget

        def _download_file(self, _url: str, target_dir: Path, *, delay_seconds: float):
            del delay_seconds
            self.site_attempts += 1
            staging = target_dir / "_tmp" / "extensionless-download.part"
            staging.parent.mkdir(parents=True, exist_ok=True)
            staging.write_bytes(raw)
            return staging, {"content-type": content_type}, url, digest, len(raw)

    article_state = {"articles": {}}
    capture = runtime._capture(
        site,
        article_state,
        DownloadCrawler(),
        tmp_path / "data",
        {},
        initial_frontier=[
            {
                "url": url,
                "kind": "document",
                "depth": 0,
                "referrer": referrer,
                "title": "Actuarial research report",
            }
        ],
    )

    assert capture["request_count"] == 1
    assert capture["errors"] == []
    assert len(capture["downloads"]) == 1
    article = capture["downloads"][0]
    assert article["detected_kind"] == kind
    assert article["content_type"] == content_type
    assert article["source_page_url"] == referrer
    assert Path(article["artifact_path"]).read_bytes() == raw
    extracted = runtime._extract_markdown(raw, content_type, article["url"], tmp_path / "data")
    assert marker.strip() in extracted


def test_extensionless_generic_zip_keeps_unsupported_type() -> None:
    raw = b"PK\x03\x04generic ZIP payload"
    url = "https://documents.example/document/download?id=123"
    assert runtime._detect_kind(raw, "application/zip", url) == "zip"


def test_invalid_json_frontier_item_is_retried_after_restart(tmp_path: Path, monkeypatch) -> None:
    state_path = tmp_path / "runtime-state.json"
    monkeypatch.setattr(runtime, "_ACTIVE_STATE_PATH", state_path)
    site = runtime.SITES["cia-research-publications"]
    rest_url = runtime._initial_frontier(site)[0]["url"]
    site_state = {
        "queue": [{"url": rest_url, "kind": "json_rest", "page": 1}],
        "cursor": 0,
        "visited": [],
    }
    state = {"articles": {}, "sites": {site.site_id: site_state}}

    class RetryJsonCrawler:
        site_attempts = 0
        site_budget = runtime.MAX_SITE_REQUESTS
        responses = [b"{invalid JSON", b"[]"]
        requests: list[str] = []

        def begin_site(self, _budget: int) -> None:
            self.site_attempts = 0

        def _request(self, url: str, **_kwargs):
            self.site_attempts += 1
            self.requests.append(url)
            return self.responses.pop(0), {"content-type": "application/json"}, url

    crawler = RetryJsonCrawler()
    first = runtime._capture(site, state, crawler, tmp_path / "data", site_state)
    assert first["errors"]
    assert site_state["visited"] == []
    assert site_state["retry_queue"][0]["attempts"] == 1

    retry = runtime._capture(site, state, crawler, tmp_path / "data", site_state)
    assert retry["errors"] == []
    assert crawler.requests == [runtime._canonical_url(rest_url)] * 2
    assert site_state.get("retry_queue", []) == []
    assert runtime._queue_identity(site_state["queue"][1]) in site_state["visited"]
    assert site_state["completed_frontier"] is True


def test_model_evidence_must_be_an_exact_short_contiguous_markdown_quote() -> None:
    category_evidence_schema = runtime._model_schema()["properties"]["categories"]["items"][
        "properties"
    ]["evidence"]
    assert category_evidence_schema["minItems"] == 1
    assert category_evidence_schema["items"]["pattern"] == runtime.EVIDENCE_SPAN_ID_PATTERN

    markdown = "AI evidence.\nPublication date: September 2026."
    model_result = {
        "categories": [{"key": "AI", "evidence": ["AI evidence.", "AI  evidence."]}],
        "date_facts": {
            "publication_date": {"evidence": "Publication date:  September 2026."},
            "revision_date": None,
            "document_version_date": None,
        },
        "date_review_unresolved": [{"field": "revision_date", "evidence": "September 2026."}],
    }

    assert runtime._model_evidence_errors(markdown, model_result) == [
        "categories[0].evidence[1]",
        "date_facts.publication_date.evidence",
    ]
    long_quote_result = {
        **model_result,
        "categories": [{"key": "AI", "evidence": ["x" * (runtime.MAX_EVIDENCE_QUOTE_CHARS + 1)]}],
    }
    assert runtime._model_evidence_errors(markdown, long_quote_result) == [
        "categories[0].evidence[0]",
        "date_facts.publication_date.evidence",
    ]
    empty_category_evidence = {
        **model_result,
        "categories": [{"key": "AI", "evidence": []}],
    }
    assert "categories[0].evidence" in runtime._model_evidence_errors(
        markdown, empty_category_evidence
    )


def test_evidence_span_ids_map_to_exact_category_and_date_source_text() -> None:
    markdown = (
        "Pricing analysis uses AI risk models.\n"
        "Publication date: September 2026.\n"
        "The revision date remains unresolved.\n"
    ) * 12
    spans = runtime._build_evidence_span_catalog(markdown)
    assert spans == runtime._build_evidence_span_catalog(markdown)
    assert spans[0]["start"] == 0
    assert spans[-1]["end"] == len(markdown)
    assert all(0 < len(row["text"]) <= runtime.MAX_EVIDENCE_SPAN_CHARS for row in spans)
    assert all(row["text"] == markdown[row["start"] : row["end"]] for row in spans)
    assert all(current["start"] <= previous["end"] for previous, current in zip(spans, spans[1:]))
    assert len({row["id"] for row in spans}) == len(spans)

    def id_for(quote: str) -> str:
        return next(row["id"] for row in spans if quote in row["text"])

    def text_for(span_id: str) -> str:
        return next(row["text"] for row in spans if row["id"] == span_id)

    pricing_id = id_for("Pricing analysis uses AI risk models.")
    ai_id = id_for("AI risk models.")
    publication_id = id_for("Publication date: September 2026.")
    unresolved_id = id_for("revision date remains unresolved.")
    raw_result = {
        "categories": [
            {"key": "Pricing", "evidence": [pricing_id]},
            {"key": "AI", "evidence": [ai_id]},
        ],
        "date_facts": {
            "publication_date": {
                "value": "2026-09",
                "precision": "month",
                "semantic": "publication",
                "evidence": publication_id,
                "location": "cover",
            },
            "revision_date": None,
            "document_version_date": None,
        },
        "date_review_unresolved": [{"field": "revision_date", "evidence": unresolved_id}],
    }
    resolved, errors = runtime._resolve_model_evidence_span_ids(raw_result, spans)
    assert errors == []
    assert resolved["categories"][0]["evidence"] == [text_for(pricing_id)]
    assert resolved["categories"][1]["evidence"] == [text_for(ai_id)]
    assert resolved["date_facts"]["publication_date"]["evidence"] == text_for(publication_id)
    assert resolved["date_review_unresolved"][0]["evidence"] == text_for(unresolved_id)
    assert "Pricing analysis uses AI risk models." in resolved["categories"][0]["evidence"][0]
    assert "AI risk models." in resolved["categories"][1]["evidence"][0]
    assert (
        "Publication date: September 2026."
        in resolved["date_facts"]["publication_date"]["evidence"]
    )
    assert "revision date remains unresolved." in resolved["date_review_unresolved"][0]["evidence"]
    assert runtime._model_evidence_errors(markdown, resolved) == []

    unknown_id = "span-" + "f" * 16
    if unknown_id in {row["id"] for row in spans}:
        unknown_id = "span-" + "0" * 16
    unknown_result = {
        **raw_result,
        "date_facts": {
            **raw_result["date_facts"],
            "publication_date": {
                **raw_result["date_facts"]["publication_date"],
                "evidence": unknown_id,
            },
        },
    }
    unresolved, unknown_paths = runtime._resolve_model_evidence_span_ids(unknown_result, spans)
    assert "date_facts.publication_date.evidence" in unknown_paths
    assert "date_facts.publication_date.evidence" in runtime._model_evidence_errors(
        markdown, unresolved
    )

    empty_result = {
        **raw_result,
        "categories": [raw_result["categories"][0], {"key": "AI", "evidence": []}],
    }
    _, empty_paths = runtime._resolve_model_evidence_span_ids(empty_result, spans)
    assert "categories[1].evidence" in empty_paths


def test_raw_only_extraction_gap_retries_same_hub_version_without_fake_ai_or_ready(
    tmp_path: Path, monkeypatch
) -> None:
    schema = "test_r7_raw_" + hashlib.sha256(str(tmp_path).encode()).hexdigest()[:12]
    hub_url = __import__("os").environ["POSTGRES_TEST_URL"]
    runtime.apply_content_hub_schema(hub_url, schema=schema)

    def use_schema(function):
        def wrapped(*args, **kwargs):
            kwargs["schema"] = schema
            return function(*args, **kwargs)

        return wrapped

    monkeypatch.setattr(
        runtime, "content_hub_read_version", use_schema(content_hub.content_hub_read_version)
    )
    monkeypatch.setattr(
        runtime, "content_hub_url_exclusion", use_schema(content_hub.content_hub_url_exclusion)
    )
    monkeypatch.setattr(runtime, "ingest_live_article", use_schema(content_hub.ingest_live_article))
    monkeypatch.setattr(
        runtime, "save_live_enrichment", use_schema(content_hub.save_live_enrichment)
    )

    state_path = tmp_path / "runtime-state.json"
    monkeypatch.setattr(runtime, "_ACTIVE_STATE_PATH", state_path)
    raw = bytes.fromhex("d0cf11e0a1b11ae1") + b"legacy Office source bytes"
    raw_sha = hashlib.sha256(raw).hexdigest()
    article_path = tmp_path / "legacy-source.doc"
    article_path.write_bytes(raw)
    url = "https://actuaries.org/app/uploads/2026/10/legacy-source.doc"
    article = {
        "site_id": "iaa",
        "discovery_sites": ["iaa"],
        "url": url,
        "source_page_url": "https://actuaries.org/publications/papers/",
        "title": "Legacy actuarial source",
        "raw_sha256": raw_sha,
        "artifact_path": str(article_path),
        "content_type": "application/msword",
        "detected_kind": "office_ole",
        "stage": "downloaded",
    }
    state = {
        "articles": {runtime._file_key(url, raw_sha): article},
        "target_acks": {},
    }

    flow = runtime._load_flow_module()
    receipt_root = tmp_path / "receipts"
    previous_stages = {
        stage: runtime._receipt_stage("succeeded", count=0) for stage in flow.COUNT_STAGES
    }
    previous_stages.update(
        {
            "china_ack": runtime._ack_stage(
                {
                    "acknowledged": 1,
                    "pending": 0,
                    "failed": 0,
                    "ready_status": "failed",
                    "ready_error": "prior native Ready failure task=known-task",
                    "ready_evidence": ["prior log=/logs/known-task.log"],
                }
            ),
            "overseas_ack": runtime._ack_stage(
                {
                    "acknowledged": 1,
                    "pending": 0,
                    "failed": 0,
                    "ready_status": "succeeded",
                }
            ),
        }
    )
    runtime._write_receipt(
        receipt_root,
        "2026-10-25",
        "iaa",
        previous_stages,
        [],
        reason=None,
    )
    monkeypatch.setattr(
        runtime,
        "_stream_target",
        lambda *_args, **_kwargs: pytest.fail("raw-only article must not call a target"),
    )

    first = runtime._process_articles(
        "iaa",
        state,
        {"data_root": str(tmp_path / "hubdata"), "state_root": str(tmp_path)},
        hub_url,
        preview=False,
        article_limit=1,
    )
    assert first["new_versions"] == 1
    assert first["ai_candidates"] == 0
    assert "extraction gap (office_ole)" in first["article_errors"][0]
    assert article["hub_version"]["status"] == "new"
    assert all(
        first["target"][target]["ready_status"] == "unknown" for target in ("china", "overseas")
    )
    assert all(first["target"][target]["acknowledged"] == 0 for target in ("china", "overseas"))
    raw_only = content_hub.content_hub_read_version(url, raw_sha, hub_url, schema=schema)
    assert raw_only["content_version_id"] == article["hub_version"]["content_version_id"]
    assert raw_only["markdown"] is None
    assert raw_only["markdown_extraction_status"] == "extraction_gap"
    assert raw_only["annotation"] is None

    def save_site_receipt(summary, date_text):
        stages = {
            stage: runtime._receipt_stage("succeeded", count=0) for stage in flow.COUNT_STAGES
        }
        article_error = "; ".join(summary["article_errors"])
        stages.update(
            {
                "downloaded": runtime._receipt_stage("succeeded", count=1),
                "new_versions": runtime._receipt_stage(
                    "partial", count=summary["new_versions"], error=article_error
                ),
                "enriched": runtime._receipt_stage("partial", count=0, error=article_error),
                "ai_candidates": runtime._receipt_stage("partial", count=0, error=article_error),
                "china_ack": runtime._ack_stage(summary["target"]["china"]),
                "overseas_ack": runtime._ack_stage(summary["target"]["overseas"]),
            }
        )
        return runtime._write_receipt(receipt_root, date_text, "iaa", stages, [], article_error)

    first_receipt = save_site_receipt(first, "2026-10-26")
    assert first_receipt["run_status"] == "partial"
    assert first_receipt["stages"]["china_ack"]["ready_status"] == "unknown"

    retry = runtime._process_articles(
        "iaa",
        state,
        {"data_root": str(tmp_path / "hubdata"), "state_root": str(tmp_path)},
        hub_url,
        preview=False,
        article_limit=1,
    )
    assert retry["new_versions"] == 0
    assert retry["ai_candidates"] == 0
    assert "extraction gap (office_ole)" in retry["article_errors"][0]
    assert (
        content_hub.content_hub_read_version(url, raw_sha, hub_url, schema=schema)[
            "content_version_id"
        ]
        == raw_only["content_version_id"]
    )
    status = content_hub.content_hub_status(hub_url, schema=schema)
    assert status["content_version_count"] == 1
    assert status["source_count"] == 1
    second_receipt = save_site_receipt(retry, "2026-10-26")
    assert second_receipt["attempt_count"] == 2

    report_process = subprocess.run(
        [
            sys.executable,
            str(runtime.FLOW_SCRIPT),
            "report",
            "--date",
            "2026-10-26",
            "--receipts",
            str(receipt_root),
            "--json",
        ],
        check=False,
        capture_output=True,
        text=True,
    )
    assert report_process.returncode == 0, report_process.stderr
    report = json.loads(report_process.stdout)
    prior_ready_case = next(
        case
        for case in report["failure_cases"]
        if case["site_id"] == "iaa" and case["stage"] == "china_ready"
    )
    assert "known-task" in prior_ready_case["error"]
    plan_process = subprocess.run(
        [
            sys.executable,
            str(runtime.FLOW_SCRIPT),
            "plan",
            "--date",
            "2026-10-27",
            "--receipts",
            str(receipt_root),
            "--json",
        ],
        check=False,
        capture_output=True,
        text=True,
    )
    assert plan_process.returncode == 0, plan_process.stderr
    assert "iaa" in json.loads(plan_process.stdout)["extra_retry_sites"]

    markdown = (
        "The same original actuarial source is now available to a supported converter. "
        "Its evidence remains bound to the original PDF or Office bytes."
    )
    monkeypatch.setattr(runtime, "_extract_markdown", lambda *_args: markdown)
    span_id = runtime._build_evidence_span_catalog(markdown)[0]["id"]
    package = {
        "title": "AI actuarial source",
        "summary": "An AI research document for actuarial practice.",
        "category": "AI",
        "keywords": ["AI", "actuarial"],
    }
    model_result = {
        "primary_category": "AI",
        "bilingual_metadata": {"zh": package, "en": package},
        "categories": [{"key": "AI", "evidence": [span_id]}],
        "date_facts": {
            "publication_date": None,
            "revision_date": None,
            "document_version_date": None,
        },
        "date_review_status": "complete",
        "date_review_reason": None,
        "date_review_unresolved": [],
    }
    monkeypatch.setattr(runtime, "_run_model", lambda *_args, **_kwargs: model_result)
    completed_counts: list[tuple[bool, int]] = []
    completed, new_version, enriched = runtime._prepare_hub_article(
        article,
        state,
        {"data_root": str(tmp_path / "hubdata"), "state_root": str(tmp_path)},
        hub_url,
        preview=False,
        on_hub_version=lambda is_new, row: completed_counts.append(
            (is_new, int(row["content_version_id"]))
        ),
    )
    assert completed["content_version_id"] == raw_only["content_version_id"]
    assert new_version is False
    assert enriched is True
    assert completed["ai_candidate"] is True
    assert completed_counts == [(False, raw_only["content_version_id"])]
    final_readback = content_hub.content_hub_read_version(url, raw_sha, hub_url, schema=schema)
    assert final_readback["content_version_id"] == raw_only["content_version_id"]
    assert final_readback["markdown"] == markdown
    assert final_readback["markdown_extraction_status"] == "available"
    assert final_readback["annotation"] is not None
    final_status = content_hub.content_hub_status(hub_url, schema=schema)
    assert final_status["content_version_count"] == 1
    assert final_status["source_count"] == 2


@pytest.mark.parametrize(
    "empty_evidence",
    [False, True],
    ids=["unknown-span-id", "empty-evidence"],
)
@pytest.mark.parametrize("correction_is_valid", [True, False])
def test_invalid_model_evidence_gets_one_specific_correction_before_hub_write(
    correction_is_valid: bool,
    empty_evidence: bool,
    tmp_path: Path,
    monkeypatch,
) -> None:
    markdown = "This is exact AI evidence."
    evidence_spans = runtime._build_evidence_span_catalog(markdown)
    valid_span_id = evidence_spans[0]["id"]
    unknown_span_id = "span-" + "f" * 16
    if unknown_span_id == valid_span_id:
        unknown_span_id = "span-" + "0" * 16
    invalid_category_evidence = [] if empty_evidence else [unknown_span_id]
    raw_sha = "a" * 64
    markdown_sha = hashlib.sha256(markdown.encode("utf-8")).hexdigest()
    initial = {
        "content_version_id": 41,
        "markdown_text_variant_id": 52,
        "markdown_sha256": markdown_sha,
        "markdown": markdown,
        "annotation": {},
        "bilingual_metadata": {},
    }
    categories = [{"key": "AI", "evidence": ["This is exact AI evidence."]}]
    model_categories = [{"key": "AI", "evidence": [valid_span_id]}]
    bilingual = {
        language: {
            "title": "AI paper",
            "summary": "AI research summary",
            "category": "AI",
            "keywords": ["AI"],
        }
        for language in ("zh", "en")
    }
    enriched = {
        **initial,
        "annotation": {
            "categories": categories,
            "date_facts": {},
            "date_review_unresolved": [],
        },
        "bilingual_metadata": bilingual,
    }
    reads = iter([initial, enriched])
    monkeypatch.setattr(runtime, "content_hub_read_version", lambda *_args: next(reads))
    monkeypatch.setattr(runtime, "_load_category_keys", lambda: ["AI"])
    responses = [
        {
            "primary_category": "AI",
            "bilingual_metadata": bilingual,
            "categories": [{"key": "AI", "evidence": invalid_category_evidence}],
            "date_facts": {
                "publication_date": None,
                "revision_date": None,
                "document_version_date": None,
            },
            "date_review_status": "complete",
            "date_review_reason": None,
            "date_review_unresolved": [],
        },
        {
            "primary_category": "AI",
            "bilingual_metadata": bilingual,
            "categories": (
                model_categories
                if correction_is_valid
                else [{"key": "AI", "evidence": [unknown_span_id]}]
            ),
            "date_facts": {
                "publication_date": None,
                "revision_date": None,
                "document_version_date": None,
            },
            "date_review_status": "complete",
            "date_review_reason": None,
            "date_review_unresolved": [],
        },
    ]
    model_calls: list[dict[str, Any]] = []

    def run_model(*_args, **kwargs):
        model_calls.append(kwargs)
        return responses.pop(0)

    saved_categories: list[list[dict[str, Any]]] = []

    def save_enrichment(**kwargs):
        saved_categories.append(kwargs["categories"])
        return {"success": True}

    monkeypatch.setattr(runtime, "_run_model", run_model)
    monkeypatch.setattr(runtime, "save_live_enrichment", save_enrichment)
    article = {
        "url": "https://example.test/ai-paper.pdf",
        "raw_sha256": raw_sha,
        "title": "AI paper",
    }

    if correction_is_valid:
        result, new_version, newly_enriched = runtime._prepare_hub_article(
            article,
            {},
            {"data_root": str(tmp_path), "state_root": str(tmp_path)},
            "unused-hub",
            False,
        )
        assert result["categories"] == categories
        assert new_version is False
        assert newly_enriched is True
        assert saved_categories == [categories]
    else:
        with pytest.raises(runtime.RuntimeErrorSafe, match="after one correction attempt"):
            runtime._prepare_hub_article(
                article,
                {},
                {"data_root": str(tmp_path), "state_root": str(tmp_path)},
                "unused-hub",
                False,
            )
        assert saved_categories == []

    assert len(model_calls) == 2
    assert "evidence_feedback" not in model_calls[0]
    expected_feedback = (
        "categories[0].evidence[0]" if not empty_evidence else "categories[0].evidence"
    )
    assert model_calls[1]["evidence_feedback"] == [expected_feedback]


@pytest.mark.parametrize("site_id", ["caa", "australia"])
def test_source_bootstrap_errors_are_recorded_in_v2_site_receipts(
    site_id: str, tmp_path: Path, monkeypatch
) -> None:
    monkeypatch.setattr(runtime, "_ACTIVE_STATE_PATH", tmp_path / "runtime-state.json")

    class BootstrapFailureCrawler:
        site_attempts = 0
        site_budget = runtime.MAX_SITE_REQUESTS
        caa_cookie = None
        algolia_key = None

        def begin_site(self, _budget: int) -> None:
            self.site_attempts = 0

        def _request(self, _url: str, **_kwargs):
            self.site_attempts += 1
            raise OSError("source bootstrap unavailable")

        def _request_post(self, *_args, **_kwargs):
            self.site_attempts += 1
            raise OSError("source endpoint unavailable")

    monkeypatch.setattr(runtime, "_new_crawler", lambda _root: BootstrapFailureCrawler())
    monkeypatch.setattr(
        runtime,
        "_process_articles",
        lambda *_args, **_kwargs: {
            "candidate_count": 0,
            "new_versions": 0,
            "duplicates": 0,
            "enriched": 0,
            "ai_candidates": 0,
            "article_errors": [],
            "target": {
                name: {
                    "acknowledged": 0,
                    "pending": 0,
                    "failed": 0,
                    "ready_status": "succeeded",
                }
                for name in ("china", "overseas")
            },
        },
    )
    state = {"sites": {}, "articles": {}, "target_acks": {}}
    if site_id == "australia":
        state["sites"][site_id] = {
            "queue": [
                {
                    "url": f"https://{runtime.ALGOLIA_HOST}/1/indexes/public/query",
                    "kind": "algolia",
                    "page": 0,
                    "content_type": runtime.AUSTRALIA_RESEARCH_TYPES[0],
                    "script_url": "https://www.actuaries.asn.au/assets/search.js",
                }
            ],
            "cursor": 0,
            "visited": [],
        }
    config = {
        "state_root": str(tmp_path),
        "receipts_root": str(tmp_path / "receipts"),
        "data_root": str(tmp_path / "data"),
    }

    result = runtime._run_site(site_id, "2026-10-25", config, state, "unused-hub")

    receipt_path = tmp_path / "receipts" / "2026-10-25" / f"{site_id}.json"
    receipt = json.loads(receipt_path.read_text(encoding="utf-8"))
    assert receipt["schema_version"] == 2
    assert receipt["run_status"] in {"partial", "failed"}
    assert result["stages"]["downloaded"]["status"] in {"partial", "failed"}
    assert "bootstrap" in result["stages"]["downloaded"]["error"].lower()


def test_exact_url_processes_canonical_same_host_redirect(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setattr(capacity, "ensure_capacity", lambda **_kwargs: None)
    monkeypatch.setattr(runtime, "_ACTIVE_STATE_PATH", tmp_path / "runtime-state.json")
    requested_url = (
        "https://actuaries.org/app/uploads/2026/10/Requested_AI_Paper_Final_Approved.pdf"
    )
    final_url = "https://actuaries.org/app/uploads/2026/10/Canonical_AI_Paper.pdf"
    raw = b"%PDF-1.7\nverified same-host redirect payload"

    class RedirectCrawler:
        site_attempts = 0
        site_budget = runtime.MAX_SITE_REQUESTS

        def begin_site(self, _budget: int) -> None:
            self.site_attempts = 0

        def _download_file(self, _url: str, staging_dir: Path, **_kwargs):
            self.site_attempts += 1
            staging_dir.mkdir(parents=True, exist_ok=True)
            staging_path = staging_dir / "redirected.pdf"
            staging_path.write_bytes(raw)
            return (
                staging_path,
                {"content-type": "application/pdf"},
                final_url,
                hashlib.sha256(raw).hexdigest(),
                len(raw),
            )

    processed_urls: list[str | None] = []
    monkeypatch.setattr(runtime, "_new_crawler", lambda _root: RedirectCrawler())

    def process_articles(*_args, **kwargs):
        processed_urls.append(kwargs["include_url"])
        return {
            "candidate_count": 1,
            "new_versions": 1,
            "duplicates": 0,
            "enriched": 1,
            "ai_candidates": 1,
            "article_errors": [],
            "target": {
                name: {
                    "acknowledged": 0,
                    "pending": 0,
                    "failed": 0,
                    "ready_status": "succeeded",
                }
                for name in ("china", "overseas")
            },
        }

    monkeypatch.setattr(runtime, "_process_articles", process_articles)
    config = {
        "state_root": str(tmp_path),
        "receipts_root": str(tmp_path / "receipts"),
        "data_root": str(tmp_path / "data"),
    }
    state = {"sites": {}, "articles": {}, "target_acks": {}}

    result = runtime._run_site(
        "iaa",
        "2026-10-25",
        config,
        state,
        "unused-hub",
        exact_url=requested_url,
        source_page_url="https://actuaries.org/publications/papers/",
    )

    assert processed_urls == [final_url]
    assert next(iter(state["articles"].values()))["url"] == final_url
    assert result["stages"]["downloaded"]["count"] == 1
    assert result["stages"]["new_versions"]["count"] == 1


@pytest.mark.parametrize(
    ("site_id", "exact_url", "source_page_url"),
    [
        (
            "cia-ica",
            "https://www.cia-ica.ca/app/themes/wicket/custom/dl_file.php?p=840374&fid=840377",
            "https://www.cia-ica.ca/publications/",
        ),
        (
            "cia-research-publications",
            "https://www.cia-ica.ca/app/themes/wicket/custom/dl_file.php?p=840374&fid=840377",
            "https://www.cia-ica.ca/publications/",
        ),
        (
            "cca",
            "https://www.ccactuaries.org/docs/default-source/papers/evaluating-the-cost-effectiveness-of-ai-driven-health-care-solutions_october-2025.pdf",
            "https://www.ccactuaries.org/professional-resources/papers",
        ),
        (
            "france",
            "https://www.institutdesactuaires.com/global/gene/link.php?doc_id=20148&fg=1",
            "https://www.institutdesactuaires.com/article/conferences-1105",
        ),
    ],
)
def test_exact_url_capture_does_not_repair_normal_site_frontier(
    tmp_path: Path,
    monkeypatch,
    site_id: str,
    exact_url: str,
    source_page_url: str,
) -> None:
    monkeypatch.setattr(capacity, "ensure_capacity", lambda **_kwargs: None)
    monkeypatch.setattr(runtime, "_ACTIVE_STATE_PATH", tmp_path / "runtime-state.json")
    raw = b"%PDF-1.7\nexact source payload"

    class ExactCrawler:
        def __init__(self) -> None:
            self.calls: list[str] = []
            self.site_attempts = 0
            self.site_budget = runtime.MAX_SITE_REQUESTS

        def begin_site(self, budget: int) -> None:
            self.site_attempts = 0
            self.site_budget = budget

        def _download_file(self, url: str, staging_dir: Path, **_kwargs):
            self.calls.append(url)
            self.site_attempts += 1
            staging_dir.mkdir(parents=True, exist_ok=True)
            staging_path = staging_dir / f"{hashlib.sha256(url.encode()).hexdigest()}.pdf"
            staging_path.write_bytes(raw)
            return (
                staging_path,
                {"content-type": "application/pdf"},
                url,
                hashlib.sha256(raw).hexdigest(),
                len(raw),
            )

        def _request(self, url: str, **_kwargs):
            self.calls.append(url)
            self.site_attempts += 1
            body = (
                "<html><body><h1>Research publications</h1><p>"
                + "Verified actuarial research listing. " * 8
                + "</p></body></html>"
            ).encode()
            return body, {"content-type": "text/html"}, url

    crawler = ExactCrawler()
    monkeypatch.setattr(runtime, "_new_crawler", lambda _root: crawler)
    processed_urls: list[str | None] = []

    def process_articles(*_args, **kwargs):
        processed_urls.append(kwargs["include_url"])
        return {
            "candidate_count": 1,
            "new_versions": 1,
            "duplicates": 0,
            "enriched": 1,
            "ai_candidates": 1,
            "article_errors": [],
            "target": {
                name: {
                    "acknowledged": 0,
                    "pending": 0,
                    "failed": 0,
                    "ready_status": "succeeded",
                }
                for name in ("china", "overseas")
            },
        }

    monkeypatch.setattr(runtime, "_process_articles", process_articles)
    site = runtime.SITES[site_id]
    normal_queue = [{"url": site.seeds[0], "kind": "page"}]
    state = {
        "sites": {
            site_id: {
                "queue": normal_queue.copy(),
                "cursor": 1,
                "visited": [runtime._queue_identity(normal_queue[0])],
                "retry_queue": [],
            }
        },
        "articles": {},
        "target_acks": {},
    }
    config = {
        "state_root": str(tmp_path),
        "receipts_root": str(tmp_path / "receipts"),
        "data_root": str(tmp_path / "data"),
    }
    canonical_exact_url = runtime._canonical_url(exact_url)

    for attempt in range(2):
        result = runtime._run_site(
            site_id,
            "2026-10-27",
            config,
            state,
            "unused-hub",
            exact_url=exact_url,
            source_page_url=source_page_url,
        )
        assert result["stages"]["downloaded"]["count"] == attempt + 1

    assert crawler.calls == [canonical_exact_url, canonical_exact_url]
    assert processed_urls == [canonical_exact_url, canonical_exact_url]
    assert state["sites"][site_id]["queue"] == normal_queue
    assert state["sites"][site_id]["cursor"] == 1
    replay_state = next(iter(state["sites"][site_id]["exact_replays"].values()))
    assert [row["url"] for row in replay_state["queue"]] == [canonical_exact_url]


def test_frontier_budget_and_restart_advance_without_replaying_first_page(
    tmp_path: Path, monkeypatch
) -> None:
    monkeypatch.setattr(capacity, "ensure_capacity", lambda **_kwargs: None)
    state_path = tmp_path / "runtime-state.json"
    monkeypatch.setattr(runtime, "_ACTIVE_STATE_PATH", state_path)
    links = "".join(
        f'<a href="/research/document-{number}.pdf">Research report {number}</a>'
        for number in range(32)
    )
    listing = f"<html><body>{links}</body></html>".encode()
    documents = [f"%PDF-1.7\nDocument {number}".encode() for number in range(32)]
    site = runtime.Site(
        "fixture", "collect", ("https://example.test/research",), ("example.test",), 0.5, 1
    )
    site_state: dict = {
        "queue": [{"url": site.seeds[0], "kind": "page", "depth": 0, "referrer": site.seeds[0]}],
        "cursor": 0,
        "visited": [],
    }
    state = {"articles": {}, "sites": {"fixture": site_state}}
    first = _FixtureCrawler(listing, documents)
    result1 = runtime._capture(site, state, first, tmp_path / "data", site_state)
    assert result1["request_count"] == 30
    assert result1["frontier_remaining"] == 3
    assert len(result1["downloads"]) == 29
    assert site_state["cursor"] == 30
    assert len(json.loads(state_path.read_text())["articles"]) == 29

    second = _FixtureCrawler(b"", documents[29:])
    result2 = runtime._capture(site, state, second, tmp_path / "data", site_state)
    assert result2["request_count"] == 3
    assert result2["frontier_remaining"] == 0
    assert len(result2["downloads"]) == 3
    assert first.requests == [site.seeds[0]]
    assert second.requests == []
    assert site_state["completed_frontier"] is True

    next_cycle = _FixtureCrawler(listing, documents)
    result3 = runtime._capture(site, state, next_cycle, tmp_path / "data", site_state)
    assert result3["request_count"] == runtime.MAX_SITE_REQUESTS
    assert next_cycle.requests == [site.seeds[0]]
    assert site_state["cycle"] == 1


def test_receipt_counts_merge_across_restart_and_target_retries_are_independent(
    tmp_path: Path, monkeypatch
) -> None:
    receipt_root = tmp_path / "receipts"
    count_stages = ("downloaded", "new_versions", "duplicates", "enriched", "ai_candidates")

    def stages(count: int, acknowledged: int = 1) -> dict:
        result = {
            name: runtime._receipt_stage("succeeded", count=count if name == "downloaded" else 0)
            for name in count_stages
        }
        result["china_ack"] = runtime._ack_stage(
            {"acknowledged": acknowledged, "pending": 0, "failed": 0, "ready_status": "succeeded"}
        )
        result["overseas_ack"] = runtime._ack_stage(
            {"acknowledged": acknowledged, "pending": 0, "failed": 0, "ready_status": "succeeded"}
        )
        return result

    first = runtime._write_receipt(receipt_root, "2026-10-25", "cas", stages(2), [], None)
    second = runtime._write_receipt(receipt_root, "2026-10-25", "cas", stages(3), [], None)
    assert first["attempt_count"] == 1
    assert second["attempt_count"] == 2
    assert second["stages"]["downloaded"]["count"] == 5
    assert second["stages"]["china_ack"]["acknowledged"] == 2

    monkeypatch.setattr(runtime, "_ACTIVE_STATE_PATH", tmp_path / "state.json")
    monkeypatch.setattr(runtime, "_atomic_json", lambda *_args, **_kwargs: None)
    raw = b"same original bytes"
    markdown = "AI research source-bound Markdown"
    item = {
        "url": "https://example.test/ai.pdf",
        "raw_sha256": hashlib.sha256(raw).hexdigest(),
        "content_version_id": 17,
        "markdown_text_variant_id": 27,
        "markdown_sha256": hashlib.sha256(markdown.encode()).hexdigest(),
        "markdown": markdown,
        "bilingual_metadata": {
            "zh": {"title": "AI", "summary": "summary", "category": "AI", "keywords": ["AI"]},
            "en": {"title": "AI", "summary": "summary", "category": "AI", "keywords": ["AI"]},
        },
        "categories": [{"key": "AI", "evidence": ["AI research"]}],
        "date_facts": {},
        "ai_candidate": True,
    }
    item.update({"annotation": {"categories": item["categories"]}, "source_primary_category": "AI"})
    article_key = runtime._file_key(item["url"], item["raw_sha256"])
    state = {
        "articles": {article_key: {**item, "stage": "downloaded", "discovery_sites": ["fixture"]}},
        "target_acks": {},
    }
    monkeypatch.setattr(
        runtime, "_prepare_hub_article", lambda *_args, **_kwargs: (item, True, True)
    )
    calls: list[str] = []

    def target_write(_config, target, items, ready_pending, **_kwargs):
        calls.append(target)
        if target == "overseas" and calls.count("overseas") == 1:
            raise RuntimeError("temporary overseas outage")
        ack = {
            "content_version_id": items[0]["content_version_id"],
            "raw_sha256": items[0]["raw_sha256"],
            "metadata_fingerprint": runtime._target_metadata_fingerprint(items[0]),
            "status": "written",
        }
        return {"acknowledgements": [ack], "failures": [], "ready_status": "succeeded"}

    monkeypatch.setattr(runtime, "_stream_target", target_write)
    config = {"state_root": str(tmp_path), "data_root": str(tmp_path)}
    first_run = runtime._process_articles(
        "fixture", state, config, "unused-hub", preview=False, article_limit=1
    )
    assert first_run["target"]["china"]["acknowledged"] == 1
    assert first_run["target"]["overseas"]["ready_status"] == "failed"
    assert state["target_acks"]["china:17:" + item["raw_sha256"]]["status"] == "acknowledged"

    second_run = runtime._process_articles(
        "fixture", state, config, "unused-hub", preview=False, article_limit=1
    )
    assert second_run["target"]["china"]["acknowledged"] == 1
    assert second_run["target"]["overseas"]["acknowledged"] == 1
    assert calls == ["china", "overseas", "overseas"]


def test_ready_only_retry_counts_prior_item_ack_and_clears_pending(
    tmp_path: Path, monkeypatch
) -> None:
    state_path = tmp_path / "state.json"
    monkeypatch.setattr(runtime, "_ACTIVE_STATE_PATH", state_path)
    raw = b"same original bytes"
    markdown = "AI research source-bound Markdown"
    item = {
        "url": "https://example.test/ai.pdf",
        "raw_sha256": hashlib.sha256(raw).hexdigest(),
        "content_version_id": 71,
        "markdown_text_variant_id": 81,
        "markdown_sha256": hashlib.sha256(markdown.encode()).hexdigest(),
        "markdown": markdown,
        "bilingual_metadata": {
            "zh": {"title": "AI", "summary": "summary", "category": "AI", "keywords": ["AI"]},
            "en": {"title": "AI", "summary": "summary", "category": "AI", "keywords": ["AI"]},
        },
        "categories": [{"key": "AI", "evidence": ["AI research"]}],
        "date_facts": {},
        "ai_candidate": True,
    }
    article_key = runtime._file_key(item["url"], item["raw_sha256"])
    article = {
        **item,
        "site_id": "iaa",
        "stage": "downloaded",
        "discovery_sites": ["fixture"],
    }
    state = {"articles": {article_key: article}, "target_acks": {}}
    monkeypatch.setattr(
        runtime, "_prepare_hub_article", lambda *_args, **_kwargs: (item, True, True)
    )
    calls: list[tuple[str, list[dict], bool]] = []
    china_ready_attempt = 0

    def target_write(_config, target, items, ready_pending, **_kwargs):
        nonlocal china_ready_attempt
        calls.append((target, list(items), ready_pending))
        acknowledgements = [
            {
                "content_version_id": item["content_version_id"],
                "raw_sha256": item["raw_sha256"],
                "metadata_fingerprint": runtime._target_metadata_fingerprint(item),
                "status": "written",
            }
            for item in items
        ]
        if target == "china":
            china_ready_attempt += 1
            if china_ready_attempt <= 2:
                task_record = {
                    "id": f"china-ai-ready-{china_ready_attempt}",
                    "type": "rag_indexing",
                    "status": "error",
                    "errors": [f"index error {china_ready_attempt}"],
                    "error_details": {"kb_id": "ai", "attempt": china_ready_attempt},
                    "log_file": f"/logs/china-ai-{china_ready_attempt}.log",
                }
                return {
                    "acknowledgements": acknowledgements,
                    "failures": [],
                    "ready_status": "failed",
                    "ready_error": {
                        "ai": f"native rag_indexing task {task_record['id']} ended with error: {task_record['errors'][0]}"
                    },
                    "ready_evidence": {
                        "ai": [json.dumps(task_record, sort_keys=True, separators=(",", ":"))]
                    },
                }
        if items:
            return {
                "acknowledgements": acknowledgements,
                "failures": [],
                "ready_status": "succeeded",
            }
        return {"acknowledgements": [], "failures": [], "ready_status": "succeeded"}

    monkeypatch.setattr(runtime, "_stream_target", target_write)
    config = {"state_root": str(tmp_path), "data_root": str(tmp_path)}
    first = runtime._process_articles(
        "fixture", state, config, "unused-hub", preview=False, article_limit=1
    )
    assert first["target"]["china"]["acknowledged"] == 1
    assert first["target"]["china"]["failed"] == 0
    first_ack = runtime._ack_stage(first["target"]["china"])
    assert first_ack["status"] == "succeeded"
    assert first_ack["failed"] == 0
    assert first_ack["ready_status"] == "failed"
    assert calls[0][1][0]["source_site"] == "iaa"

    flow = runtime._load_flow_module()
    receipt_root = tmp_path / "receipts"

    def save_receipt(summary):
        stages = {
            stage: runtime._receipt_stage("succeeded", count=0) for stage in flow.COUNT_STAGES
        }
        stages.update(
            {
                "china_ack": runtime._ack_stage(summary["target"]["china"]),
                "overseas_ack": runtime._ack_stage(summary["target"]["overseas"]),
            }
        )
        return runtime._write_receipt(receipt_root, "2026-10-26", "iaa", stages, [], reason=None)

    first_receipt = save_receipt(first)
    assert first_receipt["run_status"] == "partial"
    assert first_receipt["stages"]["china_ack"]["status"] == "succeeded"
    assert first_receipt["stages"]["china_ack"]["failed"] == 0
    assert first_receipt["stages"]["china_ack"]["ready_status"] == "failed"

    retried = runtime._process_articles(
        "fixture", state, config, "unused-hub", preview=False, article_limit=1
    )
    retry_ack = runtime._ack_stage(retried["target"]["china"])
    assert retry_ack["status"] == "succeeded"
    assert retry_ack["acknowledged"] == 1
    assert retry_ack["pending"] == 0
    assert retry_ack["failed"] == 0
    assert retry_ack["ready_status"] == "failed"
    assert state["target_acks"][f"china:71:{item['raw_sha256']}"]["ready_pending"] is True
    assert calls[2][1] == [] and calls[2][2] is True

    second_receipt = save_receipt(retried)
    assert second_receipt["attempt_count"] == 2
    assert second_receipt["stages"]["china_ack"]["failed"] == 0
    assert second_receipt["stages"]["china_ack"]["ready_status"] == "failed"
    assert "china-ai-ready-2" in second_receipt["stages"]["china_ack"]["ready_error"]
    assert (
        '"log_file":"/logs/china-ai-2.log"'
        in second_receipt["stages"]["china_ack"]["ready_evidence"][0]
    )

    report_process = subprocess.run(
        [
            sys.executable,
            str(runtime.FLOW_SCRIPT),
            "report",
            "--date",
            "2026-10-26",
            "--receipts",
            str(receipt_root),
            "--json",
        ],
        check=False,
        capture_output=True,
        text=True,
    )
    assert report_process.returncode == 0, report_process.stderr
    failure_case = next(
        case
        for case in json.loads(report_process.stdout)["failure_cases"]
        if case["site_id"] == "iaa" and case["stage"] == "china_ready"
    )
    assert "china-ai-ready-2" in failure_case["error"]
    assert any('"log_file":"/logs/china-ai-2.log"' in row for row in failure_case["evidence"])

    plan_process = subprocess.run(
        [
            sys.executable,
            str(runtime.FLOW_SCRIPT),
            "plan",
            "--date",
            "2026-10-27",
            "--receipts",
            str(receipt_root),
            "--json",
        ],
        check=False,
        capture_output=True,
        text=True,
    )
    assert plan_process.returncode == 0, plan_process.stderr
    assert "iaa" in json.loads(plan_process.stdout)["extra_retry_sites"]

    ready_recovered = runtime._process_articles(
        "fixture", state, config, "unused-hub", preview=False, article_limit=1
    )
    recovered_ack = runtime._ack_stage(ready_recovered["target"]["china"])
    assert recovered_ack["status"] == "succeeded"
    assert recovered_ack["acknowledged"] == 1
    assert recovered_ack["ready_status"] == "succeeded"
    assert state["target_acks"][f"china:71:{item['raw_sha256']}"]["ready_pending"] is False
    assert article["stage"] == "complete"


def test_hub_version_count_survives_enrichment_failure_without_retry_recount(
    tmp_path: Path, monkeypatch
) -> None:
    state_path = tmp_path / "state.json"
    monkeypatch.setattr(runtime, "_ACTIVE_STATE_PATH", state_path)
    raw_sha = hashlib.sha256(b"captured bytes").hexdigest()
    article_key = runtime._file_key("https://example.test/paper.pdf", raw_sha)
    article = {
        "url": "https://example.test/paper.pdf",
        "raw_sha256": raw_sha,
        "stage": "downloaded",
        "discovery_sites": ["cas"],
    }
    state = {"articles": {article_key: article}, "target_acks": {}}
    full = {
        "url": article["url"],
        "raw_sha256": raw_sha,
        "content_version_id": 19,
        "markdown_text_variant_id": 29,
        "markdown_sha256": hashlib.sha256(b"markdown").hexdigest(),
        "ai_candidate": False,
    }
    attempts = 0

    def prepare(_article, _state, _config, _hub, _preview, *, on_hub_version=None):
        nonlocal attempts
        attempts += 1
        on_hub_version(
            attempts == 1,
            {
                "content_version_id": 19,
                "markdown_text_variant_id": 29,
                "markdown_sha256": full["markdown_sha256"],
            },
        )
        if attempts == 1:
            raise RuntimeError("model unavailable after Hub commit")
        return full, False, True

    monkeypatch.setattr(runtime, "_prepare_hub_article", prepare)
    config = {"state_root": str(tmp_path), "data_root": str(tmp_path)}
    first = runtime._process_articles(
        "cas", state, config, "unused-hub", preview=False, article_limit=1
    )
    assert first["new_versions"] == 1
    assert first["duplicates"] == 0
    assert first["article_errors"]
    assert article["hub_version"]["status"] == "new"
    assert state_path.is_file()

    second = runtime._process_articles(
        "cas", state, config, "unused-hub", preview=False, article_limit=1
    )
    assert second["new_versions"] == 0
    assert second["duplicates"] == 1
    assert second["enriched"] == 1
    assert article["stage"] == "complete"

    receipt_root = tmp_path / "receipts"
    first_receipt = runtime._write_receipt(
        receipt_root,
        "2026-10-25",
        "cas",
        {
            "new_versions": runtime._receipt_stage(
                "partial", count=first["new_versions"], error="model unavailable"
            ),
            "duplicates": runtime._receipt_stage("succeeded", count=first["duplicates"]),
        },
        [],
        "model unavailable",
    )
    second_receipt = runtime._write_receipt(
        receipt_root,
        "2026-10-25",
        "cas",
        {
            "new_versions": runtime._receipt_stage("succeeded", count=second["new_versions"]),
            "duplicates": runtime._receipt_stage("succeeded", count=second["duplicates"]),
        },
        [],
        None,
    )
    assert first_receipt["stages"]["new_versions"]["count"] == 1
    assert second_receipt["stages"]["new_versions"]["count"] == 1
    assert second_receipt["attempt_count"] == 2


@pytest.mark.parametrize(
    ("failed_target", "expected_run_status", "expected_failure_stage", "expected_reason"),
    [
        (
            {
                "acknowledged": 0,
                "pending": 0,
                "failed": 1,
                "ready_status": "succeeded",
            },
            "failed",
            "china_ack",
            "target item acknowledgement failed",
        ),
        (
            {
                "acknowledged": 1,
                "pending": 0,
                "failed": 0,
                "ready_status": "failed",
                "ready_error": "native index build failed",
            },
            "partial",
            "china_ready",
            "native index build failed",
        ),
    ],
)
def test_target_only_failure_receipt_is_durable_and_readable_offline(
    tmp_path: Path,
    failed_target: dict[str, Any],
    expected_run_status: str,
    expected_failure_stage: str,
    expected_reason: str,
) -> None:
    receipt_root = tmp_path / "receipts"
    stages = {
        stage: runtime._receipt_stage("succeeded", count=0)
        for stage in runtime._load_flow_module().COUNT_STAGES
    }
    stages.update(
        {
            "china_ack": runtime._ack_stage(failed_target),
            "overseas_ack": runtime._ack_stage(
                {
                    "acknowledged": 1,
                    "pending": 0,
                    "failed": 0,
                    "ready_status": "succeeded",
                }
            ),
        }
    )

    receipt = runtime._write_receipt(receipt_root, "2026-10-25", "cas", stages, [], reason=None)

    saved_path = receipt_root / "2026-10-25" / "cas.json"
    saved = json.loads(saved_path.read_text(encoding="utf-8"))
    assert saved["schema_version"] == 2
    assert saved["run_status"] == expected_run_status
    assert saved["reason"] == expected_reason
    assert receipt == runtime._load_flow_module().normalize_receipt(saved, saved_path)

    report_process = subprocess.run(
        [
            sys.executable,
            str(runtime.FLOW_SCRIPT),
            "report",
            "--date",
            "2026-10-25",
            "--receipts",
            str(receipt_root),
            "--json",
        ],
        check=False,
        capture_output=True,
        text=True,
    )
    assert report_process.returncode == 0, report_process.stderr
    report = json.loads(report_process.stdout)
    assert "cas" in report["coverage"]["actual_collection_receipt_sites"]
    assert any(
        case["site_id"] == "cas" and case["stage"] == expected_failure_stage
        for case in report["failure_cases"]
    )

    plan_process = subprocess.run(
        [
            sys.executable,
            str(runtime.FLOW_SCRIPT),
            "plan",
            "--date",
            "2026-10-26",
            "--receipts",
            str(receipt_root),
            "--json",
        ],
        check=False,
        capture_output=True,
        text=True,
    )
    assert plan_process.returncode == 0, plan_process.stderr
    plan = json.loads(plan_process.stdout)
    assert "cas" in plan["extra_retry_sites"]


def test_target_item_failures_reach_summary_receipt_and_offline_case(
    tmp_path: Path, monkeypatch
) -> None:
    monkeypatch.setattr(runtime, "_ACTIVE_STATE_PATH", tmp_path / "runtime-state.json")
    url = "https://example.test/a.pdf"
    raw = b"verified target-failure article"
    markdown = "Pricing analysis uses AI methods."
    raw_sha = hashlib.sha256(raw).hexdigest()
    markdown_sha = hashlib.sha256(markdown.encode("utf-8")).hexdigest()
    item = {
        "url": url,
        "source_site": "cas",
        "raw_sha256": raw_sha,
        "content_version_id": 91,
        "markdown_text_variant_id": 101,
        "markdown_sha256": markdown_sha,
        "markdown": markdown,
        "bilingual_metadata": {
            language: {
                "title": "Pricing and AI",
                "summary": "Pricing analysis uses AI methods.",
                "category": "Pricing",
                "keywords": ["pricing", "AI"],
            }
            for language in ("zh", "en")
        },
        "categories": [
            {"key": "Pricing", "evidence": ["Pricing analysis"]},
            {"key": "AI", "evidence": ["AI methods"]},
        ],
        "date_facts": {
            "publication_date": None,
            "revision_date": None,
            "document_version_date": None,
        },
        "date_review_status": "complete",
        "date_review_reason": None,
        "date_review_unresolved": [],
        "source_signature": "a" * 64,
        "result_sha256": "b" * 64,
        "ai_candidate": True,
    }
    article_key = runtime._file_key(url, raw_sha)
    article = {
        **item,
        "site_id": "cas",
        "artifact_path": str(tmp_path / "article.pdf"),
        "stage": "downloaded",
        "discovery_sites": ["cas"],
    }
    Path(article["artifact_path"]).write_bytes(raw)
    state = {"articles": {article_key: article}, "target_acks": {}}
    monkeypatch.setattr(
        runtime, "_prepare_hub_article", lambda *_args, **_kwargs: (item, False, False)
    )
    failure = {
        "url": url,
        "error": "native Markdown upsert failed: disk full",
    }

    def target_reply(_config, target, items, _ready_pending, **_kwargs):
        if target == "china":
            return {
                "acknowledgements": [],
                "failures": [failure],
                "ready_status": "succeeded",
                "ready_error": None,
            }
        return {
            "acknowledgements": [
                {
                    "url": items[0]["url"],
                    "content_version_id": items[0]["content_version_id"],
                    "raw_sha256": items[0]["raw_sha256"],
                    "metadata_fingerprint": runtime._target_metadata_fingerprint(items[0]),
                    "status": "written",
                }
            ],
            "failures": [],
            "ready_status": "succeeded",
            "ready_error": None,
        }

    monkeypatch.setattr(runtime, "_stream_target", target_reply)
    summary = runtime._process_articles(
        "cas",
        state,
        {"data_root": str(tmp_path), "state_root": str(tmp_path)},
        "unused-hub",
        preview=False,
        article_limit=1,
    )
    target_summary = summary["target"]["china"]
    assert target_summary["failures"] == [failure]
    assert target_summary["failed"] == 1
    ack = runtime._ack_stage(target_summary)
    expected_evidence = f"{url}: {failure['error']}"
    assert expected_evidence in ack["error"]
    assert ack["evidence"] == [expected_evidence]

    flow = runtime._load_flow_module()
    stages = {stage: runtime._receipt_stage("succeeded", count=0) for stage in flow.COUNT_STAGES}
    stages.update(
        {
            "china_ack": ack,
            "overseas_ack": runtime._ack_stage(summary["target"]["overseas"]),
        }
    )
    receipt_root = tmp_path / "receipts"
    receipt = runtime._write_receipt(receipt_root, "2026-10-25", "cas", stages, [], reason=None)
    assert receipt["run_status"] == "failed"
    assert receipt["stages"]["china_ack"]["evidence"] == [expected_evidence]

    report_process = subprocess.run(
        [
            sys.executable,
            str(runtime.FLOW_SCRIPT),
            "report",
            "--date",
            "2026-10-25",
            "--receipts",
            str(receipt_root),
            "--json",
        ],
        check=False,
        capture_output=True,
        text=True,
    )
    assert report_process.returncode == 0, report_process.stderr
    report = json.loads(report_process.stdout)
    case = next(
        row
        for row in report["failure_cases"]
        if row["site_id"] == "cas" and row["stage"] == "china_ack"
    )
    assert expected_evidence in case["error"]
    assert case["evidence"] == [expected_evidence]


def test_site_receipt_download_count_uses_capture_results_not_retry_queue(
    tmp_path: Path, monkeypatch
) -> None:
    captured = {
        "downloads": [],
        "errors": [],
        "request_count": 1,
        "frontier_remaining": 0,
        "budget_exhausted": False,
    }
    monkeypatch.setattr(runtime, "_new_crawler", lambda _root: object())
    monkeypatch.setattr(runtime, "_capture", lambda *_args, **_kwargs: captured)
    monkeypatch.setattr(
        runtime,
        "_process_articles",
        lambda *_args, **_kwargs: {
            "candidate_count": 1,
            "pending_article_count": 0,
            "new_versions": 0,
            "duplicates": 1,
            "enriched": 0,
            "ai_candidates": 0,
            "article_errors": [],
            "target": {
                name: {
                    "acknowledged": 0,
                    "pending": 0,
                    "failed": 0,
                    "ready_status": "succeeded",
                }
                for name in ("china", "overseas")
            },
        },
    )

    def save_receipt(_root, _date, _site, stages, _evidence, reason):
        return {"stages": stages, "reason": reason}

    monkeypatch.setattr(runtime, "_write_receipt", save_receipt)
    config = {"state_root": str(tmp_path), "data_root": str(tmp_path)}
    state = {"sites": {}, "articles": {}, "target_acks": {}}

    result = runtime._run_site("cas", "2026-10-25", config, state, "unused-hub")
    assert result["stages"]["downloaded"]["count"] == 0
    assert result["stages"]["downloaded"]["status"] == "succeeded"

    captured["downloads"] = [{"url": "https://example.test/one.pdf"}] * 2
    result = runtime._run_site("cas", "2026-10-25", config, state, "unused-hub")
    assert result["stages"]["downloaded"]["count"] == 2


def test_cli_help_and_plan_json_contract(tmp_path: Path, monkeypatch, capsys) -> None:
    help_result = subprocess.run(
        [sys.executable, "-m", "ai_actuarial.local_collection_runtime", "--help"],
        capture_output=True,
        text=True,
        check=True,
    )
    assert "apply-hub-schema" in help_result.stdout
    assert "reconcile-targets" in help_result.stdout
    run_site_help = subprocess.run(
        [sys.executable, "-m", "ai_actuarial.local_collection_runtime", "run-site", "--help"],
        capture_output=True,
        text=True,
        check=True,
    )
    assert "--request-limit" in run_site_help.stdout
    assert "Switzerland" in run_site_help.stdout
    monkeypatch.setattr(
        runtime,
        "_load_config",
        lambda _path: {"state_root": str(tmp_path), "data_root": str(tmp_path)},
    )
    assert (
        runtime.main(["plan", "--date", "2026-10-25", "--receipts", str(tmp_path), "--json"]) == 0
    )
    output = json.loads(capsys.readouterr().out)
    assert output["date_local"] == "2026-10-25"
    assert output["baseline_size"] == 6
    assert len(output["items"]) == output["baseline_size"]


@pytest.mark.parametrize(
    ("stage_status", "expected_code"), [("failed", 2), ("partial", 3), ("succeeded", 0)]
)
def test_run_site_preview_exit_codes_follow_observed_stage_status(
    tmp_path: Path, monkeypatch, capsys, stage_status: str, expected_code: int
) -> None:
    from contextlib import nullcontext

    monkeypatch.setattr(
        runtime,
        "_load_config",
        lambda _path: {
            "state_root": str(tmp_path),
            "data_root": str(tmp_path),
            "hub_connection_file": str(tmp_path / "hub.json"),
        },
    )
    monkeypatch.setattr(runtime, "_state_path", lambda _config: tmp_path / "state.json")
    monkeypatch.setattr(runtime, "_exclusive_lock", lambda _path: nullcontext())
    monkeypatch.setattr(runtime, "_load_state", lambda _path: {"sites": {}, "articles": {}})
    monkeypatch.setattr(runtime, "_connection_url", lambda _path: "unused-hub")
    monkeypatch.setattr(
        runtime,
        "_run_site",
        lambda *_args, **_kwargs: {
            "preview": True,
            "stages": {
                "downloaded": {"status": stage_status},
                "china_ack": {"status": "succeeded", "ready_status": "unknown"},
                "overseas_ack": {"status": "succeeded", "ready_status": "unknown"},
            },
        },
    )
    code = runtime.main(
        [
            "run-site",
            "--site",
            "iaa",
            "--config",
            str(tmp_path / "config.json"),
            "--preview",
            "--json",
        ]
    )
    output = json.loads(capsys.readouterr().out)
    assert code == expected_code
    assert output["preview"] is True


@pytest.mark.parametrize(
    ("site_id", "page_url", "listing", "expected_url", "rejected_url"),
    [
        (
            "cia-research-publications",
            "https://www.cia-ica.ca/publications/226099e/",
            b'<a href="mailto:research@example.ca">Email</a><a href="tel:+14165551212">Call</a>'
            b'<a href="javascript:void(0)">Menu</a><a href="/app/themes/wicket/custom/dl_file.php?p=840374&amp;fid=840377">Download publication</a>'
            b'<a href="/private-files/226099e.pdf">Download publication</a>',
            "https://www.cia-ica.ca/app/themes/wicket/custom/dl_file.php?fid=840377&p=840374",
            "https://www.cia-ica.ca/private-files/226099e.pdf",
        ),
        (
            "cca",
            "https://www.ccactuaries.org/professional-resources/papers",
            b'<a href="mailto:info@ccactuaries.org">Email</a><a href="tel:555">Call</a>'
            b'<a href="javascript:openMenu()">Menu</a><a href="/docs/default-source/papers/evaluating-ai.pdf">AI actuarial research paper</a>'
            b'<a href="https://my.ccactuaries.org/Security/Sign-In?returnurl=%2Fpapers">Sign In</a>',
            "https://www.ccactuaries.org/docs/default-source/papers/evaluating-ai.pdf",
            "https://my.ccactuaries.org/Security/Sign-In?returnurl=%2Fpapers",
        ),
        (
            "france",
            "https://www.institutdesactuaires.com/article/conferences-1105",
            b'<a href="mailto:info@institutdesactuaires.com">Email</a><a href="tel:555">Call</a>'
            b'<a href="javascript:void(0)">Menu</a><a href="/global/gene/link.php?doc_id=20148&amp;fg=1">Presentation AI 2026</a>'
            b'<a href="/global/gene/link.php?doc_id=1&amp;fg=1">Old unrelated link</a>',
            "https://www.institutdesactuaires.com/global/gene/link.php?doc_id=20148&fg=1",
            "https://www.institutdesactuaires.com/global/gene/link.php?doc_id=1&fg=1",
        ),
    ],
)
def test_observed_source_links_skip_non_http_and_queue_only_public_documents(
    site_id: str,
    page_url: str,
    listing: bytes,
    expected_url: str,
    rejected_url: str,
    tmp_path: Path,
    monkeypatch,
) -> None:
    monkeypatch.setattr(runtime, "_ACTIVE_STATE_PATH", tmp_path / "state.json")

    class OneRequestCrawler(_FixtureCrawler):
        def begin_site(self, _budget: int) -> None:
            self.site_attempts = 0
            self.site_budget = 1

    site = runtime.SITES[site_id]
    site_state = {
        "queue": [{"url": page_url, "kind": "page", "depth": 0, "referrer": page_url}],
        "cursor": 0,
        "visited": [],
    }
    state = {"articles": {}, "sites": {site_id: site_state}}
    result = runtime._capture(
        site, state, OneRequestCrawler(listing, []), tmp_path / "data", site_state
    )

    urls = {row["url"] for row in site_state["queue"]}
    assert result["request_count"] == 1
    assert expected_url in urls
    assert rejected_url not in urls
    assert not any(
        urlsplit_scheme.startswith(("mailto", "tel", "javascript"))
        for urlsplit_scheme in (row["url"].split(":", 1)[0].lower() for row in site_state["queue"])
    )
    if site_id == "cia-research-publications":
        assert runtime._looks_site_document(expected_url, "Download publication", site)
        assert not runtime._looks_site_document(rejected_url, "Download publication", site)
    if site_id == "france":
        assert runtime._relevant_page(
            "https://www.institutdesactuaires.com/article/conferences-1105", "Conférences", site
        )
        assert not runtime._looks_site_document(rejected_url, "Old unrelated link", site)


@pytest.mark.parametrize(
    ("site_id", "obsolete_url", "expected_urls"),
    [
        (
            "cca",
            "https://www.ccactuaries.org/papers",
            (
                "https://www.ccactuaries.org/professional-resources/papers",
                "https://www.ccactuaries.org/professional-resources/technical-resource-library-main",
            ),
        ),
        (
            "france",
            "https://www.institutdesactuaires.com/global/gene/link.php?doc_id=1&fg=1",
            (
                "https://www.institutdesactuaires.com/se-documenter/memoires/memoires-d-actuariat-4651",
                "https://www.institutdesactuaires.com/se-documenter/supports-des-presentations-39",
            ),
        ),
    ],
)
def test_legacy_site_frontier_repair_is_idempotent_and_keeps_cursor_and_articles(
    site_id: str, obsolete_url: str, expected_urls: tuple[str, ...]
) -> None:
    site = runtime.SITES[site_id]
    consumed = "https://www.institutdesactuaires.com/" if site_id == "france" else site.seeds[0]
    pending = f"https://{site.hosts[0]}/research/other-valid-item"
    state = {
        "articles": {"keep": {"stage": "complete", "url": "https://example.test/a.pdf"}},
        "sites": {
            site_id: {
                "queue": [
                    {"url": consumed, "kind": "page"},
                    {"url": obsolete_url, "kind": "page"},
                    {"url": pending, "kind": "page"},
                ],
                "cursor": 1,
                "visited": [runtime._queue_identity({"url": consumed, "kind": "page"})],
                "retry_queue": [
                    {"url": obsolete_url, "kind": "page", "attempts": 2},
                    {"url": "javascript:void(0)", "kind": "page", "attempts": 1},
                ],
            }
        },
    }
    site_state = state["sites"][site_id]

    runtime._repair_saved_frontier(site, site_state)
    after_first = json.loads(json.dumps(site_state))
    runtime._repair_saved_frontier(site, site_state)

    queue_urls = [row["url"] for row in site_state["queue"]]
    assert site_state == after_first
    assert site_state["cursor"] == 1
    assert queue_urls[0] == consumed
    assert queue_urls[1:3] == list(expected_urls)
    assert queue_urls[-1] == pending
    assert obsolete_url not in queue_urls
    assert all(
        urlsplit_scheme in {"http", "https"}
        for urlsplit_scheme in (
            row["url"].split(":", 1)[0].lower() for row in site_state["retry_queue"]
        )
    )
    assert state["articles"]["keep"]["stage"] == "complete"


def test_cia_legacy_frontier_adds_verified_publications_page_without_resetting_rest_cursor() -> (
    None
):
    site = runtime.SITES["cia-research-publications"]
    rest = runtime._initial_frontier(site)[1]
    later = {
        "url": "https://www.cia-ica.ca/wp-json/wp/v2/publication?page=60",
        "kind": "json_rest",
        "page": 60,
    }
    site_state = {"queue": [rest, later], "cursor": 1, "visited": [runtime._queue_identity(rest)]}

    runtime._repair_saved_frontier(site, site_state)
    after_first = json.loads(json.dumps(site_state))
    runtime._repair_saved_frontier(site, site_state)

    assert site_state == after_first
    assert site_state["cursor"] == 1
    assert site_state["queue"][0] == rest
    assert site_state["queue"][1]["url"] == "https://www.cia-ica.ca/publications/"
    assert site_state["queue"][2] == later


def test_swiss_sixty_request_budget_stops_at_limit_and_resumes_at_saved_cursor(
    tmp_path: Path, monkeypatch
) -> None:
    monkeypatch.setattr(runtime, "_ACTIVE_STATE_PATH", tmp_path / "state.json")
    monkeypatch.setattr(capacity, "ensure_capacity", lambda **_kwargs: None)
    site = runtime.SITES["swiss-association-of-actuaries"]
    listing = "".join(
        f'<a href="/docs/research-{number}.pdf">AI actuarial research {number}</a>'
        for number in range(61)
    ).encode()

    class SwissBudgetCrawler:
        def __init__(self, html_body: bytes) -> None:
            self.listing = html_body
            self.site_attempts = 0
            self.site_budget = 30
            self.requests: list[str] = []

        def begin_site(self, budget: int) -> None:
            self.site_attempts = 0
            self.site_budget = budget

        def _request(self, url: str, **_kwargs):
            self.site_attempts += 1
            self.requests.append(url)
            return self.listing, {"content-type": "text/html"}, url

        def _download_file(self, url: str, staging_dir: Path, **_kwargs):
            self.site_attempts += 1
            self.requests.append(url)
            raw = b"%PDF-1.7\nSwiss fixture article"
            staging_dir.mkdir(parents=True, exist_ok=True)
            path = staging_dir / f"{hashlib.sha256(url.encode()).hexdigest()}.pdf"
            path.write_bytes(raw)
            return (
                path,
                {"content-type": "application/pdf"},
                url,
                hashlib.sha256(raw).hexdigest(),
                len(raw),
            )

    page = "https://www.actuaries.ch/en/download"
    site_state = {
        "queue": [{"url": page, "kind": "page", "depth": 0, "referrer": page}],
        "cursor": 0,
        "visited": [],
    }
    state = {"articles": {}, "sites": {site.site_id: site_state}}
    first_crawler = SwissBudgetCrawler(listing)

    first = runtime._capture(
        site, state, first_crawler, tmp_path / "data", site_state, request_limit=60
    )

    assert first["request_count"] == 60
    assert len(first_crawler.requests) == 60
    assert first["budget_exhausted"] is True
    assert site_state["cursor"] == 60
    assert first["frontier_remaining"] == 2
    second_crawler = SwissBudgetCrawler(b"")
    second = runtime._capture(
        site, state, second_crawler, tmp_path / "data", site_state, request_limit=60
    )
    assert second["request_count"] == 2
    assert second_crawler.requests == [row["url"] for row in site_state["queue"][60:62]]
    assert site_state["completed_frontier"] is True
    with pytest.raises(runtime.RuntimeErrorSafe, match="between 1 and 30"):
        runtime._capture(
            runtime.SITES["cca"],
            {"articles": {}},
            SwissBudgetCrawler(b""),
            tmp_path / "data",
            {},
            request_limit=31,
        )


def test_budget_receipt_stays_partial_without_researching_source_and_reports_pending_articles(
    tmp_path: Path,
) -> None:
    flow = runtime._load_flow_module()
    receipt_root = tmp_path / "receipts"

    def stages(download_error: str | None = None):
        result = {
            name: runtime._receipt_stage("succeeded", count=(29 if name == "downloaded" else 0))
            for name in flow.COUNT_STAGES
        }
        result["downloaded"] = runtime._receipt_stage(
            "partial",
            count=29,
            error=download_error,
            evidence=["HTTP requests=30/30; frontier_remaining=5"],
            request_budget={"used": 30, "cap": 30, "paused": True, "frontier_remaining": 5},
        )
        result["new_versions"] = runtime._receipt_stage("succeeded", count=3, pending_count=17)
        result["enriched"] = runtime._receipt_stage("succeeded", count=1, pending_count=17)
        for target in ("china", "overseas"):
            result[f"{target}_ack"] = runtime._ack_stage(
                {"acknowledged": 1, "pending": 0, "failed": 0, "ready_status": "succeeded"}
            )
        return result

    pure_budget = runtime._write_receipt(
        receipt_root,
        "2026-10-27",
        "cas",
        stages(),
        ["request budget reached; cursor retained"],
        "site reached its bounded HTTP request budget; discovery cursor was retained",
    )
    assert pure_budget["run_status"] == "partial"
    assert pure_budget["stages"]["downloaded"]["error"] is None
    assert pure_budget["stages"]["downloaded"]["request_budget"]["used"] == 30
    assert pure_budget["stages"]["enriched"]["pending_count"] == 17
    receipts = flow.read_receipts(receipt_root)
    report = flow.make_report(__import__("datetime").date(2026, 10, 27), receipts)
    assert report["failure_cases"] == []
    assert report["research_queue"] == []
    cas = next(row for row in report["site_statuses"] if row["site_id"] == "cas")
    assert cas["state"] == "partial"
    assert cas["pending_article_count"] == 17
    assert cas["request_budget"] == {
        "used": 30,
        "cap": 30,
        "paused": True,
        "frontier_remaining": 5,
    }

    error_receipt = runtime._write_receipt(
        receipt_root,
        "2026-10-27",
        "cas",
        stages("https://www.casact.org: TimeoutError: request timed out"),
        [],
        "https://www.casact.org: TimeoutError: request timed out",
    )
    assert error_receipt["attempt_count"] == 2
    assert "TimeoutError" in error_receipt["stages"]["downloaded"]["error"]
    report = flow.make_report(
        __import__("datetime").date(2026, 10, 27), flow.read_receipts(receipt_root)
    )
    assert any(case["site_id"] == "cas" for case in report["failure_cases"])


def test_site_receipt_keeps_source_error_when_request_budget_is_also_exhausted(
    tmp_path: Path, monkeypatch
) -> None:
    captured = {
        "downloads": [],
        "errors": ["https://www.casact.org/research/: TimeoutError: connection reset"],
        "request_count": 30,
        "request_limit": 30,
        "frontier_remaining": 5,
        "budget_exhausted": True,
    }
    monkeypatch.setattr(runtime, "_new_crawler", lambda _root: object())
    monkeypatch.setattr(runtime, "_capture", lambda *_args, **_kwargs: captured)
    monkeypatch.setattr(
        runtime,
        "_process_articles",
        lambda *_args, **_kwargs: {
            "candidate_count": 0,
            "pending_article_count": 12,
            "new_versions": 0,
            "duplicates": 0,
            "enriched": 0,
            "ai_candidates": 0,
            "article_errors": [],
            "target": {
                name: {
                    "acknowledged": 0,
                    "pending": 0,
                    "failed": 0,
                    "ready_status": "succeeded",
                }
                for name in ("china", "overseas")
            },
        },
    )
    config = {"state_root": str(tmp_path), "data_root": str(tmp_path)}
    state = {"sites": {}, "articles": {}, "target_acks": {}}

    receipt = runtime._run_site("cas", "2026-10-27", config, state, "unused-hub")
    downloaded = receipt["stages"]["downloaded"]
    assert receipt["run_status"] == "failed"
    assert "TimeoutError: connection reset" in downloaded["error"]
    assert "budget reached" in downloaded["error"]
    assert downloaded["request_budget"] == {
        "used": 30,
        "cap": 30,
        "paused": True,
        "frontier_remaining": 5,
    }
    assert receipt["stages"]["enriched"]["pending_count"] == 12

    report = runtime._load_flow_module().make_report(
        __import__("datetime").date(2026, 10, 27),
        runtime._load_flow_module().read_receipts(tmp_path / "receipts"),
    )
    source_case = next(
        case
        for case in report["failure_cases"]
        if case["site_id"] == "cas" and case["stage"] == "downloaded"
    )
    assert "TimeoutError: connection reset" in source_case["error"]
    assert "budget reached" in source_case["error"]
    cas = next(row for row in report["site_statuses"] if row["site_id"] == "cas")
    assert cas["pending_article_count"] == 12


def test_article_model_limit_is_separate_from_bounded_reuse_and_target_retries(
    tmp_path: Path, monkeypatch
) -> None:
    monkeypatch.setattr(runtime, "_ACTIVE_STATE_PATH", tmp_path / "state.json")
    site_id = "swiss-association-of-actuaries"
    all_items: dict[str, dict[str, Any]] = {}
    state_articles: dict[str, dict[str, Any]] = {}
    for number in range(runtime.MAX_ARTICLE_SCAN_PER_SITE_RUN + 1):
        raw_sha = hashlib.sha256(f"raw-{number}".encode()).hexdigest()
        markdown = f"AI source evidence {number}"
        categories = [{"key": "AI", "evidence": [markdown]}]
        item = {
            "url": f"https://www.actuaries.ch/docs/research-{number}.pdf",
            "raw_sha256": raw_sha,
            "content_version_id": number + 1,
            "markdown_text_variant_id": number + 100,
            "markdown_sha256": hashlib.sha256(markdown.encode()).hexdigest(),
            "markdown": markdown,
            "source_site": site_id,
            "source_page_url": "https://www.actuaries.ch/en/download",
            "title": f"AI actuarial research {number}",
            "bilingual_metadata": {
                "zh": {"title": "AI", "summary": "摘要", "category": "AI", "keywords": ["AI"]},
                "en": {"title": "AI", "summary": "summary", "category": "AI", "keywords": ["AI"]},
            },
            "categories": categories,
            "annotation": {"categories": categories},
            "date_facts": {},
            "date_review_status": "reviewed",
            "date_review_unresolved": [],
            "ai_candidate": True,
        }
        all_items[item["url"]] = item
        stage = "hub_verified" if number < runtime.MAX_ARTICLE_SCAN_PER_SITE_RUN else "downloaded"
        state_articles[runtime._file_key(item["url"], raw_sha)] = {
            **item,
            "stage": stage,
            "enrichment": item if stage == "hub_verified" else None,
            "discovery_sites": [site_id],
            "artifact_path": str(tmp_path / f"raw-{number}.pdf"),
        }

    state = {"articles": state_articles, "target_acks": {}, "sites": {}}
    candidate_url = (
        f"https://www.actuaries.ch/docs/research-{runtime.MAX_ARTICLE_SCAN_PER_SITE_RUN}.pdf"
    )
    monkeypatch.setattr(
        runtime,
        "content_hub_read_version",
        lambda url, _sha, _hub: None if url == candidate_url else all_items.get(url),
    )
    prepared: list[int] = []

    def prepare(article, *_args, **_kwargs):
        item = all_items[article["url"]]
        prepared.append(item["content_version_id"])
        return (
            item,
            item["content_version_id"] == runtime.MAX_ARTICLE_SCAN_PER_SITE_RUN + 1,
            item["content_version_id"] == runtime.MAX_ARTICLE_SCAN_PER_SITE_RUN + 1,
        )

    monkeypatch.setattr(runtime, "_prepare_hub_article", prepare)

    def target_write(_config, _target, _items, _ready_pending, **_kwargs):
        raise RuntimeError("bounded fixture leaves target retry pending")

    monkeypatch.setattr(runtime, "_stream_target", target_write)
    first = runtime._process_articles(
        site_id,
        state,
        {"state_root": str(tmp_path), "data_root": str(tmp_path)},
        "unused-hub",
        preview=False,
        article_limit=1,
    )
    assert first["reuse_candidate_count"] == runtime.MAX_CHEAP_ARTICLES_PER_SITE_RUN
    assert first["enrichment_candidate_count"] == 0
    assert first["pending_article_count"] == 9
    assert state["article_cursors"][site_id] == runtime._article_identity(
        list(state_articles.values())[runtime.MAX_CHEAP_ARTICLES_PER_SITE_RUN - 1]
    )

    prepared.clear()
    second = runtime._process_articles(
        site_id,
        state,
        {"state_root": str(tmp_path), "data_root": str(tmp_path)},
        "unused-hub",
        preview=False,
        article_limit=1,
    )
    assert second["enrichment_candidate_count"] == 1
    assert second["new_versions"] == 1
    assert runtime.MAX_ARTICLE_SCAN_PER_SITE_RUN + 1 in prepared
    assert second["pending_article_count"] == 8


@pytest.mark.parametrize("mixed_model_and_reuse", [False, True])
def test_retry_cursor_stops_before_first_deferred_scanned_article(
    tmp_path: Path, monkeypatch, mixed_model_and_reuse: bool
) -> None:
    monkeypatch.setattr(runtime, "_ACTIVE_STATE_PATH", tmp_path / "state.json")
    site_id = "swiss-association-of-actuaries"
    records: list[tuple[str, int, dict[str, Any]]] = []
    for number in range(runtime.MAX_ARTICLE_SCAN_PER_SITE_RUN):
        raw_sha = hashlib.sha256(f"retry-cursor-raw-{number}".encode()).hexdigest()
        markdown = f"AI source evidence for retry record {number}"
        categories = [{"key": "AI", "evidence": [markdown]}]
        url = f"https://www.actuaries.ch/docs/retry-{number:02d}.pdf"
        item = {
            "url": url,
            "raw_sha256": raw_sha,
            "content_version_id": number + 100,
            "markdown_text_variant_id": number + 200,
            "markdown_sha256": hashlib.sha256(markdown.encode()).hexdigest(),
            "markdown": markdown,
            "source_site": site_id,
            "source_page_url": "https://www.actuaries.ch/en/download",
            "title": f"AI actuarial research {number}",
            "bilingual_metadata": {
                "zh": {"title": "AI", "summary": "摘要", "category": "AI", "keywords": ["AI"]},
                "en": {
                    "title": "AI",
                    "summary": "Summary",
                    "category": "AI",
                    "keywords": ["AI"],
                },
            },
            "categories": categories,
            "annotation": {"categories": categories},
            "date_facts": {},
            "date_review_status": "reviewed",
            "date_review_unresolved": [],
            "ai_candidate": True,
        }
        records.append((runtime._file_key(url, raw_sha), number, item))
    records.sort(key=lambda row: row[0])
    model_urls = {
        item["url"]
        for position, (_, _, item) in enumerate(records)
        if mixed_model_and_reuse and position < 25
    }
    state_articles: dict[str, dict[str, Any]] = {}
    for position, (key, index, item) in enumerate(records):
        reusable = not mixed_model_and_reuse or position >= 25
        state_articles[key] = {
            **item,
            "stage": "hub_verified" if reusable else "downloaded",
            "enrichment": item if reusable else None,
            "discovery_sites": [site_id],
            "artifact_path": str(tmp_path / f"retry-{index}.pdf"),
        }
    record_urls = [item["url"] for _, _, item in records]
    state_path = tmp_path / "state.json"
    state: dict[str, Any] = {
        "schema_version": 1,
        "sites": {},
        "articles": state_articles,
        "target_acks": {},
    }
    runtime._atomic_json(state_path, state)
    monkeypatch.setattr(runtime, "content_hub_read_version", lambda *_args: None)
    model_attempts: list[list[str]] = []
    target_attempts: list[dict[str, list[str]]] = []

    def prepare(article, *_args, **_kwargs):
        if article["url"] in model_urls:
            model_attempts[-1].append(article["url"])
            raise RuntimeError("fixture model timeout")
        return article["enrichment"], False, False

    monkeypatch.setattr(runtime, "_prepare_hub_article", prepare)

    def target_write(_config, target, items, _ready_pending, **_kwargs):
        target_attempts[-1][target] = [item["url"] for item in items]
        raise RuntimeError("fixture target outage")

    monkeypatch.setattr(runtime, "_stream_target", target_write)
    config = {"state_root": str(tmp_path), "data_root": str(tmp_path)}

    def process_persisted_round() -> dict[str, Any]:
        nonlocal state
        model_attempts.append([])
        target_attempts.append({})
        result = runtime._process_articles(
            site_id,
            state,
            config,
            "unused-hub",
            preview=False,
            article_limit=8,
        )
        state = runtime._load_state(state_path)
        return result

    first = process_persisted_round()
    first_selected = first["reuse_candidate_count"] + first["enrichment_candidate_count"]
    assert first_selected == (
        15 if mixed_model_and_reuse else runtime.MAX_CHEAP_ARTICLES_PER_SITE_RUN
    )
    first_deferred_index = 8 if mixed_model_and_reuse else runtime.MAX_CHEAP_ARTICLES_PER_SITE_RUN
    expected_first_anchor = record_urls[first_deferred_index - 1]
    first_cursor_url = next(
        row["url"]
        for row in state["articles"].values()
        if runtime._article_identity(row) == state["article_cursors"][site_id]
    )
    assert first_cursor_url == expected_first_anchor

    second = process_persisted_round()
    second_target_urls = set(target_attempts[1].get("china", []))
    if mixed_model_and_reuse:
        assert set(model_attempts[0]) == set(record_urls[:8])
        expected_model_urls = set(record_urls[8:16])
        assert set(model_attempts[1]) == expected_model_urls
        assert target_attempts[0]["china"] == record_urls[25:32]
        assert target_attempts[1]["china"] == record_urls[25:32]
        assert second["reuse_candidate_count"] == 7
        assert second["enrichment_candidate_count"] == 8
    else:
        assert target_attempts[0]["china"] == record_urls[:24]
        assert len(target_attempts[0]["china"]) == 24
        assert len(target_attempts[1]["china"]) == 24
        assert set(record_urls[24:32]).issubset(second_target_urls)

    if mixed_model_and_reuse:
        third = process_persisted_round()
        assert third["enrichment_candidate_count"] == 8
        assert set(model_attempts[2]) == set(record_urls[16:24])
        assert target_attempts[2]["china"] == record_urls[25:32]


def test_completed_reuse_cursor_survives_restart_and_rotates_to_new_articles(
    tmp_path: Path, monkeypatch
) -> None:
    monkeypatch.setattr(runtime, "_ACTIVE_STATE_PATH", tmp_path / "state.json")
    site_id = "swiss-association-of-actuaries"
    state_articles: dict[str, dict[str, Any]] = {}
    records: dict[str, dict[str, Any]] = {}

    def make_item(kind: str, number: int) -> dict[str, Any]:
        raw_sha = hashlib.sha256(f"{kind}-{number}".encode()).hexdigest()
        markdown = f"AI evidence {kind} {number}"
        categories = [{"key": "AI", "evidence": [markdown]}]
        url = f"https://www.actuaries.ch/docs/{kind}-{number}.pdf"
        item = {
            "url": url,
            "raw_sha256": raw_sha,
            "content_version_id": f"{kind}-{number}",
            "markdown_text_variant_id": f"variant-{kind}-{number}",
            "markdown_sha256": hashlib.sha256(markdown.encode()).hexdigest(),
            "markdown": markdown,
            "source_site": site_id,
            "source_page_url": "https://www.actuaries.ch/en/download",
            "title": f"AI actuarial research {kind} {number}",
            "bilingual_metadata": {
                "zh": {
                    "title": "AI 研究",
                    "summary": "摘要",
                    "category": "AI",
                    "keywords": ["AI"],
                },
                "en": {
                    "title": "AI research",
                    "summary": "Summary",
                    "category": "AI",
                    "keywords": ["AI"],
                },
            },
            "categories": categories,
            "annotation": {"categories": categories},
            "date_facts": {},
            "date_review_status": "reviewed",
            "date_review_unresolved": [],
            "ai_candidate": True,
        }
        records[url] = item
        if kind == "old":
            stage = "pending_retry"
            row = {**item, "stage": stage, "last_error": "previous model timeout"}
        elif kind == "reuse":
            stage = "hub_verified"
            row = {**item, "stage": stage, "enrichment": item}
        else:
            stage = "downloaded"
            row = {**item, "stage": stage}
        row["discovery_sites"] = [site_id]
        row["artifact_path"] = str(tmp_path / f"{kind}-{number}.pdf")
        state_articles[f"{ {'old': '00', 'reuse': '10', 'new': '20'}[kind] }-{number:02d}"] = row
        return item

    for number in range(8):
        make_item("old", number)
    for number in range(24):
        make_item("reuse", number)
    for number in range(8):
        make_item("new", number)

    state_path = tmp_path / "state.json"
    runtime._atomic_json(
        state_path,
        {"schema_version": 1, "sites": {}, "articles": state_articles, "target_acks": {}},
    )
    state = runtime._load_state(state_path)
    monkeypatch.setattr(runtime, "content_hub_read_version", lambda *_args: None)

    attempted_old: list[str] = []
    attempted_new: list[str] = []

    def prepare(article, *_args, **_kwargs):
        item = records[article["url"]]
        if "old-" in article["url"]:
            attempted_old.append(article["url"])
            raise RuntimeError("fixture model timeout")
        if "new-" in article["url"]:
            attempted_new.append(article["url"])
            return item, True, True
        return item, False, False

    monkeypatch.setattr(runtime, "_prepare_hub_article", prepare)

    def target_write(_config, _target, items, _ready_pending, **_kwargs):
        return {
            "acknowledgements": [
                {
                    "content_version_id": item["content_version_id"],
                    "raw_sha256": item["raw_sha256"],
                    "status": "written",
                    "metadata_fingerprint": runtime._target_metadata_fingerprint(item),
                }
                for item in items
            ],
            "ready_status": "succeeded",
            "failures": [],
        }

    monkeypatch.setattr(runtime, "_stream_target", target_write)
    config = {"state_root": str(tmp_path), "data_root": str(tmp_path)}

    first = runtime._process_articles(
        site_id, state, config, "unused-hub", preview=False, article_limit=8
    )
    assert first["reuse_candidate_count"] == 24
    assert first["enrichment_candidate_count"] == 8
    assert len(attempted_old) == 8
    assert len(attempted_new) == 0
    anchor = state["article_cursors"][site_id]
    anchor_article = next(
        row for row in state["articles"].values() if runtime._article_identity(row) == anchor
    )
    assert anchor_article["stage"] == "complete"

    state = runtime._load_state(state_path)
    second = runtime._process_articles(
        site_id, state, config, "unused-hub", preview=False, article_limit=8
    )
    assert second["enrichment_candidate_count"] == 8
    assert second["new_versions"] == 8
    assert len(attempted_new) == 8
    assert len(attempted_old) == 8
    assert all(
        row["stage"] == "pending_retry" and row["last_error"] == "fixture model timeout"
        for row in state["articles"].values()
        if "old-" in row["url"]
    )

    site_cursor = state["article_cursors"][site_id]
    exact_retry_url = next(
        row["url"] for row in state["articles"].values() if "old-0.pdf" in row["url"]
    )
    exact = runtime._process_articles(
        site_id,
        state,
        config,
        "unused-hub",
        preview=False,
        article_limit=1,
        include_url=exact_retry_url,
    )
    assert exact["enrichment_candidate_count"] == 1
    assert state["article_cursors"][site_id] == site_cursor
    assert len(attempted_old) == 9

    state = runtime._load_state(state_path)
    third = runtime._process_articles(
        site_id, state, config, "unused-hub", preview=False, article_limit=8
    )
    assert third["enrichment_candidate_count"] == 8
    assert len(attempted_old) == 17
    assert all(
        row["stage"] == "pending_retry"
        for row in state["articles"].values()
        if "old-" in row["url"]
    )


def test_exact_url_pending_count_stays_site_wide_in_daily_report(
    tmp_path: Path, monkeypatch
) -> None:
    site_id = "cas"
    exact_url = "https://www.casact.org/research/exact.pdf"
    articles: dict[str, dict[str, Any]] = {}
    for number in range(40):
        url = exact_url if number == 0 else f"https://www.casact.org/research/{number}.pdf"
        raw_sha = hashlib.sha256(f"article-{number}".encode()).hexdigest()
        article = {
            "url": url,
            "raw_sha256": raw_sha,
            "stage": "pending_retry",
            "discovery_sites": [site_id],
            "last_error": "previous retry still pending",
        }
        articles[runtime._file_key(url, raw_sha)] = article
    cursor = runtime._article_identity(next(iter(articles.values())))
    state_path = tmp_path / "state.json"
    monkeypatch.setattr(runtime, "_ACTIVE_STATE_PATH", state_path)
    monkeypatch.setattr(runtime, "content_hub_read_version", lambda *_args: None)
    state = {"articles": articles, "article_cursors": {site_id: cursor}, "target_acks": {}}
    runtime._atomic_json(state_path, state)

    attempted: list[str] = []

    def fail_selected_article(article, *_args, **_kwargs):
        attempted.append(article["url"])
        raise RuntimeError("fixture model timeout")

    monkeypatch.setattr(runtime, "_prepare_hub_article", fail_selected_article)
    config = {"data_root": str(tmp_path), "state_root": str(tmp_path)}

    def run_exact() -> dict[str, Any]:
        return runtime._process_articles(
            site_id,
            state,
            config,
            "unused-hub",
            preview=False,
            article_limit=1,
            include_url=exact_url,
        )

    first = run_exact()
    second = run_exact()
    assert first["candidate_count"] == second["candidate_count"] == 1
    assert first["pending_article_count"] == second["pending_article_count"] == 39
    assert attempted == [exact_url, exact_url]
    assert state["article_cursors"][site_id] == cursor
    assert sum(row["stage"] != "complete" for row in state["articles"].values()) == 40

    flow = runtime._load_flow_module()
    receipt_root = tmp_path / "receipts"

    def write_attempt(summary: dict[str, Any], pending_count: int, error: str | None = None):
        stages = {
            stage: runtime._receipt_stage("succeeded", count=0) for stage in flow.COUNT_STAGES
        }
        for stage in ("new_versions", "enriched", "ai_candidates"):
            stages[stage] = runtime._receipt_stage(
                "partial" if error else "succeeded",
                count=summary.get(stage, 0),
                error=error,
                pending_count=pending_count,
            )
        for target in ("china", "overseas"):
            stages[f"{target}_ack"] = runtime._ack_stage(summary["target"][target])
        return runtime._write_receipt(
            receipt_root,
            "2026-10-27",
            site_id,
            stages,
            [],
            error,
        )

    prior_stages = {
        stage: runtime._receipt_stage("succeeded", count=0) for stage in flow.COUNT_STAGES
    }
    prior_stages["new_versions"] = runtime._receipt_stage("succeeded", count=0, pending_count=38)
    prior_stages["enriched"] = runtime._receipt_stage("succeeded", count=0, pending_count=38)
    for target in ("china", "overseas"):
        prior_stages[f"{target}_ack"] = runtime._ack_stage(
            {"acknowledged": 0, "pending": 0, "failed": 0, "ready_status": "succeeded"}
        )
    runtime._write_receipt(receipt_root, "2026-10-27", site_id, prior_stages, [], None)

    for summary in (first, second):
        error = "; ".join(summary["article_errors"])
        receipt = write_attempt(summary, summary["pending_article_count"], error)
    assert receipt["attempt_count"] == 3
    assert receipt["stages"]["new_versions"]["pending_count"] == 39
    assert receipt["stages"]["enriched"]["pending_count"] == 39

    report = flow.make_report(
        __import__("datetime").date(2026, 10, 27), flow.read_receipts(receipt_root)
    )
    cas = next(row for row in report["site_statuses"] if row["site_id"] == site_id)
    assert cas["pending_article_count"] == 39


def test_streamed_target_command_is_shell_safe_noninteractive_and_supports_sudo(
    monkeypatch,
) -> None:
    monkeypatch.setattr(
        runtime,
        "TARGET_WORKER",
        type("WorkerFile", (), {"read_bytes": lambda self: b"worker-code"})(),
    )
    seen: dict[str, object] = {}

    def fake_run(command, *, input, text, capture_output, timeout):
        seen.update(
            {
                "command": command,
                "input": input,
                "text": text,
                "capture_output": capture_output,
                "timeout": timeout,
            }
        )
        return subprocess.CompletedProcess(
            command, 0, '{"target":"china","failures":[],"ready_status":"not_required"}\n', ""
        )

    monkeypatch.setattr(runtime.subprocess, "run", fake_run)
    result = runtime._stream_target(
        {"targets": {"china": {"host": "existing-host", "sudo": True, "container": "ai-api"}}},
        "china",
        [],
        False,
    )
    assert result["target"] == "china"
    command = seen["command"]
    assert command[:9] == [
        "ssh",
        "-o",
        "BatchMode=yes",
        "-o",
        "ConnectTimeout=20",
        "-o",
        "ServerAliveInterval=30",
        "-o",
        "ServerAliveCountMax=3",
    ]
    assert command[9] == "existing-host"
    remote_tokens = shlex.split(command[10])
    assert remote_tokens[:5] == ["sudo", "-n", "docker", "exec", "-i"]
    assert remote_tokens[5:7] == ["ai-api", "python"]
    assert remote_tokens[7] == "-c"
    assert "<local-target-worker>" in remote_tokens[8]
    lines = str(seen["input"]).splitlines()
    assert base64.b64decode(lines[0]) == b"worker-code"
    assert json.loads(lines[1])["target"] == "china"
    assert seen["timeout"] == 21600


def test_streamed_target_timeout_uses_existing_worker_failure_path(monkeypatch) -> None:
    monkeypatch.setattr(
        runtime,
        "TARGET_WORKER",
        type("WorkerFile", (), {"read_bytes": lambda self: b"worker-code"})(),
    )
    seen: dict[str, object] = {}

    def fake_run(command, *, input, text, capture_output, timeout):
        seen["command"] = command
        seen["timeout"] = timeout
        raise subprocess.TimeoutExpired(command, timeout)

    monkeypatch.setattr(runtime.subprocess, "run", fake_run)

    with pytest.raises(
        runtime.RuntimeErrorSafe,
        match="china target worker could not complete: TimeoutExpired",
    ):
        runtime._stream_target(
            {"targets": {"china": {"host": "existing-host"}}}, "china", [], False
        )

    assert seen["timeout"] == 21600
    assert "ServerAliveInterval=30" in seen["command"]
    assert "ServerAliveCountMax=3" in seen["command"]


@pytest.mark.parametrize(
    ("reply_text", "returncode", "stderr", "error_text", "error_details"),
    [
        (
            '{"schema_version":1,"error":"database unavailable"}\n',
            2,
            "",
            "database unavailable",
            None,
        ),
        (
            "",
            255,
            "ssh: connect to host configured-china-host port 22: Connection timed out",
            "china target worker returned no JSON result (exit 255)",
            {
                "exit_code": 255,
                "stderr": "ssh: connect to host configured-china-host port 22: Connection timed out",
            },
        ),
    ],
)
def test_target_worker_startup_error_keeps_configured_identity_and_receipt_evidence(
    tmp_path: Path,
    monkeypatch,
    reply_text: str,
    returncode: int,
    stderr: str,
    error_text: str,
    error_details: dict[str, Any] | None,
) -> None:
    monkeypatch.setattr(runtime, "_ACTIVE_STATE_PATH", tmp_path / "runtime-state.json")
    raw = b"target startup failure source bytes"
    url = "https://example.test/target-startup-failure.pdf"
    markdown = "Actuarial AI source text used for target startup failure verification."
    raw_sha = hashlib.sha256(raw).hexdigest()
    artifact = tmp_path / "target-startup-failure.pdf"
    artifact.write_bytes(raw)
    item = {
        "url": url,
        "source_site": "iaa",
        "raw_sha256": raw_sha,
        "content_version_id": 211,
        "markdown_text_variant_id": 311,
        "markdown_sha256": hashlib.sha256(markdown.encode()).hexdigest(),
        "markdown": markdown,
        "bilingual_metadata": {
            language: {
                "title": "AI actuarial source",
                "summary": "Actuarial AI source summary.",
                "category": "AI",
                "keywords": ["actuarial", "AI"],
            }
            for language in ("zh", "en")
        },
        "categories": [{"key": "AI", "evidence": ["Actuarial AI source"]}],
        "date_facts": {},
        "ai_candidate": True,
        "artifact_path": str(artifact),
    }
    key = runtime._file_key(url, raw_sha)
    article = {**item, "site_id": "iaa", "stage": "downloaded", "discovery_sites": ["iaa"]}
    state = {"articles": {key: article}, "target_acks": {}}
    monkeypatch.setattr(
        runtime, "_prepare_hub_article", lambda *_args, **_kwargs: (item, False, False)
    )
    seen_commands: list[list[str]] = []
    real_stream_target = runtime._stream_target

    def target_transport(config, target, items, ready_pending, **kwargs):
        if target == "china":
            return real_stream_target(config, target, items, ready_pending, **kwargs)
        return {
            "target": target,
            "acknowledgements": [
                {
                    "content_version_id": item["content_version_id"],
                    "raw_sha256": item["raw_sha256"],
                    "metadata_fingerprint": runtime._target_metadata_fingerprint(item),
                    "status": "written",
                }
            ],
            "failures": [],
            "ready_status": "succeeded",
        }

    def fake_ssh(command, **_kwargs):
        seen_commands.append(command)
        return subprocess.CompletedProcess(command, returncode, reply_text, stderr)

    monkeypatch.setattr(runtime, "_stream_target", target_transport)
    config = {
        "data_root": str(tmp_path),
        "state_root": str(tmp_path),
        "targets": {"china": {"host": "configured-china-host"}},
    }
    with monkeypatch.context() as context:
        context.setattr(runtime.subprocess, "run", fake_ssh)
        summary = runtime._process_articles(
            "iaa", state, config, "unused-hub", preview=False, article_limit=1
        )

    assert len(seen_commands) == 1
    assert summary["target"]["china"]["failed"] == 1
    assert summary["target"]["china"]["acknowledged"] == 0
    assert error_text in summary["target"]["china"]["ready_error"]
    evidence = json.loads(summary["target"]["china"]["ready_evidence"][0])
    expected_evidence = {
        "configured_target": "china",
        "error": error_text,
        "worker_target": None,
    }
    if error_details is not None:
        expected_evidence["error_details"] = error_details
    assert evidence == expected_evidence
    china_ack = runtime._ack_stage(summary["target"]["china"])
    assert china_ack["status"] == "failed"
    assert china_ack["error"] == summary["target"]["china"]["ready_error"]

    flow = runtime._load_flow_module()
    stages = {stage: runtime._receipt_stage("succeeded", count=0) for stage in flow.COUNT_STAGES}
    stages.update(
        {
            "china_ack": china_ack,
            "overseas_ack": runtime._ack_stage(summary["target"]["overseas"]),
        }
    )
    receipt_root = tmp_path / "receipts"
    receipt = runtime._write_receipt(receipt_root, "2026-10-26", "iaa", stages, [], reason=None)
    assert receipt["run_status"] == "failed"
    assert error_text in receipt["reason"]
    assert receipt["stages"]["china_ack"]["ready_evidence"] == [
        summary["target"]["china"]["ready_evidence"][0]
    ]
    report_process = subprocess.run(
        [
            sys.executable,
            str(runtime.FLOW_SCRIPT),
            "report",
            "--date",
            "2026-10-26",
            "--receipts",
            str(receipt_root),
            "--json",
        ],
        check=False,
        capture_output=True,
        text=True,
    )
    assert report_process.returncode == 0, report_process.stderr
    ready_case = next(
        case
        for case in json.loads(report_process.stdout)["failure_cases"]
        if case["site_id"] == "iaa" and case["stage"] == "china_ready"
    )
    assert error_text in ready_case["error"]
    diagnostic_text = ready_case["evidence"][0]
    if diagnostic_text.startswith("ai: "):
        diagnostic_text = diagnostic_text.partition(": ")[2]
    diagnostic = json.loads(diagnostic_text)
    assert diagnostic["configured_target"] == "china"
    if error_details is not None:
        assert diagnostic["error_details"] == error_details


@pytest.mark.parametrize(
    ("site_id", "document_url", "referrer", "title"),
    [
        (
            "cia-research-publications",
            "https://www.cia-ica.ca/app/themes/wicket/custom/dl_file.php?fid=840377&p=840374",
            "https://www.cia-ica.ca/publications/",
            "AI actuarial research paper",
        ),
        (
            "cca",
            "https://www.ccactuaries.org/docs/default-source/papers/evaluating-ai-health-solutions.pdf",
            "https://www.ccactuaries.org/professional-resources/papers",
            "Evaluating AI health solutions research paper",
        ),
        (
            "france",
            "https://www.institutdesactuaires.com/global/gene/link.php?doc_id=20148&fg=1",
            "https://www.institutdesactuaires.com/article/conferences-1105",
            "The human factor in AI adoption",
        ),
    ],
)
def test_pending_document_priority_preserves_prefix_and_nonpriority_order(
    site_id: str, document_url: str, referrer: str, title: str
) -> None:
    site = runtime.SITES[site_id]
    prefix = [
        {"url": f"https://{site.hosts[0]}/already-visited-{number}", "kind": "page"}
        for number in range(2)
    ]
    nav_before = {"url": f"https://{site.hosts[0]}/research-nav-before", "kind": "page"}
    private = {
        "url": f"https://{site.hosts[0]}/private-files/research.pdf",
        "kind": "document",
        "referrer": referrer,
        "title": title,
    }
    document = {
        "url": document_url,
        "kind": "document",
        "referrer": referrer,
        "title": title,
    }
    nav_after = {"url": f"https://{site.hosts[0]}/research-nav-after", "kind": "page"}
    queue = [*prefix, nav_before, private, document, nav_after]

    runtime._prioritize_pending_documents(site, queue, cursor=len(prefix))

    assert queue[: len(prefix)] == prefix
    assert queue[len(prefix) :] == [document, nav_before, private, nav_after]

    swiss_queue = [*prefix, nav_before, private, document, nav_after]
    original_swiss_queue = list(swiss_queue)
    runtime._prioritize_pending_documents(
        runtime.SITES["swiss-association-of-actuaries"],
        swiss_queue,
        cursor=len(prefix),
    )
    assert swiss_queue == original_swiss_queue


class _PriorityQueueCrawler:
    def __init__(self, page_bodies: dict[str, bytes] | None = None) -> None:
        self.page_bodies = page_bodies or {}
        self.site_attempts = 0
        self.site_budget = 0
        self.calls: list[str] = []

    def begin_site(self, budget: int) -> None:
        self.site_attempts = 0
        self.site_budget = budget

    def _request(self, url: str, **_kwargs):
        self.site_attempts += 1
        self.calls.append(url)
        body = self.page_bodies.get(
            url,
            (
                "<html><body><h1>Research publications</h1><p>"
                + "Verified actuarial research listing. " * 8
                + "</p></body></html>"
            ).encode(),
        )
        return body, {"content-type": "text/html"}, url

    def _download_file(self, url: str, staging_dir: Path, **_kwargs):
        self.site_attempts += 1
        self.calls.append(url)
        raw = b"%PDF-1.7\nverified public research attachment"
        digest = hashlib.sha256(raw).hexdigest()
        staging_dir.mkdir(parents=True, exist_ok=True)
        staging_path = staging_dir / f"{digest}.pdf"
        staging_path.write_bytes(raw)
        return (
            staging_path,
            {"content-type": "application/pdf"},
            url,
            digest,
            len(raw),
        )


def test_france_v2_bootstrap_migrates_v1_queue_and_runs_before_pending_documents(
    tmp_path: Path, monkeypatch
) -> None:
    monkeypatch.setattr(capacity, "ensure_capacity", lambda **_kwargs: None)
    monkeypatch.setattr(runtime, "_ACTIVE_STATE_PATH", tmp_path / "runtime-state.json")
    site = runtime.SITES["france"]
    prefix = [
        {"url": f"https://www.institutdesactuaires.com/visited/{n}", "kind": "page"}
        for n in range(2)
    ]
    old_paths = [
        "https://www.institutdesactuaires.com/se-documenter/memoire-d-actuariat-38",
        "https://www.institutdesactuaires.com/se-documenter/memoires-38",
        "https://www.institutdesactuaires.com/global/gene/link.php?doc_id=1&fg=1",
    ]
    documents = [
        {
            "url": f"https://www.institutdesactuaires.com/docs/research-{n}.pdf",
            "kind": "document",
            "title": "Mémoire actuariel",
            "referrer": "https://www.institutdesactuaires.com/article/conferences-1105",
        }
        for n in range(5)
    ]
    valid_retry = {
        "url": "https://www.institutdesactuaires.com/article/conferences-1105",
        "kind": "page",
        "attempts": 1,
        "last_error": "kept retry",
    }
    queue = [*prefix, *[{"url": url, "kind": "page"} for url in old_paths], *documents]
    prefix_before = json.loads(json.dumps(prefix))
    documents_before = json.loads(json.dumps(documents))
    site_state = {
        "queue": queue,
        "cursor": len(prefix),
        "visited": [runtime._queue_identity(row) for row in prefix],
        "retry_queue": [
            {"url": old_paths[1], "kind": "page", "attempts": 2},
            valid_retry,
        ],
        "frontier_repair_france_v1": True,
    }
    state = {"articles": {"prior": {"stage": "complete", "url": "https://example.test/a"}}}
    crawler = _PriorityQueueCrawler()

    result = runtime._capture(site, state, crawler, tmp_path / "data", site_state, request_limit=1)

    assert crawler.calls == [site.seeds[0]]
    assert result["request_count"] == 1
    assert site_state["frontier_repair_france_v1"] is True
    assert site_state["frontier_repair_france_v2"] is True
    assert site_state["cursor"] == len(prefix) + 1
    assert site_state["queue"][: len(prefix)] == prefix_before
    assert [
        row["url"]
        for row in site_state["queue"][len(prefix) + 1 : len(prefix) + 1 + len(documents)]
    ] == [row["url"] for row in documents_before]
    assert not set(old_paths).intersection(row["url"] for row in site_state["queue"])
    assert not set(old_paths).intersection(row["url"] for row in site_state.get("retry_queue", []))
    preserved_retry = next(row for row in site_state["queue"] if row["url"] == valid_retry["url"])
    assert preserved_retry["attempts"] == valid_retry["attempts"]
    assert preserved_retry["last_error"] == valid_retry["last_error"]
    assert sum(row["url"] == site.seeds[0] for row in site_state["queue"]) == 1
    assert state["articles"]["prior"]["stage"] == "complete"

    migrated = json.loads(json.dumps(site_state))
    runtime._repair_saved_frontier(site, site_state)
    assert site_state == migrated


@pytest.mark.parametrize("placement", ["queue", "retry", "visited"])
def test_france_v2_migration_promotes_existing_seed_or_skips_visited_seed(
    placement: str,
) -> None:
    site = runtime.SITES["france"]
    current = {"url": site.seeds[0], "kind": "page", "depth": 1}
    already_visited = placement == "visited"
    site_state = {
        "queue": [current] if placement != "retry" else [],
        "cursor": 1 if already_visited else 0,
        "visited": [runtime._queue_identity(current)] if already_visited else [],
        "retry_queue": [current] if placement == "retry" else [],
        "frontier_repair_france_v1": True,
    }

    runtime._repair_saved_frontier(site, site_state)

    assert site_state["frontier_repair_france_v2"] is True
    copies = [*site_state["queue"], *site_state["retry_queue"]]
    assert sum(row["url"] == site.seeds[0] for row in copies) == 1
    if already_visited:
        assert "frontier_priority" not in current
    else:
        assert current["frontier_priority"] == "france_memoirs_v2"
        assert current["depth"] == 0


def test_france_current_memoir_listing_discovers_public_detail_pdf(
    tmp_path: Path, monkeypatch
) -> None:
    monkeypatch.setattr(capacity, "ensure_capacity", lambda **_kwargs: None)
    monkeypatch.setattr(runtime, "_ACTIVE_STATE_PATH", tmp_path / "runtime-state.json")
    site = runtime.SITES["france"]
    listing_url = site.seeds[0]
    detail_url = f"{listing_url}?id=4846c6057c963bcd1c19c2921023050a"
    document_url = (
        "https://www.institutdesactuaires.com/docs/mem/" "92120aac7893c2700f17b737734ddad1.pdf"
    )
    listing = (
        '<html><body><h1>Mémoires d’actuariat</h1><a href="?id=4846c6057c963bcd1c19c2921023050a">'
        "Mémoire d’actuariat : intelligence artificielle</a></body></html>"
    ).encode()
    detail = (
        "<html><body><h1>Mémoire d’actuariat</h1><p>"
        + "Recherche actuarielle et intelligence artificielle. " * 6
        + f'<a href="{document_url}">Télécharger le mémoire</a></p></body></html>'
    ).encode()
    crawler = _PriorityQueueCrawler({listing_url: listing, detail_url: detail})
    state = {"articles": {}}
    site_state: dict[str, Any] = {"queue": [], "cursor": 0, "visited": [], "retry_queue": []}

    result = runtime._capture(
        site,
        state,
        crawler,
        tmp_path / "data",
        site_state,
        request_limit=3,
        initial_frontier=[{"url": listing_url, "kind": "page", "depth": 0}],
    )

    assert crawler.calls == [listing_url, detail_url, document_url]
    assert result["request_count"] == 3
    assert [article["url"] for article in result["downloads"]] == [detail_url, document_url]
    assert site_state["cursor"] == 3


def test_existing_public_document_jumps_over_navigation_and_restarts_without_duplicates(
    tmp_path: Path, monkeypatch
) -> None:
    monkeypatch.setattr(capacity, "ensure_capacity", lambda **_kwargs: None)
    monkeypatch.setattr(runtime, "_ACTIVE_STATE_PATH", tmp_path / "runtime-state.json")
    site = runtime.SITES["cia-research-publications"]
    document_url = "https://www.cia-ica.ca/app/themes/wicket/custom/dl_file.php?fid=840377&p=840374"
    prefix = [
        {"url": f"https://www.cia-ica.ca/visited-prefix-{number}", "kind": "page"}
        for number in range(2)
    ]
    prefix_identities = {runtime._queue_identity(item) for item in prefix}
    navigation = [
        {
            "url": f"https://www.cia-ica.ca/publications/archive-{number}/",
            "kind": "page",
            "depth": 1,
            "referrer": "https://www.cia-ica.ca/publications/",
            "title": f"Research publications {number}",
        }
        for number in range(100)
    ]
    document = {
        "url": document_url,
        "kind": "document",
        "depth": 1,
        "referrer": "https://www.cia-ica.ca/publications/",
        "title": "AI actuarial research paper",
    }
    retry = {
        "url": "https://www.cia-ica.ca/publications/retry-page/",
        "kind": "page",
        "depth": 1,
        "referrer": "https://www.cia-ica.ca/publications/",
        "title": "Research publications retry",
        "last_error": "previous timeout",
        "attempts": 1,
    }
    queue = [*prefix, *navigation, document]
    site_state = {
        "queue": queue,
        "cursor": len(prefix),
        "visited": sorted(prefix_identities),
        "retry_queue": [retry],
        "frontier_repair_cia-research-publications_v1": True,
    }
    state = {"articles": {}}
    prefix_before = json.loads(json.dumps(queue[: len(prefix)]))
    retry_before = dict(retry)
    first_crawler = _PriorityQueueCrawler()

    first = runtime._capture(
        site,
        state,
        first_crawler,
        tmp_path / "data",
        site_state,
        request_limit=30,
    )

    page_urls = [item["url"] for item in navigation]
    assert first["request_count"] == 30
    assert first["budget_exhausted"] is True
    assert first_crawler.calls == [document_url, *page_urls[:29]]
    assert len(first_crawler.calls) == 30
    assert site_state["cursor"] == len(prefix) + 30
    assert site_state["queue"][: len(prefix)] == prefix_before
    assert prefix_identities.issubset(set(site_state["visited"]))
    retry_after = next(item for item in site_state["queue"] if item["url"] == retry["url"])
    assert retry_after["last_error"] == retry_before["last_error"]
    assert retry_after["attempts"] == retry_before["attempts"]
    assert len(state["articles"]) == 30

    second_crawler = _PriorityQueueCrawler()
    second = runtime._capture(
        site,
        state,
        second_crawler,
        tmp_path / "data",
        site_state,
        request_limit=30,
    )
    assert second["request_count"] == 30
    assert second_crawler.calls == page_urls[29:59]
    assert len(set(first_crawler.calls + second_crawler.calls)) == 60
    assert site_state["cursor"] == len(prefix) + 60
    assert site_state["queue"][: len(prefix)] == prefix_before


def test_document_discovered_from_page_is_next_and_restart_skips_visited_work(
    tmp_path: Path, monkeypatch
) -> None:
    monkeypatch.setattr(capacity, "ensure_capacity", lambda **_kwargs: None)
    monkeypatch.setattr(runtime, "_ACTIVE_STATE_PATH", tmp_path / "runtime-state.json")
    site = runtime.SITES["cia-research-publications"]
    page_url = "https://www.cia-ica.ca/publications/current/"
    document_url = "https://www.cia-ica.ca/app/themes/wicket/custom/dl_file.php?fid=840377&p=840374"
    navigation_urls = [
        f"https://www.cia-ica.ca/publications/archive-{number}/" for number in range(2)
    ]
    page_body = (
        f'<html><body><h1>Research publications</h1><a href="{document_url}">'
        "AI actuarial research paper</a><p>"
        + "Verified actuarial research. " * 8
        + "</p></body></html>"
    ).encode()
    crawler = _PriorityQueueCrawler({page_url: page_body})
    site_state = {
        "queue": [
            {
                "url": page_url,
                "kind": "page",
                "depth": 1,
                "referrer": "https://www.cia-ica.ca/publications/",
                "title": "Research publications",
            },
            *[
                {
                    "url": url,
                    "kind": "page",
                    "depth": 1,
                    "referrer": page_url,
                    "title": "Research publications archive",
                }
                for url in navigation_urls
            ],
        ],
        "cursor": 0,
        "visited": [],
        "frontier_repair_cia-research-publications_v1": True,
    }
    state = {"articles": {}}

    first = runtime._capture(
        site,
        state,
        crawler,
        tmp_path / "data",
        site_state,
        request_limit=2,
    )
    assert crawler.calls == [page_url, document_url]
    assert first["request_count"] == 2
    assert first["budget_exhausted"] is True
    assert site_state["queue"][site_state["cursor"] :]
    assert [item["url"] for item in site_state["queue"][site_state["cursor"] :]] == navigation_urls

    resumed_crawler = _PriorityQueueCrawler()
    resumed = runtime._capture(
        site,
        state,
        resumed_crawler,
        tmp_path / "data",
        site_state,
        request_limit=2,
    )
    assert resumed_crawler.calls == navigation_urls
    assert resumed["request_count"] == 2
    assert set(crawler.calls).isdisjoint(resumed_crawler.calls)
